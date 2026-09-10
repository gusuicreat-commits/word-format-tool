"""Synthetic papers: verify saved OOXML, actual styles, and execution reports."""

import base64
import json
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt
from lxml import etree

import format_docx as formatter


def element(tag, text=None, **attributes):
    node = OxmlElement(tag)
    if text is not None:
        node.text = text
    for key, value in attributes.items():
        node.set(qn(f"w:{key}"), value)
    return node


def add_link(paragraph, text="Reference link"):
    link = element("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to("https://example.org/paper", RT.HYPERLINK, is_external=True))
    run = element("w:r")
    run.append(element("w:t", text))
    link.append(run)
    paragraph._p.append(link)


def add_numbering(paragraph):
    numbering = element("w:numPr")
    numbering.append(element("w:ilvl", val="0"))
    numbering.append(element("w:numId", val="1"))
    paragraph._p.get_or_add_pPr().append(numbering)


def make_paper():
    doc = Document()
    for text in [
        "Synthetic research paper", "摘要", "本文研究 AI 2026 在排版中的应用。",
        "关键词：排版；验证", "Abstract", "This study checks document formatting.",
        "Keywords: formatting; verification",
    ]:
        doc.add_paragraph(text)
    heading = doc.add_paragraph("Research methods", style="Heading 1")
    add_numbering(heading)
    doc.add_paragraph("1.1 方法", style="Heading 2")
    doc.add_paragraph("1.1.1 数据", style="Heading 3")
    doc.add_paragraph("普通正文包含 AI 2026 和中文。")
    link = doc.add_paragraph("来源 2026：")
    add_link(link)
    field = doc.add_paragraph("Cross reference 1: ")
    simple = element("w:fldSimple", instr=" REF target ")
    run = element("w:r")
    run.append(element("w:t", "1"))
    simple.append(run)
    field._p.append(simple)
    equation = doc.add_paragraph("Equation 1: ")
    math = element("m:oMath")
    math_run = element("m:r")
    math_run.append(element("m:t", "x=1"))
    math.append(math_run)
    equation._p.append(math)
    bookmark = doc.add_paragraph("Bookmark target 2026")
    bookmark._p.insert(0, element("w:bookmarkStart", id="0", name="target"))
    bookmark._p.append(element("w:bookmarkEnd", id="0"))
    image = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aL1sAAAAASUVORK5CYII=")
    doc.add_paragraph().add_run().add_picture(BytesIO(image), width=Inches(0.5))
    doc.add_paragraph("图1 测试图")
    doc.add_paragraph("表1 测试表")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "合并表头"
    table.cell(1, 0).text = "数据 A"
    table.cell(1, 1).text = "数据 B"
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] Smith. Formatting methods. 2026.")
    doc.add_paragraph("附录")
    doc.add_paragraph("附录正文。")
    doc.sections[0].header.paragraphs[0].text = "Synthetic header"
    footer = doc.sections[0].footer.paragraphs[0]
    page = element("w:fldSimple", instr=" PAGE ")
    footer._p.append(page)
    return doc


def protected_snapshot(file):
    with ZipFile(file) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        paragraphs = ["".join(p.xpath(".//w:t/text() | .//m:t/text()", namespaces=root.nsmap))
                      for p in root.iter(qn("w:p"))]
        names = {"hyperlink", "fldSimple", "fldChar", "instrText", "oMath", "drawing",
                 "bookmarkStart", "bookmarkEnd", "numPr", "gridSpan", "vMerge", "tblGrid", "tab", "br",
                 "footnoteReference", "endnoteReference", "pict", "object"}
        structures = [etree.tostring(node, method="c14n") for node in root.iter()
                      if formatter.get_xml_local_name(node) in names]
        parts = {name: archive.read(name) for name in archive.namelist()
                 if not name.endswith("/") and (
                     "/media/" in "/" + name
                     or name.startswith(("word/header", "word/footer", "word/footnotes", "word/endnotes", "word/comments", "word/embeddings/"))
                     or name.endswith(".rels") or name == "word/numbering.xml")}
        return paragraphs, structures, parts


class DocumentCorpusTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def execute(self, doc, override=None, template="default"):
        source, output = self.root / "input.docx", self.root / "output.docx"
        doc.save(source)
        override_path = None
        if override:
            override_path = self.root / "override.json"
            override_path.write_text(json.dumps(override), encoding="utf-8")
        rules = formatter.get_final_format_rules(template, override_path, print_warnings=False)
        report = formatter.init_report(source, output, rules)
        formatter.format_document(source, output, rules, report, overwrite=True)
        return Document(output), report, source, output

    def test_representative_paper_preserves_content_and_package_structures(self):
        for template in ("default", "course_paper"):
            with self.subTest(template=template):
                output, report, source_path, output_path = self.execute(make_paper(), {
                    "styles": {"body": {"size_pt": 13, "line_spacing": 2}},
                    "latin_digit_format": {"font": "Times New Roman", "scope": "body"},
                }, template)
                self.assertEqual(protected_snapshot(source_path), protected_snapshot(output_path))
                body = next(p for p in output.paragraphs if p.text.startswith("普通正文"))
                self.assertEqual(body.paragraph_format.line_spacing, 2)
                self.assertTrue(all(r.font.size.pt == 13 for r in body.runs))
                self.assertEqual(report["stats"]["table"], 1)
                for key in ("abstract_cn", "abstract_en", "keywords_cn", "keywords_en", "heading"):
                    self.assertEqual(report["module_status"][key]["action"], "formatted")
                self.assertTrue(any("复杂结构" in w["message"] for w in report["warnings"]))

    def test_toc_style_paragraph_is_really_protected(self):
        doc = Document()
        doc.add_paragraph("Paper title")
        doc.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
        toc = doc.add_paragraph("Research methods\t3", style="TOC 1")
        toc.runs[0].font.size = Pt(9)
        toc.paragraph_format.left_indent = Pt(18)
        before = etree.tostring(toc._p)
        doc.add_paragraph("Research methods", style="Heading 1")
        output, report, _, _ = self.execute(doc)
        self.assertEqual(etree.tostring(output.paragraphs[1]._p), before)
        self.assertEqual(report["module_status"]["toc"]["action"], "protected")
        entry = next(p for p in report["paragraphs"] if p["index"] == 1)
        self.assertIsNone(entry["applied_style"])

    def test_toc_field_without_toc_style_is_protected(self):
        doc = Document()
        doc.add_paragraph("Paper title")
        toc = doc.add_paragraph()
        for kind in ("begin", "separate"):
            toc.add_run()._r.append(element("w:fldChar", fldCharType=kind))
            if kind == "begin":
                toc.add_run()._r.append(element("w:instrText", ' TOC \\o "1-3" '))
        toc.add_run("Research methods 3").font.size = Pt(9)
        toc.add_run()._r.append(element("w:fldChar", fldCharType="end"))
        before = etree.tostring(toc._p)
        output, report, _, _ = self.execute(doc)
        self.assertEqual(etree.tostring(output.paragraphs[1]._p), before)
        self.assertEqual(report["module_status"]["toc"]["action"], "protected")

    def test_multiparagraph_toc_with_nested_field_preserves_body_boundary(self):
        doc = Document()
        doc.add_paragraph("Paper title")
        start = doc.add_paragraph()
        start.add_run()._r.append(element("w:fldChar", fldCharType="begin"))
        start.add_run()._r.append(element("w:instrText", " TO"))
        start.add_run()._r.append(element("w:instrText", 'C \\o "1-3" '))
        start.add_run()._r.append(element("w:fldChar", fldCharType="separate"))
        entry = doc.add_paragraph("Research methods ")
        entry.add_run()._r.append(element("w:fldChar", fldCharType="begin"))
        entry.add_run()._r.append(element("w:instrText", " PAGEREF target "))
        entry.add_run()._r.append(element("w:fldChar", fldCharType="separate"))
        entry.add_run("3")
        entry.add_run()._r.append(element("w:fldChar", fldCharType="end"))
        end = doc.add_paragraph()
        end.add_run()._r.append(element("w:fldChar", fldCharType="end"))
        before = [etree.tostring(p._p) for p in doc.paragraphs[1:]]
        doc.add_paragraph("Body after TOC")
        output, report, _, _ = self.execute(doc, {"styles": {"body": {"size_pt": 13}}})
        self.assertEqual([etree.tostring(p._p) for p in output.paragraphs[1:4]], before)
        self.assertEqual(output.paragraphs[4].runs[0].font.size.pt, 13)
        self.assertEqual(report["module_status"]["toc"]["count"], 3)

    def test_detailed_styles_and_report_match_saved_document(self):
        output, report, source, target = self.execute(make_paper(), {"styles": {
            "abstract_cn_content": {"size_pt": 11},
            "abstract_en_content": {"size_pt": 10},
            "keywords_cn_label": {"bold": True, "size_pt": 12},
            "keywords_cn_content": {"bold": False, "size_pt": 11},
            "keywords_en_label": {"bold": True, "size_pt": 12},
            "keywords_en_content": {"bold": False, "size_pt": 10},
            "reference_item": {"size_pt": 10.5, "hanging_indent_chars": 2},
        }})
        self.assertEqual(protected_snapshot(source), protected_snapshot(target))
        by_text = {p.text: p for p in output.paragraphs}
        self.assertEqual(by_text["本文研究 AI 2026 在排版中的应用。"].runs[0].font.size.pt, 11)
        self.assertEqual(by_text["This study checks document formatting."].runs[0].font.size.pt, 10)
        for text, content_size in (("关键词：排版；验证", 11), ("Keywords: formatting; verification", 10)):
            paragraph = by_text[text]
            self.assertTrue(paragraph.runs[0].bold)
            self.assertFalse(paragraph.runs[-1].bold)
            self.assertEqual(paragraph.runs[-1].font.size.pt, content_size)
        reference = by_text["[1] Smith. Formatting methods. 2026."]
        self.assertEqual(reference.paragraph_format.first_line_indent.pt, -21)
        self.assertEqual(reference.paragraph_format.left_indent.pt, 21)
        appendix = next(p for p in report["paragraphs"] if p["text_preview"] == "附录正文。")
        self.assertEqual(appendix["detected_type"], "body")
        counts = report["final_rules_debug"]["debug_summary"]["execution_counts"]
        self.assertEqual(counts["keywords_cn_label_count"], 1)
        self.assertEqual(counts["keywords_en_content_count"], 1)
        self.assertEqual(counts["reference_hanging_indent"], 1)

    def test_merged_cell_nested_table_is_counted_once(self):
        doc = Document()
        doc.add_paragraph("Paper title")
        table = doc.add_table(rows=1, cols=2)
        cell = table.cell(0, 0).merge(table.cell(0, 1))
        cell.add_table(rows=1, cols=1).cell(0, 0).text = "Nested data"
        _, report, source, output = self.execute(doc)
        self.assertEqual(protected_snapshot(source), protected_snapshot(output))
        self.assertEqual(report["stats"]["table"], 2)
        self.assertEqual(len(report["tables"]), 2)

    def test_chinese_abstract_specific_rule_inherits_legacy_defaults(self):
        doc = Document()
        doc.add_paragraph("Paper title")
        doc.add_paragraph("摘要")
        doc.add_paragraph("摘要正文。")
        output, report, _, _ = self.execute(doc, {"styles": {"abstract_cn_title": {"size_pt": 17}}})
        title = output.paragraphs[1]
        self.assertEqual(title.runs[0].font.size.pt, 17)
        self.assertTrue(title.runs[0].bold)
        self.assertEqual(title.alignment, 1)
        self.assertEqual(output.paragraphs[2].runs[0].font.size.pt, 12)
        entry = next(p for p in report["paragraphs"] if p["index"] == 1)
        self.assertEqual(entry["applied_style"], "abstract_cn_title")

    def test_complex_mixed_run_is_not_reported_as_fully_formatted(self):
        doc = Document()
        doc.add_paragraph("Paper title")
        paragraph = doc.add_paragraph("中文 AI\t2026\n正文")
        _, report, source, output = self.execute(doc, {
            "styles": {"body": {"size_pt": 12}},
            "latin_digit_format": {"font": "Times New Roman", "size_pt": 9, "scope": "body"},
        })
        self.assertEqual(protected_snapshot(source), protected_snapshot(output))
        self.assertEqual(Document(output).paragraphs[1].runs[0].font.size.pt, 12)
        self.assertEqual(report["module_status"]["latin_digit_format"]["action"], "skipped")

    def test_real_paper_body_and_english_references_keep_numbering(self):
        for reference_title in ("References", "REFERENCES:", "Bibliography"):
            with self.subTest(reference_title=reference_title):
                doc = Document()
                doc.add_paragraph("Paper title")
                body = doc.add_paragraph("Data source", style="List Number")
                add_numbering(body)
                doc.add_paragraph(reference_title)
                reference = doc.add_paragraph("Author. Article. 2026.", style="List Number")
                add_numbering(reference)
                doc.add_paragraph("Appendix A")
                doc.add_paragraph("Supporting data")
                output, report, source, target = self.execute(doc)
                self.assertEqual(protected_snapshot(source), protected_snapshot(target))
                self.assertEqual(output.paragraphs[1].style.name, "List Number")
                self.assertEqual(output.paragraphs[3].style.name, "List Number")
                self.assertEqual(report["stats"]["reference_title"], 1)
                self.assertEqual(report["stats"]["reference_item"], 1)
                self.assertEqual(report["paragraphs"][-1]["detected_type"], "body")

    def test_caption_reference_prose_keeps_body_style_and_numbering(self):
        for text in ("图3-2为硬件时序图，图中的信号用于选择数据。", "表1显示各组结果，下面进行分析。"):
            with self.subTest(text=text):
                doc = Document()
                doc.add_paragraph("Paper title")
                paragraph = doc.add_paragraph(text)
                add_numbering(paragraph)
                output, report, source, target = self.execute(doc)
                self.assertEqual(protected_snapshot(source), protected_snapshot(target))
                self.assertEqual(report["paragraphs"][1]["detected_type"], "body")
                self.assertEqual(output.paragraphs[1].runs[0].font.size.pt, 12)

    def test_long_bold_title_and_named_sections(self):
        doc = Document()
        doc.add_paragraph().add_run("Trials and tribulations of implementing a complex intervention in everyday clinical practice").bold = True
        doc.add_paragraph("Alice Smith and Bob Jones").runs[0].bold = True
        doc.add_paragraph("Methods")
        doc.add_paragraph("Participants")
        doc.add_paragraph("The methods used in this study are described below.")
        _, report, _, _ = self.execute(doc)
        self.assertEqual([p["detected_type"] for p in report["paragraphs"]],
                         ["paper_title", "body", "heading_1", "heading_2", "body"])

    def test_explicit_line_spacing_does_not_snap_to_source_grid(self):
        doc = Document()
        doc.add_paragraph("Paper title")
        paragraph = doc.add_paragraph("Body text under a document grid.")
        paragraph._p.get_or_add_pPr().append(element("w:snapToGrid"))
        doc.sections[0]._sectPr.append(element("w:docGrid", type="lines", linePitch="600"))
        output, _, _, _ = self.execute(doc)
        self.assertEqual(output.paragraphs[1]._p.pPr.find(qn("w:snapToGrid")).get(qn("w:val")), "0")
        self.assertEqual(output.paragraphs[1].paragraph_format.line_spacing, 1.5)

    def test_structured_abstract_labels_are_not_main_sections(self):
        doc = Document()
        for text in ("Paper title", "Abstract", "Background", "Study context.", "Methods", "Study methods.",
                     "Results", "Study results.", "Keywords: study", "Methods", "Main body."):
            doc.add_paragraph(text)
        _, report, _, _ = self.execute(doc)
        types = [p["detected_type"] for p in report["paragraphs"]]
        self.assertEqual(types[2:8], ["abstract_en_content"] * 6)
        self.assertEqual(types[9:], ["heading_1", "body"])

    def test_unstructured_abstract_ends_at_results(self):
        doc = Document()
        for text in ("Paper title", "ABSTRACT", "Summary of the research.", "RESULTS", "Main body."):
            doc.add_paragraph(text)
        _, report, _, _ = self.execute(doc)
        self.assertEqual([p["detected_type"] for p in report["paragraphs"]][2:],
                         ["abstract_en_content", "heading_1", "body"])

    def test_bold_english_numbered_heading_not_plain_list(self):
        doc = Document()
        doc.add_paragraph("Paper title")
        doc.add_paragraph("1. Research summary").runs[0].bold = True
        doc.add_paragraph("2. Collect the data")
        doc.add_paragraph("3. This is a complete sentence.").runs[0].bold = True
        _, report, _, _ = self.execute(doc)
        self.assertEqual([p["detected_type"] for p in report["paragraphs"]][1:],
                         ["heading_1", "body", "body"])

    def test_bracketed_chinese_abstract_and_keywords(self):
        doc = Document()
        doc.add_paragraph("论文标题")
        doc.add_paragraph("【摘  要】本文研究通信系统。")
        doc.add_paragraph("【关键词】通信；系统")
        output, report, source, target = self.execute(doc, {"styles": {
            "keywords_cn_label": {"bold": True}, "keywords_cn_content": {"bold": False},
        }})
        self.assertEqual(protected_snapshot(source), protected_snapshot(target))
        self.assertEqual(report["paragraphs"][1]["detected_type"], "abstract_content")
        self.assertEqual(report["paragraphs"][2]["detected_type"], "keywords")
        self.assertEqual(output.paragraphs[2].runs[0].text, "【关键词】")
        self.assertTrue(output.paragraphs[2].runs[0].bold)


if __name__ == "__main__":
    unittest.main()
