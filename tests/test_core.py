import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

import format_docx


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def make_context():
    return {
        "first_non_empty_seen": True,
        "in_reference_section": False,
        "in_abstract_section": False,
        "stats": {},
    }


def add_direct_numbering(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = format_docx.OxmlElement("w:numPr")
    ilvl = format_docx.OxmlElement("w:ilvl")
    ilvl.set(format_docx.qn("w:val"), "0")
    num_id = format_docx.OxmlElement("w:numId")
    num_id.set(format_docx.qn("w:val"), "1")
    num_pr.append(ilvl)
    num_pr.append(num_id)
    p_pr.append(num_pr)


def has_direct_numbering(paragraph):
    p_pr = paragraph._p.pPr
    return p_pr is not None and p_pr.find(format_docx.qn("w:numPr")) is not None


def get_run_font_mapping(run, field):
    r_pr = run._r.rPr
    if r_pr is None or r_pr.rFonts is None:
        return None
    return r_pr.rFonts.get(format_docx.qn(f"w:{field}"))


class CoreTests(unittest.TestCase):
    def format_paragraphs_for_report(self, paragraphs, override=None, add_table=False):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            input_path = temp_path / "input.docx"
            output_path = temp_path / "output.docx"
            override_path = temp_path / "override.json"

            doc = Document()
            for text in paragraphs:
                doc.add_paragraph(text)
            if add_table:
                table = doc.add_table(rows=1, cols=1)
                table.cell(0, 0).text = "表格文字"
            doc.save(str(input_path))

            if override is not None:
                override_path.write_text(
                    json.dumps(override, ensure_ascii=False),
                    encoding="utf-8",
                )

            rules = format_docx.get_final_format_rules(
                "default",
                str(override_path) if override is not None else None,
                print_warnings=False,
            )
            report = format_docx.init_report(input_path, output_path, rules)
            format_docx.format_document(input_path, output_path, rules, report, overwrite=True)
            return report

    def test_detect_chinese_heading_1(self):
        result = format_docx.detect_paragraph_type("一、研究背景", 0, make_context())
        self.assertEqual(result, "heading_1")

    def test_detect_chinese_heading_2(self):
        result = format_docx.detect_paragraph_type("（一）研究意义", 0, make_context())
        self.assertEqual(result, "heading_2")

    def test_detect_number_heading_2(self):
        result = format_docx.detect_paragraph_type("1.1 研究现状", 0, make_context())
        self.assertEqual(result, "heading_2")

    def test_detect_heading_from_word_heading_style_without_number_text(self):
        doc = Document()
        paragraph = doc.add_paragraph("研究背景", style="Heading 2")
        result = format_docx.detect_paragraph_type(
            paragraph.text.strip(),
            0,
            make_context(),
            paragraph,
        )
        self.assertEqual(result, "heading_2")

    def test_detect_table_caption(self):
        result = format_docx.detect_paragraph_type("表1 论文格式处理流程示例", 0, make_context())
        self.assertEqual(result, "table_caption")

    def test_detect_table_caption_with_space(self):
        result = format_docx.detect_paragraph_type("表 1 论文格式处理流程示例", 0, make_context())
        self.assertEqual(result, "table_caption")

    def test_detect_figure_caption(self):
        result = format_docx.detect_paragraph_type("图1 系统架构图", 0, make_context())
        self.assertEqual(result, "figure_caption")

    def test_detect_figure_caption_with_dash_number(self):
        result = format_docx.detect_paragraph_type("图 1-1 用户操作流程图", 0, make_context())
        self.assertEqual(result, "figure_caption")

    def test_reference_section_caption_like_text_stays_reference_item(self):
        context = make_context()
        context["in_reference_section"] = True
        result = format_docx.detect_paragraph_type("图1 系统架构图", 0, context)
        self.assertEqual(result, "reference_item")

    def test_reference_mode_exits_on_appendix(self):
        context = make_context()
        self.assertEqual(
            format_docx.detect_paragraph_type("参考文献", 0, context),
            "reference_title",
        )
        self.assertEqual(
            format_docx.detect_paragraph_type("[1] 张三. 测试文献. 2024.", 1, context),
            "reference_item",
        )
        self.assertEqual(
            format_docx.detect_paragraph_type(
                "1. 注意：这一条在参考文献区域内，虽然以数字编号开头，但不应该被识别成正文标题。",
                2,
                context,
            ),
            "reference_item",
        )

        appendix_type = format_docx.detect_paragraph_type("附录", 3, context)
        self.assertNotEqual(appendix_type, "reference_item")
        self.assertFalse(context["in_reference_section"])
        self.assertNotEqual(
            format_docx.detect_paragraph_type("一、访谈提纲", 4, context),
            "reference_item",
        )
        self.assertEqual(
            format_docx.detect_paragraph_type(
                "1. 你平时会在哪些学习场景中使用AI工具？",
                5,
                context,
            ),
            "body",
        )

    def test_reference_mode_exits_on_english_appendix_titles(self):
        for text in ["Appendix", "Appendix A", "Acknowledgement", "Acknowledgments"]:
            context = make_context()
            context["in_reference_section"] = True
            self.assertNotEqual(
                format_docx.detect_paragraph_type(text, 0, context),
                "reference_item",
            )
            self.assertFalse(context["in_reference_section"])

    def test_module_status_for_simple_paper(self):
        report = self.format_paragraphs_for_report(
            [
                "测试论文标题",
                "摘要：本文用于测试模块状态。",
                "关键词：模块状态；参考文献",
                "一、研究背景",
                "这是一段普通正文。",
                "参考文献",
                "[1] 张三. 测试文献. 2024.",
            ]
        )
        module_status = report["module_status"]

        self.assertEqual(module_status["paper_title"]["status"], "detected")
        self.assertEqual(module_status["abstract_cn"]["status"], "detected")
        self.assertEqual(module_status["keywords_cn"]["status"], "detected")
        self.assertEqual(module_status["body"]["status"], "detected")
        self.assertEqual(module_status["reference"]["status"], "detected")
        self.assertEqual(module_status["abstract_en"]["status"], "not_detected")
        self.assertEqual(module_status["keywords_en"]["status"], "not_detected")
        self.assertEqual(module_status["toc"]["status"], "not_detected")
        self.assertIn(
            "module_status",
            report["final_rules_debug"]["debug_summary"],
        )

    def test_module_status_detects_english_abstract_and_keywords(self):
        report = self.format_paragraphs_for_report(
            [
                "测试论文标题",
                "Abstract",
                "This paragraph is an English abstract.",
                "Keywords: AI; learning",
            ]
        )
        module_status = report["module_status"]

        self.assertEqual(module_status["abstract_en"]["status"], "detected")
        self.assertEqual(module_status["abstract_en"]["count"], 2)
        self.assertEqual(module_status["abstract_en"]["action"], "formatted")
        self.assertEqual(module_status["keywords_en"]["status"], "detected")
        self.assertEqual(module_status["keywords_en"]["count"], 1)
        self.assertEqual(module_status["keywords_en"]["action"], "formatted")

    def test_module_status_detects_and_protects_toc_entries(self):
        report = self.format_paragraphs_for_report(
            [
                "测试论文标题",
                "目录",
                "一、研究背景 ........ 1",
                "1.1 研究说明   2",
                "一、研究背景",
                "这是一段正文。",
            ]
        )
        module_status = report["module_status"]
        detected_types = [item["detected_type"] for item in report["paragraphs"]]

        self.assertEqual(module_status["toc"]["status"], "detected")
        self.assertEqual(module_status["toc"]["action"], "protected")
        self.assertEqual(detected_types[1], "body")
        self.assertEqual(detected_types[2], "body")
        self.assertEqual(detected_types[3], "body")
        self.assertEqual(detected_types[4], "heading_1")

    def test_module_status_for_realistic_paper_structure(self):
        report = self.format_paragraphs_for_report(
            [
                "财务透明度对企业融资成本的影响——以中小企业为例",
                "摘要：本文研究财务透明度与企业融资成本的关系。",
                "关键词：财务透明度；融资成本；中小企业",
                "Abstract",
                "This paper studies financial transparency and financing cost.",
                "Keywords",
                "financial transparency; financing cost",
                "目录",
                "一、绪论 2",
                "1.1 研究背景 3",
                "1.1.1 财务透明度的概念与发展 4",
                "一、绪论",
                "1.1 研究背景",
                "1.1.1 三级标题",
                "这是正文内容，用于验证正文模块仍然被检测。",
                "参考文献",
                "[1] 张三. 财务透明度研究. 2024.",
            ]
        )
        module_status = report["module_status"]

        self.assertEqual(module_status["abstract_cn"]["status"], "detected")
        self.assertEqual(module_status["keywords_cn"]["status"], "detected")
        self.assertEqual(module_status["abstract_en"]["status"], "detected")
        self.assertEqual(module_status["keywords_en"]["status"], "detected")
        self.assertEqual(module_status["toc"]["status"], "detected")
        self.assertEqual(module_status["toc"]["action"], "protected")
        self.assertEqual(module_status["heading"]["status"], "detected")
        self.assertGreater(module_status["heading"]["count"], 0)
        self.assertEqual(module_status["body"]["status"], "detected")
        self.assertEqual(module_status["reference"]["status"], "detected")

    def test_module_status_heading_uses_word_heading_styles(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            input_path = temp_path / "input.docx"
            output_path = temp_path / "output.docx"

            doc = Document()
            doc.add_paragraph("论文标题")
            doc.add_paragraph("研究背景", style="Heading 2")
            doc.add_paragraph("正文内容")
            doc.save(str(input_path))

            rules = format_docx.get_final_format_rules("default", None, print_warnings=False)
            report = format_docx.init_report(input_path, output_path, rules)
            format_docx.format_document(input_path, output_path, rules, report, overwrite=True)

        module_status = report["module_status"]
        self.assertEqual(module_status["heading"]["status"], "detected")
        self.assertGreater(module_status["heading"]["count"], 0)

    def test_module_status_records_unsupported_header_footer_and_page_number(self):
        report = self.format_paragraphs_for_report(
            ["测试论文标题", "这是一段正文。"],
            override={
                "name": "module_requirements",
                "module_requirements": {
                    "header_footer": True,
                    "page_number": True,
                },
                "styles": {"body": {"font": "宋体"}},
            },
        )
        module_status = report["module_status"]
        warning_messages = [item["message"] for item in report["warnings"]]

        self.assertEqual(
            module_status["header_footer"]["status"],
            "detected_but_not_supported",
        )
        self.assertEqual(module_status["header_footer"]["action"], "warning_only")
        self.assertEqual(
            module_status["page_number"]["status"],
            "detected_but_not_supported",
        )
        self.assertEqual(module_status["page_number"]["action"], "warning_only")
        self.assertTrue(any("页眉页脚" in message for message in warning_messages))
        self.assertTrue(any("页码" in message for message in warning_messages))

    def test_module_status_for_plain_paper_keeps_missing_modules_explicit(self):
        report = self.format_paragraphs_for_report(["测试论文标题", "这是一段正文。"])
        module_status = report["module_status"]

        for module_key in format_docx.MODULE_STATUS_KEYS:
            self.assertIn(module_key, module_status)
        self.assertEqual(module_status["header_footer"]["status"], "not_requested")
        self.assertEqual(module_status["page_number"]["status"], "not_requested")
        self.assertEqual(module_status["latin_digit_format"]["status"], "not_requested")
        self.assertEqual(module_status["latin_digit_format"]["action"], "skipped")
        self.assertEqual(module_status["abstract_en"]["status"], "not_detected")
        self.assertEqual(module_status["toc"]["status"], "not_detected")

    def test_module_status_records_latin_digit_format_processing(self):
        report = self.format_paragraphs_for_report(
            ["测试论文标题", "正文包含 AI 和 2025。"],
            override={
                "name": "latin_digit_teacher_rule",
                "latin_digit_format": {
                    "font": "Times New Roman",
                    "size_cn": "小四",
                    "scope": "global",
                },
            },
        )
        module_status = report["module_status"]["latin_digit_format"]

        self.assertEqual(module_status["status"], "detected")
        self.assertEqual(module_status["action"], "formatted")
        self.assertGreater(module_status["count"], 0)
        self.assertEqual(
            report["final_rules_debug"]["debug_summary"]["module_status"][
                "latin_digit_format"
            ]["action"],
            "formatted",
        )

    def test_module_status_uses_actual_execution_counts(self):
        report = self.format_paragraphs_for_report(
            [
                "测试论文标题",
                "Abstract",
                "This paragraph is an English abstract.",
                "Keywords: Clinical nursing; Recovery",
                "参考文献",
                "[1] Xiaofeng L. 护理研究[J]. Medicine, 2025, 104(37): e44505.",
            ],
            override={
                "name": "module_status_execution",
                "styles": {
                    "abstract_en_title": {"font": "Times New Roman", "size_cn": "小四"},
                    "abstract_en_content": {"font": "Times New Roman", "size_cn": "小四"},
                    "keywords_en": {"font": "Times New Roman", "size_cn": "小四"},
                    "keywords_en_label": {
                        "font": "Times New Roman",
                        "size_cn": "小四",
                        "bold": True,
                    },
                    "keywords_en_content": {
                        "font": "Times New Roman",
                        "size_cn": "小四",
                        "bold": False,
                    },
                    "reference_item": {
                        "font": "宋体",
                        "size_cn": "五号",
                        "hanging_indent_chars": 2,
                    },
                },
                "reference_latin_digit_format": {
                    "font": "Times New Roman",
                    "size_cn": "五号",
                    "scope": "reference",
                },
                "unsupported_modules": {
                    "header_footer": {
                        "status": "detected_but_not_supported",
                        "action": "warning_only",
                        "note": "检测到页眉页脚要求，但当前版本暂不处理。",
                    },
                    "page_number": {
                        "status": "detected_but_not_supported",
                        "action": "warning_only",
                        "note": "检测到页码要求，但当前版本暂不处理。",
                    },
                },
            },
        )
        module_status = report["module_status"]

        self.assertEqual(module_status["abstract_en"]["action"], "formatted")
        self.assertEqual(module_status["keywords_en"]["action"], "formatted")
        self.assertEqual(module_status["reference"]["action"], "formatted")
        self.assertEqual(module_status["reference_latin_digit_format"]["action"], "formatted")
        self.assertEqual(module_status["header_footer"]["action"], "warning_only")
        self.assertEqual(module_status["page_number"]["action"], "warning_only")
        debug_summary = report["final_rules_debug"]["debug_summary"]
        self.assertGreater(
            debug_summary["keyword_label_content"]["keywords_en_label_count"],
            0,
        )
        self.assertGreater(
            debug_summary["reference_execution"]["reference_latin_digit_format_count"],
            0,
        )
        self.assertGreater(
            debug_summary["reference_execution"]["reference_hanging_indent_count"],
            0,
        )

    def test_normalize_font_size_xiaosi(self):
        self.assertEqual(format_docx.normalize_font_size("小四"), 12)

    def test_normalize_font_size_wuhao(self):
        self.assertEqual(format_docx.normalize_font_size("五号"), 10.5)

    def test_normalize_alignment_center(self):
        self.assertEqual(format_docx.normalize_alignment("居中"), "center")

    def test_normalize_alignment_justify(self):
        self.assertEqual(format_docx.normalize_alignment("两端对齐"), "justify")

    def test_normalize_line_spacing_one_point_five(self):
        self.assertEqual(format_docx.normalize_line_spacing("1.5倍"), 1.5)

    def test_normalize_line_spacing_single(self):
        self.assertEqual(format_docx.normalize_line_spacing("单倍行距"), 1.0)

    def test_normalize_indent_two_characters(self):
        self.assertEqual(format_docx.normalize_indent("2字符"), 24)

    def test_first_line_indent_chars_and_points_are_normalized(self):
        normalized = format_docx.normalize_format_rules(
            {
                "styles": {
                    "body": {
                        "font": "宋体",
                        "size_cn": "小四",
                        "first_line_indent_chars": 2,
                        "first_line_indent_pt": 24,
                    },
                    "reference_item": {
                        "font": "宋体",
                        "size_cn": "五号",
                        "first_line_indent_chars": 2,
                        "first_line_indent_pt": 12,
                    },
                }
            },
            fill_defaults=False,
        )

        self.assertEqual(normalized["styles"]["body"]["first_line_indent_pt"], 24)
        body_summary = [
            item
            for item in normalized["_indent_normalization"]
            if item["style"] == "body"
        ][0]
        self.assertEqual(body_summary["status"], "equivalent")
        self.assertFalse(
            any(
                "styles.body.first_line_indent" in warning
                for warning in normalized["_template_warnings"]
            )
        )
        self.assertTrue(
            any(
                "styles.reference_item.first_line_indent_chars" in warning
                for warning in normalized["_template_warnings"]
            )
        )

    def test_normalize_page_margin_cm_string(self):
        self.assertEqual(format_docx.normalize_page_margin("3cm"), 3)

    def test_normalize_color_black(self):
        self.assertEqual(format_docx.normalize_color("黑色"), "000000")

    def test_validate_format_rules_accepts_color(self):
        template = {
            "version": "1.0",
            "name": "test",
            "page": {},
            "styles": {"body": {"color": "000000"}},
        }
        errors, _ = format_docx.validate_format_rules(template)
        self.assertEqual(errors, [])

    def test_validate_format_rules_accepts_underline(self):
        template = {
            "version": "1.0",
            "name": "test",
            "page": {},
            "styles": {"body": {"underline": False}},
        }
        errors, _ = format_docx.validate_format_rules(template)
        self.assertEqual(errors, [])

    def test_validate_format_rules_rejects_bad_underline(self):
        template = {
            "version": "1.0",
            "name": "bad",
            "page": {},
            "styles": {"body": {"underline": "no"}},
        }
        errors, _ = format_docx.validate_format_rules(template)
        self.assertTrue(any("styles.body.underline" in error for error in errors))

    def test_validate_format_rules_accepts_keep_with_next(self):
        template = {
            "version": "1.0",
            "name": "test",
            "page": {},
            "styles": {"body": {}, "table_caption": {"keep_with_next": True}},
        }
        errors, _ = format_docx.validate_format_rules(template)
        self.assertEqual(errors, [])

    def test_validate_format_rules_accepts_keep_together(self):
        template = {
            "version": "1.0",
            "name": "test",
            "page": {},
            "styles": {"body": {}, "figure_caption": {"keep_together": True}},
        }
        errors, _ = format_docx.validate_format_rules(template)
        self.assertEqual(errors, [])

    def test_validate_format_rules_accepts_v04_parse_fields(self):
        template = {
            "version": "1.0",
            "name": "test",
            "page": {},
            "styles": {
                "body": {"font": "宋体"},
                "abstract_cn_title": {"font": "黑体", "bold": True},
                "abstract_cn_content": {
                    "font": "宋体",
                    "first_line_indent_chars": 2,
                },
                "keywords_cn_label": {"font": "黑体", "bold": True},
                "keywords_cn_content": {"font": "宋体", "bold": False},
                "keywords_en_label": {"font": "Times New Roman", "bold": True},
                "keywords_en_content": {"font": "Times New Roman", "bold": False},
            },
            "latin_digit_format": {
                "font": "Times New Roman",
                "scope": "heading",
            },
            "reference_latin_digit_format": {
                "font": "Times New Roman",
                "size_cn": "五号",
                "scope": "reference",
            },
            "toc": {"action": "protect"},
            "unsupported_modules": {
                "header_footer": {
                    "status": "detected_but_not_supported",
                    "action": "warning_only",
                    "note": "检测到页眉页脚要求，但当前版本暂不处理。",
                },
                "page_number": {
                    "status": "detected_but_not_supported",
                    "action": "warning_only",
                    "note": "检测到页码要求，但当前版本暂不处理。",
                },
            },
        }
        normalized = format_docx.normalize_format_rules(template)
        errors, _ = format_docx.validate_format_rules(normalized)
        self.assertEqual(errors, [])
        self.assertEqual(
            normalized["reference_latin_digit_format"]["size_pt"],
            10.5,
        )
        self.assertEqual(
            normalized["styles"]["abstract_cn_content"]["first_line_indent_pt"],
            24,
        )

    def test_validate_format_rules_rejects_bad_keep_with_next(self):
        template = {
            "version": "1.0",
            "name": "bad",
            "page": {},
            "styles": {"body": {}, "table_caption": {"keep_with_next": "yes"}},
        }
        errors, _ = format_docx.validate_format_rules(template)
        self.assertTrue(
            any(
                "styles.table_caption.keep_with_next 应为布尔值 true 或 false"
                in error
                for error in errors
            )
        )

    def test_validate_format_rules_rejects_bad_color(self):
        template = {
            "version": "1.0",
            "name": "bad",
            "page": {},
            "styles": {"body": {"color": "black"}},
        }
        errors, _ = format_docx.validate_format_rules(template)
        self.assertTrue(any("styles.body.color" in error for error in errors))

    def test_apply_run_font_accepts_color(self):
        doc = Document()
        run = doc.add_paragraph().add_run("测试")
        format_docx.apply_run_font(
            run,
            {
                "font": "宋体",
                "size_pt": 12,
                "color": "000000",
                "bold": False,
                "italic": False,
            },
        )
        self.assertEqual(str(run.font.color.rgb), "000000")

    def test_format_normal_paragraphs_preserves_body_underline_without_explicit_rule(self):
        doc = Document()
        doc.add_paragraph("\u8bba\u6587\u6807\u9898")
        paragraph = doc.add_paragraph("\u8fd9\u662f\u6b63\u6587\u5185\u5bb9")
        paragraph.runs[0].font.underline = True
        template = {
            "styles": {
                "paper_title": {"font": "\u5b8b\u4f53", "size_pt": 12},
                "body": {"font": "\u5b8b\u4f53", "size_pt": 12},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        format_docx.format_normal_paragraphs(doc, template, report)

        self.assertTrue(paragraph.runs[0].font.underline)

    def test_format_normal_paragraphs_clears_body_underline_when_explicit_false(self):
        doc = Document()
        doc.add_paragraph("\u8bba\u6587\u6807\u9898")
        paragraph = doc.add_paragraph("\u8fd9\u662f\u6b63\u6587\u5185\u5bb9")
        paragraph.runs[0].font.underline = True
        template = {
            "styles": {
                "paper_title": {"font": "\u5b8b\u4f53", "size_pt": 12},
                "body": {"font": "\u5b8b\u4f53", "size_pt": 12, "underline": False},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        format_docx.format_normal_paragraphs(doc, template, report)

        self.assertFalse(paragraph.runs[0].font.underline)

    def test_format_normal_paragraphs_preserves_title_underline_without_explicit_rule(self):
        doc = Document()
        paragraph = doc.add_paragraph("\u8bba\u6587\u6807\u9898")
        paragraph.runs[0].font.underline = True
        template = {
            "styles": {
                "paper_title": {"font": "\u5b8b\u4f53", "size_pt": 12},
                "body": {"font": "\u5b8b\u4f53", "size_pt": 12},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        format_docx.format_normal_paragraphs(doc, template, report)

        self.assertTrue(paragraph.runs[0].font.underline)

    def test_format_tables_clears_table_text_underline(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        paragraph = table.cell(0, 0).paragraphs[0]
        paragraph.text = "\u8868\u683c\u6587\u5b57"
        paragraph.runs[0].font.underline = True
        template = {
            "styles": {
                "body": {"font": "\u5b8b\u4f53", "size_pt": 12},
                "table_text": {"font": "\u5b8b\u4f53", "size_pt": 10.5, "underline": False},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"tables": [], "warnings": []}

        format_docx.format_tables(doc, template, report)

        self.assertFalse(paragraph.runs[0].font.underline)

    def test_format_normal_paragraphs_clears_reference_item_underline(self):
        doc = Document()
        doc.add_paragraph("\u53c2\u8003\u6587\u732e")
        paragraph = doc.add_paragraph("[1] \u5f20\u4e09. \u6d4b\u8bd5\u6587\u732e. 2024.")
        paragraph.runs[0].font.underline = True
        template = {
            "styles": {
                "body": {"font": "\u5b8b\u4f53", "size_pt": 12},
                "reference_title": {"font": "\u9ed1\u4f53", "size_pt": 14},
                "reference_item": {"font": "\u5b8b\u4f53", "size_pt": 10.5, "underline": False},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        format_docx.format_normal_paragraphs(doc, template, report)

        self.assertEqual(report["paragraphs"][1]["detected_type"], "reference_item")
        self.assertFalse(paragraph.runs[0].font.underline)

    def test_explicit_underline_true_is_preserved(self):
        doc = Document()
        doc.add_paragraph("\u8bba\u6587\u6807\u9898")
        paragraph = doc.add_paragraph("\u8fd9\u662f\u6b63\u6587\u5185\u5bb9")
        paragraph.runs[0].font.underline = False
        template = {
            "styles": {
                "paper_title": {"font": "\u5b8b\u4f53", "size_pt": 12},
                "body": {
                    "font": "\u5b8b\u4f53",
                    "size_pt": 12,
                    "bold": True,
                    "italic": True,
                    "underline": True,
                },
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        format_docx.format_normal_paragraphs(doc, template, report)

        self.assertTrue(paragraph.runs[0].font.underline)
        self.assertTrue(paragraph.runs[0].bold)
        self.assertTrue(paragraph.runs[0].italic)

    def test_format_preserves_word_heading_style_and_numbering(self):
        doc = Document()
        doc.add_paragraph("论文标题")
        heading_1 = doc.add_paragraph("绪论", style="Heading 1")
        heading_2 = doc.add_paragraph("研究背景", style="Heading 2")
        heading_3 = doc.add_paragraph("理论意义", style="Heading 3")
        for paragraph in (heading_1, heading_2, heading_3):
            add_direct_numbering(paragraph)

        template = {
            "styles": {
                "paper_title": {"font": "宋体", "size_pt": 12},
                "body": {"font": "宋体", "size_pt": 12},
                "heading_1": {"font": "黑体", "size_pt": 15},
                "heading_2": {"font": "黑体", "size_pt": 14},
                "heading_3": {"font": "黑体", "size_pt": 12},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        format_docx.format_normal_paragraphs(doc, template, report)

        self.assertEqual(heading_1.style.name, "Heading 1")
        self.assertEqual(heading_2.style.name, "Heading 2")
        self.assertEqual(heading_3.style.name, "Heading 3")
        self.assertTrue(has_direct_numbering(heading_1))
        self.assertTrue(has_direct_numbering(heading_2))
        self.assertTrue(has_direct_numbering(heading_3))

    def test_numbered_heading_level_3_overrides_heading_2_style(self):
        doc = Document()
        doc.add_paragraph("论文标题")
        paragraph = doc.add_paragraph("1.1.1 财务透明度的概念与发展", style="Heading 2")
        original_text = paragraph.text
        template = {
            "styles": {
                "paper_title": {"font": "宋体", "size_pt": 12},
                "body": {"font": "宋体", "size_pt": 12},
                "heading_1": {
                    "font": "黑体",
                    "size_pt": 15,
                    "bold": True,
                    "space_before_pt": 12,
                    "space_before_lines": 0.5,
                    "space_after_lines": 0.5,
                },
                "heading_2": {"font": "黑体", "size_pt": 14, "bold": True},
                "heading_3": {"font": "黑体", "size_pt": 12, "bold": True},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        context = format_docx.format_normal_paragraphs(doc, template, report)

        self.assertEqual(paragraph.text, original_text)
        self.assertEqual(report["paragraphs"][1]["detected_type"], "heading_3")
        self.assertEqual(paragraph.style.name, "Heading 3")
        self.assertAlmostEqual(paragraph.runs[0].font.size.pt, 12)
        self.assertTrue(paragraph.runs[0].bold)
        self.assertEqual(context["heading_style_conflicts"][0]["final_level"], "heading_3")

    def test_numbered_heading_level_2_overrides_wrong_heading_1_style_and_keeps_numpr(self):
        doc = Document()
        doc.add_paragraph("论文标题")
        paragraph = doc.add_paragraph("3.1 中小企业融资困难", style="Heading 1")
        add_direct_numbering(paragraph)
        original_text = paragraph.text
        template = {
            "styles": {
                "paper_title": {"font": "宋体", "size_pt": 12},
                "body": {"font": "宋体", "size_pt": 12},
                "heading_1": {
                    "font": "黑体",
                    "size_pt": 15,
                    "bold": True,
                    "space_before_pt": 12,
                    "space_before_lines": 0.5,
                    "space_after_lines": 0.5,
                },
                "heading_2": {"font": "黑体", "size_pt": 14, "bold": True},
                "heading_3": {"font": "黑体", "size_pt": 12, "bold": True},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        context = format_docx.format_normal_paragraphs(doc, template, report)

        self.assertEqual(paragraph.text, original_text)
        self.assertEqual(report["paragraphs"][1]["detected_type"], "heading_2")
        self.assertEqual(paragraph.style.name, "Heading 2")
        self.assertAlmostEqual(paragraph.runs[0].font.size.pt, 14)
        self.assertTrue(has_direct_numbering(paragraph))
        self.assertEqual(context["heading_style_conflicts"][0]["original_style_level"], "heading_1")
        self.assertEqual(context["heading_style_conflicts"][0]["final_level"], "heading_2")
        self.assertTrue(any("标题层级冲突" in item["message"] for item in report["warnings"]))

    def test_chinese_numbered_heading_from_normal_syncs_to_heading_1(self):
        doc = Document()
        doc.add_paragraph("论文标题")
        paragraph = doc.add_paragraph("四、提高财务透明度降低中小企业融资成本的对策")
        original_text = paragraph.text
        template = {
            "styles": {
                "paper_title": {"font": "宋体", "size_pt": 12},
                "body": {"font": "宋体", "size_pt": 12},
                "heading_1": {
                    "font": "黑体",
                    "size_pt": 15,
                    "bold": True,
                    "space_before_pt": 12,
                    "space_before_lines": 0.5,
                    "space_after_lines": 0.5,
                },
                "heading_2": {"font": "黑体", "size_pt": 14, "bold": True},
                "heading_3": {"font": "黑体", "size_pt": 12, "bold": True},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        format_docx.format_normal_paragraphs(doc, template, report)

        self.assertEqual(paragraph.text, original_text)
        self.assertEqual(report["paragraphs"][1]["detected_type"], "heading_1")
        self.assertEqual(paragraph.style.name, "Heading 1")
        self.assertAlmostEqual(paragraph.runs[0].font.size.pt, 15)
        self.assertTrue(paragraph.runs[0].bold)
        self.assertAlmostEqual(paragraph.paragraph_format.space_before.pt, 7.5)
        self.assertAlmostEqual(paragraph.paragraph_format.space_after.pt, 7.5)

    def test_arabic_level_1_headings_are_detected_in_numeric_heading_system(self):
        doc = Document()
        doc.add_paragraph("论文标题")
        h1_plain = doc.add_paragraph("1 绪论")
        h2 = doc.add_paragraph("1.1 研究背景")
        h3 = doc.add_paragraph("1.1.1 研究目的")
        body_list = doc.add_paragraph("1. 加强财务信息披露")
        h1_dot = doc.add_paragraph("2. 相关理论")
        h2_second = doc.add_paragraph("2.1 理论基础")
        h1_comma = doc.add_paragraph("3、研究方法")
        h2_third = doc.add_paragraph("3.1 样本来源")
        template = {
            "styles": {
                "paper_title": {"font": "宋体", "size_pt": 12},
                "body": {"font": "宋体", "size_pt": 12},
                "heading_1": {"font": "黑体", "size_pt": 15, "bold": True},
                "heading_2": {"font": "黑体", "size_pt": 14, "bold": True},
                "heading_3": {"font": "黑体", "size_pt": 12, "bold": True},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        context = format_docx.format_normal_paragraphs(doc, template, report)
        detected_types = [item["detected_type"] for item in report["paragraphs"]]

        self.assertTrue(context["heading_numbering_system"]["arabic_heading_system"])
        self.assertEqual(detected_types[1], "heading_1")
        self.assertEqual(detected_types[2], "heading_2")
        self.assertEqual(detected_types[3], "heading_3")
        self.assertEqual(detected_types[4], "body")
        self.assertEqual(detected_types[5], "heading_1")
        self.assertEqual(detected_types[6], "heading_2")
        self.assertEqual(detected_types[7], "heading_1")
        self.assertEqual(detected_types[8], "heading_2")
        self.assertEqual(body_list.style.name, "Normal")
        for paragraph in (h1_plain, h1_dot, h1_comma):
            self.assertEqual(paragraph.style.name, "Heading 1")
            self.assertAlmostEqual(paragraph.runs[0].font.size.pt, 15)
        self.assertEqual(h2.style.name, "Heading 2")
        self.assertEqual(h2_second.style.name, "Heading 2")
        self.assertEqual(h2_third.style.name, "Heading 2")
        self.assertEqual(h3.style.name, "Heading 3")

    def test_chapter_heading_forms_are_level_1_without_numeric_subheadings(self):
        doc = Document()
        doc.add_paragraph("论文标题")
        chapter_digit = doc.add_paragraph("第1章 绪论")
        chapter_chinese = doc.add_paragraph("第一章 研究背景")
        body_sentence = doc.add_paragraph("第一章为绪论。本章主要介绍研究背景和研究意义。")
        template = {
            "styles": {
                "paper_title": {"font": "宋体", "size_pt": 12},
                "body": {"font": "宋体", "size_pt": 12},
                "heading_1": {"font": "黑体", "size_pt": 15, "bold": True},
                "heading_2": {"font": "黑体", "size_pt": 14, "bold": True},
                "heading_3": {"font": "黑体", "size_pt": 12, "bold": True},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        format_docx.format_normal_paragraphs(doc, template, report)
        detected_types = [item["detected_type"] for item in report["paragraphs"]]

        self.assertEqual(detected_types[1], "heading_1")
        self.assertEqual(detected_types[2], "heading_1")
        self.assertEqual(detected_types[3], "body")
        self.assertEqual(chapter_digit.style.name, "Heading 1")
        self.assertEqual(chapter_chinese.style.name, "Heading 1")
        self.assertEqual(body_sentence.text, "第一章为绪论。本章主要介绍研究背景和研究意义。")

    def test_single_arabic_numbered_body_item_is_not_forced_to_heading_1(self):
        doc = Document()
        doc.add_paragraph("论文标题")
        body = doc.add_paragraph("这是一段正文说明。")
        item = doc.add_paragraph("1. 加强财务信息披露")
        template = {
            "styles": {
                "paper_title": {"font": "宋体", "size_pt": 12},
                "body": {"font": "宋体", "size_pt": 12},
                "heading_1": {"font": "黑体", "size_pt": 15, "bold": True},
                "heading_2": {"font": "黑体", "size_pt": 14, "bold": True},
                "heading_3": {"font": "黑体", "size_pt": 12, "bold": True},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        context = format_docx.format_normal_paragraphs(doc, template, report)
        detected_types = [item["detected_type"] for item in report["paragraphs"]]

        self.assertEqual(body.text, "这是一段正文说明。")
        self.assertEqual(item.text, "1. 加强财务信息披露")
        self.assertEqual(detected_types[2], "body")
        self.assertEqual(item.style.name, "Normal")
        self.assertFalse(context["heading_numbering_system"]["arabic_heading_system"])
        self.assertTrue(context["arabic_heading_level1_skipped"])

    def test_toc_entries_are_not_reclassified_as_numbered_headings(self):
        report = self.format_paragraphs_for_report(
            [
                "论文标题",
                "目录",
                "1.1.1 财务透明度的概念与发展   3",
                "1.1.1 财务透明度的概念与发展",
            ]
        )
        detected_types = [item["detected_type"] for item in report["paragraphs"]]

        self.assertEqual(detected_types[1], "body")
        self.assertEqual(detected_types[2], "body")
        self.assertEqual(detected_types[3], "heading_3")
        self.assertEqual(report["module_status"]["toc"]["action"], "protected")

    def test_latin_digit_format_splits_mixed_body_run_without_changing_text(self):
        doc = Document()
        doc.add_paragraph("论文标题")
        paragraph = doc.add_paragraph()
        paragraph.add_run("人工智能 AI 2.0 技术推动 education reform in 2025。")
        original_text = paragraph.text
        template = {
            "styles": {
                "paper_title": {"font": "黑体", "size_pt": 15},
                "body": {"font": "宋体", "size_pt": 12},
            },
            "latin_digit_format": {
                "font": "Times New Roman",
                "size_pt": 12,
                "scope": "global",
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        format_docx.format_normal_paragraphs(doc, template, report)

        self.assertEqual(paragraph.text, original_text)
        self.assertGreater(len(paragraph.runs), 1)
        latin_runs = [
            run for run in paragraph.runs if format_docx.LATIN_DIGIT_TEXT_PATTERN.search(run.text)
        ]
        self.assertEqual(
            "".join(run.text for run in latin_runs),
            "AI2.0educationreformin2025",
        )
        for run in latin_runs:
            self.assertEqual(get_run_font_mapping(run, "ascii"), "Times New Roman")
            self.assertEqual(get_run_font_mapping(run, "hAnsi"), "Times New Roman")
            self.assertAlmostEqual(run.font.size.pt, 12)
            self.assertFalse(run.font.underline)

        chinese_runs = [run for run in paragraph.runs if "人工智能" in run.text]
        self.assertEqual(len(chinese_runs), 1)
        self.assertEqual(get_run_font_mapping(chinese_runs[0], "eastAsia"), "宋体")

    def test_latin_digit_format_preserves_heading_style_and_numbering(self):
        doc = Document()
        doc.add_paragraph("论文标题")
        heading = doc.add_paragraph("AI技术在2025年的发展", style="Heading 2")
        add_direct_numbering(heading)
        original_text = heading.text
        template = {
            "styles": {
                "paper_title": {"font": "黑体", "size_pt": 15},
                "body": {"font": "宋体", "size_pt": 12},
                "heading_2": {"font": "黑体", "size_pt": 14},
            },
            "latin_digit_format": {"font": "Times New Roman", "scope": "global"},
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        format_docx.format_normal_paragraphs(doc, template, report)

        self.assertEqual(heading.text, original_text)
        self.assertEqual(heading.style.name, "Heading 2")
        self.assertTrue(has_direct_numbering(heading))
        latin_runs = [
            run for run in heading.runs if format_docx.LATIN_DIGIT_TEXT_PATTERN.search(run.text)
        ]
        self.assertEqual("".join(run.text for run in latin_runs), "AI2025")
        for run in latin_runs:
            self.assertEqual(get_run_font_mapping(run, "ascii"), "Times New Roman")
            self.assertEqual(get_run_font_mapping(run, "hAnsi"), "Times New Roman")

    def test_latin_digit_format_protects_toc_text(self):
        doc = Document()
        doc.add_paragraph("论文标题")
        doc.add_paragraph("目录")
        toc_entry = doc.add_paragraph("1.1 AI研究背景 2")
        original_text = toc_entry.text
        template = {
            "styles": {
                "paper_title": {"font": "黑体", "size_pt": 15},
                "body": {"font": "宋体", "size_pt": 12},
            },
            "latin_digit_format": {"font": "Times New Roman", "scope": "global"},
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        context = format_docx.format_normal_paragraphs(doc, template, report)

        self.assertEqual(toc_entry.text, original_text)
        self.assertEqual(len(toc_entry.runs), 1)
        self.assertNotEqual(get_run_font_mapping(toc_entry.runs[0], "ascii"), "Times New Roman")
        self.assertGreater(context["module_counts"].get("toc", 0), 0)
        self.assertEqual(context["module_counts"].get("latin_digit_format", 0), 0)

    def test_keywords_cn_label_content_run_formatting(self):
        doc = Document()
        doc.add_paragraph("论文标题")
        paragraph = doc.add_paragraph("关键词：临床护理干预；术后恢复；心理护理")
        original_text = paragraph.text
        template = {
            "styles": {
                "paper_title": {"font": "黑体", "size_pt": 15},
                "body": {"font": "宋体", "size_pt": 12},
                "keywords": {"font": "宋体", "size_pt": 12, "bold": False},
                "keywords_cn_label": {"font": "黑体", "size_pt": 12, "bold": True},
                "keywords_cn_content": {"font": "宋体", "size_pt": 12, "bold": False},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        context = format_docx.format_normal_paragraphs(doc, template, report)

        self.assertEqual(paragraph.text, original_text)
        self.assertEqual(paragraph.runs[0].text, "关键词：")
        self.assertEqual(get_run_font_mapping(paragraph.runs[0], "eastAsia"), "黑体")
        self.assertTrue(paragraph.runs[0].bold)
        content_text = "".join(run.text for run in paragraph.runs[1:])
        self.assertEqual(content_text, "临床护理干预；术后恢复；心理护理")
        for run in paragraph.runs[1:]:
            self.assertEqual(get_run_font_mapping(run, "eastAsia"), "宋体")
            self.assertFalse(run.bold)
        self.assertEqual(context["module_counts"].get("keywords_cn_label_count"), 1)
        self.assertEqual(context["module_counts"].get("keywords_cn_content_count"), 1)

    def test_keywords_en_label_content_run_formatting(self):
        doc = Document()
        doc.add_paragraph("论文标题")
        paragraph = doc.add_paragraph(
            "Keywords: Clinical nursing intervention; Postoperative recovery"
        )
        original_text = paragraph.text
        template = {
            "styles": {
                "paper_title": {"font": "黑体", "size_pt": 15},
                "body": {"font": "宋体", "size_pt": 12},
                "keywords": {"font": "宋体", "size_pt": 12, "bold": False},
                "keywords_en": {"font": "Times New Roman", "size_pt": 12},
                "keywords_en_label": {
                    "font": "Times New Roman",
                    "size_pt": 12,
                    "bold": True,
                },
                "keywords_en_content": {
                    "font": "Times New Roman",
                    "size_pt": 12,
                    "bold": False,
                },
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        context = format_docx.format_normal_paragraphs(doc, template, report)

        self.assertEqual(paragraph.text, original_text)
        self.assertEqual(paragraph.runs[0].text, "Keywords:")
        self.assertEqual(get_run_font_mapping(paragraph.runs[0], "ascii"), "Times New Roman")
        self.assertTrue(paragraph.runs[0].bold)
        content_text = "".join(run.text for run in paragraph.runs[1:])
        self.assertEqual(content_text, " Clinical nursing intervention; Postoperative recovery")
        for run in paragraph.runs[1:]:
            self.assertEqual(get_run_font_mapping(run, "ascii"), "Times New Roman")
            self.assertFalse(run.bold)
        self.assertEqual(context["module_counts"].get("keywords_en_label_count"), 1)
        self.assertEqual(context["module_counts"].get("keywords_en_content_count"), 1)

    def test_reference_latin_digit_format_applies_to_reference_items(self):
        doc = Document()
        doc.add_paragraph("论文标题")
        doc.add_paragraph("参考文献")
        paragraph = doc.add_paragraph(
            "[9] Xiaofeng L, Cai T, Yi Z, et al. 护理研究[J]. Medicine, 2025, 104(37): e44505."
        )
        original_text = paragraph.text
        template = {
            "styles": {
                "paper_title": {"font": "黑体", "size_pt": 15},
                "body": {"font": "宋体", "size_pt": 12},
                "reference_title": {"font": "黑体", "size_pt": 14},
                "reference_item": {"font": "宋体", "size_pt": 10.5},
            },
            "reference_latin_digit_format": {
                "font": "Times New Roman",
                "size_pt": 10.5,
                "scope": "reference",
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        context = format_docx.format_normal_paragraphs(doc, template, report)

        self.assertEqual(paragraph.text, original_text)
        latin_runs = [
            run
            for run in paragraph.runs
            if format_docx.REFERENCE_LATIN_DIGIT_TEXT_PATTERN.search(run.text)
        ]
        self.assertGreater(len(latin_runs), 0)
        for run in latin_runs:
            self.assertEqual(get_run_font_mapping(run, "ascii"), "Times New Roman")
            self.assertEqual(get_run_font_mapping(run, "hAnsi"), "Times New Roman")
            self.assertAlmostEqual(run.font.size.pt, 10.5)
        chinese_runs = [run for run in paragraph.runs if "护理研究" in run.text]
        self.assertEqual(len(chinese_runs), 1)
        self.assertEqual(get_run_font_mapping(chinese_runs[0], "eastAsia"), "宋体")
        self.assertGreater(context["module_counts"].get("reference_latin_digit_format", 0), 0)

    def test_reference_hanging_indent_applies_only_to_items(self):
        doc = Document()
        doc.add_paragraph("论文标题")
        title = doc.add_paragraph("参考文献")
        first_item = doc.add_paragraph("[1] 张三. 测试文献. 2024.")
        second_item = doc.add_paragraph("[2] 李四. 另一篇文献. 2025.")
        template = {
            "styles": {
                "paper_title": {"font": "黑体", "size_pt": 15},
                "body": {"font": "宋体", "size_pt": 12},
                "reference_title": {"font": "黑体", "size_pt": 14},
                "reference_item": {
                    "font": "宋体",
                    "size_pt": 10.5,
                    "hanging_indent_chars": 2,
                },
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        context = format_docx.format_normal_paragraphs(doc, template, report)

        self.assertEqual(first_item.text, "[1] 张三. 测试文献. 2024.")
        self.assertEqual(second_item.text, "[2] 李四. 另一篇文献. 2025.")
        self.assertNotAlmostEqual(title.paragraph_format.first_line_indent.pt, -21)
        for item in (first_item, second_item):
            self.assertAlmostEqual(item.paragraph_format.left_indent.pt, 21)
            self.assertAlmostEqual(item.paragraph_format.first_line_indent.pt, -21)
        self.assertEqual(context["module_counts"].get("reference_hanging_indent"), 2)

    def test_format_document_preserves_paragraph_text_order(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            input_path = temp_path / "input.docx"
            output_path = temp_path / "output.docx"

            original_texts = [
                "论文标题",
                "摘要",
                "这是一段中文摘要。",
                "关键词：格式；测试",
                "绪论",
                "研究背景",
                "正文内容不应被修改。",
                "参考文献",
                "[1] 张三. 测试文献. 2024.",
            ]
            doc = Document()
            doc.add_paragraph(original_texts[0])
            doc.add_paragraph(original_texts[1], style="Heading 1")
            doc.add_paragraph(original_texts[2])
            doc.add_paragraph(original_texts[3])
            doc.add_paragraph(original_texts[4], style="Heading 1")
            doc.add_paragraph(original_texts[5], style="Heading 2")
            doc.add_paragraph(original_texts[6])
            doc.add_paragraph(original_texts[7])
            doc.add_paragraph(original_texts[8])
            doc.save(str(input_path))

            rules = format_docx.get_final_format_rules("default", None, print_warnings=False)
            report = format_docx.init_report(input_path, output_path, rules)
            format_docx.format_document(input_path, output_path, rules, report, overwrite=True)
            output_doc = Document(str(output_path))

        self.assertEqual([p.text for p in output_doc.paragraphs], original_texts)

    def test_english_abstract_rules_do_not_override_chinese_abstract(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            input_path = temp_path / "input.docx"
            output_path = temp_path / "output.docx"
            override_path = temp_path / "override.json"

            doc = Document()
            doc.add_paragraph("测试论文标题")
            doc.add_paragraph("摘要")
            abstract_paragraph = doc.add_paragraph("这是中文摘要正文。")
            doc.add_paragraph("关键词：中文；摘要")
            doc.save(str(input_path))

            override_path.write_text(
                json.dumps(
                    {
                        "name": "english_only",
                        "styles": {
                            "abstract_en_content": {
                                "font": "Times New Roman",
                                "bold": True,
                            },
                            "keywords_en": {
                                "font": "Times New Roman",
                                "bold": True,
                            },
                        },
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            rules = format_docx.get_final_format_rules(
                "default",
                str(override_path),
                print_warnings=False,
            )
            report = format_docx.init_report(input_path, output_path, rules)
            format_docx.format_document(input_path, output_path, rules, report, overwrite=True)

            output_doc = Document(str(output_path))

        output_abstract = output_doc.paragraphs[2]
        output_keywords = output_doc.paragraphs[3]
        self.assertEqual(output_abstract.text, abstract_paragraph.text)
        self.assertNotEqual(output_abstract.runs[0].font.name, "Times New Roman")
        self.assertFalse(output_abstract.runs[0].bold)
        self.assertNotEqual(output_keywords.runs[0].font.name, "Times New Roman")
        self.assertFalse(output_keywords.runs[0].bold)

    def test_apply_paragraph_style_sets_pagination_controls(self):
        doc = Document()
        paragraph = doc.add_paragraph("表1 论文格式处理流程示例")
        format_docx.apply_paragraph_style(
            paragraph,
            {
                "font": "宋体",
                "size_pt": 10.5,
                "keep_with_next": True,
                "keep_together": True,
            },
        )
        self.assertTrue(paragraph.paragraph_format.keep_with_next)
        self.assertTrue(paragraph.paragraph_format.keep_together)

    def test_format_normal_paragraphs_applies_table_caption_pagination_default(self):
        doc = Document()
        paragraph = doc.add_paragraph("表1 论文格式处理流程示例")
        template = {
            "styles": {
                "body": {"font": "宋体", "size_pt": 12},
                "table_caption": {"font": "宋体", "size_pt": 10.5},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        format_docx.format_normal_paragraphs(doc, template, report)

        self.assertTrue(paragraph.paragraph_format.keep_with_next)
        self.assertTrue(paragraph.paragraph_format.keep_together)
        self.assertEqual(report["paragraphs"][0]["detected_type"], "table_caption")
        self.assertEqual(
            report["paragraphs"][0]["pagination"],
            {"keep_with_next": True, "keep_together": True},
        )

    def test_format_normal_paragraphs_applies_figure_caption_pagination_default(self):
        doc = Document()
        paragraph = doc.add_paragraph("图1 系统架构图")
        template = {
            "styles": {
                "body": {"font": "宋体", "size_pt": 12},
                "figure_caption": {"font": "宋体", "size_pt": 10.5},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        format_docx.format_normal_paragraphs(doc, template, report)

        self.assertTrue(paragraph.paragraph_format.keep_with_next)
        self.assertTrue(paragraph.paragraph_format.keep_together)
        self.assertEqual(report["paragraphs"][0]["detected_type"], "figure_caption")
        self.assertEqual(
            report["paragraphs"][0]["pagination"],
            {"keep_with_next": True, "keep_together": True},
        )

    def test_format_normal_paragraphs_clears_caption_direct_numbering(self):
        doc = Document()
        paragraph = doc.add_paragraph("\u88681 \u8bba\u6587\u683c\u5f0f\u5904\u7406\u6d41\u7a0b\u793a\u4f8b")
        add_direct_numbering(paragraph)
        self.assertIsNotNone(paragraph._p.pPr.find(format_docx.qn("w:numPr")))
        template = {
            "styles": {
                "body": {"font": "\u5b8b\u4f53", "size_pt": 12},
                "table_caption": {"font": "\u5b8b\u4f53", "size_pt": 10.5},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        format_docx.format_normal_paragraphs(doc, template, report)

        self.assertIsNone(paragraph._p.pPr.find(format_docx.qn("w:numPr")))

    def test_format_normal_paragraphs_clears_figure_caption_direct_numbering(self):
        doc = Document()
        paragraph = doc.add_paragraph("\u56fe1 \u7cfb\u7edf\u603b\u4f53\u6d41\u7a0b\u56fe")
        add_direct_numbering(paragraph)
        self.assertIsNotNone(paragraph._p.pPr.find(format_docx.qn("w:numPr")))
        template = {
            "styles": {
                "body": {"font": "\u5b8b\u4f53", "size_pt": 12},
                "figure_caption": {"font": "\u5b8b\u4f53", "size_pt": 10.5},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        format_docx.format_normal_paragraphs(doc, template, report)

        self.assertIsNone(paragraph._p.pPr.find(format_docx.qn("w:numPr")))
        self.assertTrue(paragraph.paragraph_format.keep_with_next)
        self.assertTrue(paragraph.paragraph_format.keep_together)

    def test_format_normal_paragraphs_clears_caption_list_style(self):
        doc = Document()
        paragraph = doc.add_paragraph("\u56fe1 \u7cfb\u7edf\u67b6\u6784\u56fe")
        paragraph.style = "List Bullet"
        self.assertEqual(paragraph.style.name, "List Bullet")
        template = {
            "styles": {
                "body": {"font": "\u5b8b\u4f53", "size_pt": 12},
                "figure_caption": {"font": "\u5b8b\u4f53", "size_pt": 10.5},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        format_docx.format_normal_paragraphs(doc, template, report)

        self.assertEqual(paragraph.style.name, "Normal")

    def test_format_normal_paragraphs_preserves_body_list_paragraph_style(self):
        doc = Document()
        doc.add_paragraph("Paper title")
        paragraph = doc.add_paragraph("\u8fd9\u662f\u4e00\u6bb5\u6b63\u6587")
        paragraph.style = "List Paragraph"
        self.assertEqual(paragraph.style.name, "List Paragraph")
        template = {
            "styles": {
                "body": {"font": "\u5b8b\u4f53", "size_pt": 12},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "warnings": []}

        format_docx.format_normal_paragraphs(doc, template, report)

        self.assertEqual(paragraph.text, "\u8fd9\u662f\u4e00\u6bb5\u6b63\u6587")
        self.assertEqual(paragraph.style.name, "List Paragraph")

    def test_clear_paragraph_numbering_keeps_manual_number_text(self):
        doc = Document()
        paragraph = doc.add_paragraph("1.1 \u6280\u672f\u8def\u7ebf")
        add_direct_numbering(paragraph)

        format_docx.clear_paragraph_numbering(paragraph)

        self.assertEqual(paragraph.text, "1.1 \u6280\u672f\u8def\u7ebf")
        self.assertIsNone(paragraph._p.pPr.find(format_docx.qn("w:numPr")))

    def test_format_tables_keeps_low_risk_caption_and_row_pagination(self):
        doc = Document()
        caption = doc.add_paragraph("\u88681 \u8bba\u6587\u683c\u5f0f\u5904\u7406\u6d41\u7a0b\u793a\u4f8b")
        table = doc.add_table(rows=3, cols=2)
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                cell.text = f"r{row_index}c{cell_index}"
        template = {
            "styles": {
                "body": {"font": "\u5b8b\u4f53", "size_pt": 12},
                "table_caption": {"font": "\u5b8b\u4f53", "size_pt": 10.5},
                "table_text": {
                    "font": "\u5b8b\u4f53",
                    "size_pt": 10.5,
                    "keep_with_next": True,
                    "keep_together": True,
                },
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "tables": [], "warnings": []}

        format_docx.format_normal_paragraphs(doc, template, report)
        format_docx.format_tables(doc, template, report)

        self.assertTrue(caption.paragraph_format.keep_with_next)
        self.assertTrue(caption.paragraph_format.keep_together)
        for row in table.rows:
            self.assertIsNotNone(row._tr.trPr.find(format_docx.qn("w:cantSplit")))
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    self.assertIsNot(paragraph.paragraph_format.keep_with_next, True)
                    self.assertIsNot(paragraph.paragraph_format.keep_together, True)

    def test_format_tables_preserves_table_cell_numbering_and_list_style(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        left = table.cell(0, 0).paragraphs[0]
        right = table.cell(0, 1).paragraphs[0]
        left.text = "\u9636\u6bb5"
        right.text = "\u8f93\u51fa"
        add_direct_numbering(left)
        right.style = "List Bullet"
        self.assertIsNotNone(left._p.pPr.find(format_docx.qn("w:numPr")))
        self.assertEqual(right.style.name, "List Bullet")
        template = {
            "styles": {
                "body": {"font": "\u5b8b\u4f53", "size_pt": 12},
                "table_text": {"font": "\u5b8b\u4f53", "size_pt": 10.5},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"tables": [], "warnings": []}

        format_docx.format_tables(doc, template, report)

        self.assertEqual(left.text, "\u9636\u6bb5")
        self.assertEqual(right.text, "\u8f93\u51fa")
        self.assertIsNotNone(left._p.pPr.find(format_docx.qn("w:numPr")))
        self.assertEqual(right.style.name, "List Bullet")

    def test_clears_empty_numbered_paragraph_between_caption_and_table(self):
        doc = Document()
        doc.add_paragraph("\u88681 \u8bba\u6587\u683c\u5f0f\u5904\u7406\u6d41\u7a0b\u793a\u4f8b")
        empty = doc.add_paragraph("")
        add_direct_numbering(empty)
        doc.add_table(rows=1, cols=1)
        template = {
            "styles": {
                "body": {"font": "\u5b8b\u4f53", "size_pt": 12},
                "table_caption": {"font": "\u5b8b\u4f53", "size_pt": 10.5},
                "table_text": {"font": "\u5b8b\u4f53", "size_pt": 10.5},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"paragraphs": [], "tables": [], "warnings": []}

        format_docx.format_normal_paragraphs(doc, template, report)
        format_docx.format_tables(doc, template, report)

        self.assertEqual(empty.text, "")
        self.assertIsNone(empty._p.pPr.find(format_docx.qn("w:numPr")))

    def test_clears_empty_numbered_paragraph_adjacent_to_table(self):
        doc = Document()
        doc.add_table(rows=1, cols=1)
        empty = doc.add_paragraph("")
        add_direct_numbering(empty)
        template = {
            "styles": {
                "body": {"font": "\u5b8b\u4f53", "size_pt": 12},
                "table_text": {"font": "\u5b8b\u4f53", "size_pt": 10.5},
            },
            "_runtime": {"fallback_count": 0, "warned_missing_styles": set()},
        }
        report = {"tables": [], "warnings": []}

        format_docx.format_tables(doc, template, report)

        self.assertEqual(empty.text, "")
        self.assertIsNone(empty._p.pPr.find(format_docx.qn("w:numPr")))

    def test_validate_format_rules_accepts_valid_template(self):
        template = {
            "version": "1.0",
            "name": "test",
            "description": "测试模板",
            "page": {"top_margin_cm": 2.5},
            "styles": {
                "body": {
                    "font": "宋体",
                    "size_pt": 12,
                    "bold": False,
                    "alignment": "justify",
                    "line_spacing": 1.5,
                    "first_line_indent_pt": 24,
                    "space_before_pt": 0,
                    "space_after_pt": 0,
                }
            },
        }
        errors, warnings = format_docx.validate_format_rules(template)
        self.assertEqual(errors, [])
        self.assertEqual(warnings, [])

    def test_validate_format_rules_rejects_bad_size_pt(self):
        template = {
            "version": "1.0",
            "name": "bad",
            "page": {},
            "styles": {"body": {"size_pt": "12"}},
        }
        errors, _ = format_docx.validate_format_rules(template)
        self.assertTrue(any("styles.body.size_pt" in error for error in errors))

    def test_validate_format_rules_rejects_bad_alignment(self):
        template = {
            "version": "1.0",
            "name": "bad",
            "page": {},
            "styles": {"body": {"alignment": "middle"}},
        }
        errors, _ = format_docx.validate_format_rules(template)
        self.assertTrue(any("styles.body.alignment" in error for error in errors))

    def test_validate_override_rules_accepts_partial_body_font(self):
        override = {"styles": {"body": {"font": "微软雅黑"}}}
        normalized = format_docx.normalize_format_rules(override, fill_defaults=False)
        errors, _ = format_docx.validate_override_rules(normalized)
        self.assertEqual(errors, [])

    def test_validate_override_rules_does_not_require_page(self):
        errors, _ = format_docx.validate_override_rules({"styles": {"heading_1": {"font": "黑体"}}})
        self.assertEqual(errors, [])

    def test_validate_override_rules_does_not_require_styles_body(self):
        errors, _ = format_docx.validate_override_rules({"styles": {"heading_1": {"alignment": "center"}}})
        self.assertEqual(errors, [])

    def test_validate_override_rules_rejects_bad_alignment(self):
        errors, _ = format_docx.validate_override_rules({"styles": {"body": {"alignment": "middle"}}})
        self.assertTrue(any("styles.body.alignment" in error for error in errors))

    def test_validate_override_rules_accepts_underline(self):
        errors, _ = format_docx.validate_override_rules(
            {"styles": {"body": {"underline": True}}}
        )
        self.assertEqual(errors, [])

    def test_validate_override_rules_accepts_caption_pagination_controls(self):
        errors, _ = format_docx.validate_override_rules(
            {
                "styles": {
                    "table_caption": {
                        "keep_with_next": True,
                        "keep_together": True,
                    }
                }
            }
        )
        self.assertEqual(errors, [])

    def test_merge_format_rules_preserves_override_warnings(self):
        base = format_docx.load_template("default")
        warning = "正文 font 存在冲突：宋体 vs 微软雅黑，已暂按后者覆盖。"
        override = format_docx.normalize_format_rules(
            {
                "name": "teacher",
                "warnings": [warning],
                "styles": {"body": {"font": "微软雅黑"}},
            },
            fill_defaults=False,
        )

        merged, metadata = format_docx.merge_format_rules(base, override)
        debug_summary = format_docx.build_debug_summary(merged)

        self.assertIn(warning, metadata["warnings"])
        self.assertIn(warning, merged["_template_warnings"])
        self.assertIn(warning, debug_summary["override_warnings"])

    def test_merge_format_rules_overrides_body_font(self):
        base = format_docx.load_template("default")
        override = format_docx.normalize_format_rules(
            {"name": "teacher", "styles": {"body": {"font": "微软雅黑"}}},
            fill_defaults=False,
        )
        merged, metadata = format_docx.merge_format_rules(base, override)
        self.assertEqual(merged["styles"]["body"]["font"], "微软雅黑")
        self.assertIn("styles.body.font", metadata["overridden_fields"])

    def test_merge_format_rules_keeps_uncovered_fields(self):
        base = format_docx.load_template("default")
        base_line_spacing = base["styles"]["body"]["line_spacing"]
        override = format_docx.normalize_format_rules(
            {"name": "teacher", "styles": {"body": {"font": "微软雅黑"}}},
            fill_defaults=False,
        )
        merged, _ = format_docx.merge_format_rules(base, override)
        self.assertEqual(merged["styles"]["body"]["line_spacing"], base_line_spacing)

    def test_merge_format_rules_overrides_page_left_margin(self):
        base = format_docx.load_template("default")
        override = {"name": "teacher", "page": {"left_margin_cm": 3.2}}
        merged, metadata = format_docx.merge_format_rules(base, override)
        self.assertEqual(merged["page"]["left_margin_cm"], 3.2)
        self.assertIn("page.left_margin_cm", metadata["overridden_fields"])

    def test_merge_format_rules_does_not_mutate_base(self):
        base = format_docx.load_template("default")
        original_font = base["styles"]["body"]["font"]
        override = {"name": "teacher", "styles": {"body": {"font": "微软雅黑"}}}
        format_docx.merge_format_rules(base, override)
        self.assertEqual(base["styles"]["body"]["font"], original_font)

    def test_merge_format_rules_returns_overridden_fields(self):
        base = format_docx.load_template("default")
        override = format_docx.normalize_format_rules(
            {
                "name": "teacher",
                "styles": {
                    "body": {"font": "微软雅黑", "size_cn": "五号"},
                    "heading_1": {"alignment": "居中"},
                },
            },
            fill_defaults=False,
        )
        merged, metadata = format_docx.merge_format_rules(base, override)
        self.assertEqual(merged["styles"]["body"]["size_pt"], 10.5)
        self.assertIn("styles.body.font", metadata["overridden_fields"])
        self.assertIn("styles.body.size_pt", metadata["overridden_fields"])
        self.assertIn("styles.heading_1.alignment", metadata["overridden_fields"])

    def test_merge_format_rules_preserves_v04_parse_metadata(self):
        base = format_docx.load_template("default")
        override = format_docx.normalize_override_rules(
            {
                "name": "v04_parse",
                "styles": {
                    "abstract_cn_title": {
                        "font": "黑体",
                        "size_cn": "小四",
                    },
                    "reference_item": {
                        "font": "宋体",
                        "size_cn": "五号",
                    },
                },
                "reference_latin_digit_format": {
                    "font": "Times New Roman",
                    "size_cn": "五号",
                    "scope": "reference",
                },
                "toc": {"action": "protect"},
                "unsupported_modules": {
                    "page_number": {
                        "status": "detected_but_not_supported",
                        "action": "warning_only",
                        "note": "检测到页码要求，但当前版本暂不处理。",
                    }
                },
                "parser_metadata": {
                    "parser_mode": "kimi",
                    "raw_ai_rules": {"name": "raw"},
                    "validated_rules": {"name": "validated"},
                    "conflict_result": {"hasConflicts": False},
                },
            }
        )

        merged, metadata = format_docx.merge_format_rules(base, override)
        self.assertEqual(merged["reference_latin_digit_format"]["font"], "Times New Roman")
        self.assertEqual(merged["reference_latin_digit_format"]["size_pt"], 10.5)
        self.assertEqual(merged["toc"]["action"], "protect")
        self.assertIn("page_number", merged["unsupported_modules"])
        self.assertEqual(merged["parser_metadata"]["parser_mode"], "kimi")
        self.assertIn("reference_latin_digit_format.font", metadata["overridden_fields"])
        debug_summary = format_docx.build_debug_summary(merged)
        self.assertEqual(debug_summary["parser_mode"], "kimi")
        self.assertIn("unsupported_modules", debug_summary)

    def test_load_template_default(self):
        template = format_docx.load_template("default")
        self.assertEqual(template["name"], "default")
        self.assertIn("body", template["styles"])

    def test_validate_template_accepts_valid_template(self):
        template = {
            "version": "1.0",
            "name": "test",
            "page": {"top_margin_cm": 2.5},
            "styles": {
                "body": {
                    "font": "宋体",
                    "size_pt": 12,
                    "bold": False,
                    "alignment": "justify",
                    "line_spacing": 1.5,
                    "first_line_indent_pt": 24,
                    "space_before_pt": 0,
                    "space_after_pt": 0,
                }
            },
        }
        format_docx.validate_template(template)

    def test_validate_template_rejects_bad_style_type(self):
        template = {
            "version": "1.0",
            "name": "test",
            "page": {},
            "styles": {"body": {"size_pt": "12"}},
        }
        with self.assertRaises(format_docx.TemplateError):
            format_docx.validate_template(template)

    def test_validate_input_file_missing_raises(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            missing = Path(temp_dir) / "missing.docx"
            with self.assertRaises(format_docx.UserFacingError):
                format_docx.validate_input_file(missing)

    def test_validate_output_path_same_path_raises(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            input_path = Path(temp_dir) / "input.docx"
            input_path.write_bytes(b"fake")
            with self.assertRaises(format_docx.UserFacingError):
                format_docx.validate_output_path(input_path, input_path, overwrite=True)

    def test_make_text_preview_truncates(self):
        text = "这是一个用于测试截断功能的长文本" * 5
        preview = format_docx.make_text_preview(text, max_len=10)
        self.assertEqual(len(preview), 13)
        self.assertTrue(preview.endswith("..."))

    def test_save_report_generates_json(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            report_path = Path(temp_dir) / "reports" / "report.json"
            report = {
                "version": "test",
                "input_file": "in.docx",
                "output_file": "out.docx",
                "template": {"name": "default", "description": "test"},
                "stats": {},
                "paragraphs": [],
                "tables": [],
                "warnings": [],
            }
            format_docx.save_report(report, report_path)
            self.assertTrue(report_path.exists())
            data = json.loads(report_path.read_text(encoding="utf-8"))
            self.assertEqual(data["version"], "test")

    def test_old_flow_without_report_still_works(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            input_path = temp_path / "input.docx"
            output_path = temp_path / "output.docx"

            doc = Document()
            doc.add_paragraph("测试论文标题")
            doc.add_paragraph("一、研究背景")
            doc.add_paragraph("这是一段用于测试旧流程的正文。")
            doc.save(str(input_path))

            result = subprocess.run(
                [
                    sys.executable,
                    str(PROJECT_ROOT / "format_docx.py"),
                    str(input_path),
                    str(output_path),
                    "--overwrite",
                ],
                cwd=str(PROJECT_ROOT),
                text=True,
                capture_output=True,
                check=False,
            )

            self.assertEqual(result.returncode, 0, msg=result.stderr)
            self.assertTrue(output_path.exists())
            self.assertFalse((temp_path / "report.json").exists())

    def test_validate_template_command_default_passes(self):
        result = subprocess.run(
            [
                sys.executable,
                str(PROJECT_ROOT / "format_docx.py"),
                "--validate-template",
                "templates/default.json",
            ],
            cwd=str(PROJECT_ROOT),
            text=True,
            capture_output=True,
            check=False,
        )

        self.assertEqual(result.returncode, 0, msg=result.stderr)
        self.assertIn("模板校验通过", result.stdout)

    def test_merge_rules_command_exports_final_rules(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            final_rules_path = Path(temp_dir) / "final_rules.json"
            result = subprocess.run(
                [
                    sys.executable,
                    str(PROJECT_ROOT / "format_docx.py"),
                    "--template",
                    "default",
                    "--override",
                    "overrides/teacher_override.json",
                    "--export-final-rules",
                    str(final_rules_path),
                    "--merge-rules",
                ],
                cwd=str(PROJECT_ROOT),
                text=True,
                capture_output=True,
                check=False,
            )

            self.assertEqual(result.returncode, 0, msg=result.stderr)
            self.assertTrue(final_rules_path.exists())
            final_rules = json.loads(final_rules_path.read_text(encoding="utf-8"))
            errors, _ = format_docx.validate_format_rules(final_rules)
            self.assertEqual(errors, [])

    def test_merge_rules_command_exports_debug_rules_for_margin_and_references(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            override_path = temp_path / "override.json"
            debug_rules_path = temp_path / "debug_rules.json"
            override_path.write_text(
                json.dumps(
                    {
                        "name": "teacher",
                        "page": {"left_margin_cm": "3cm"},
                        "styles": {
                            "reference_title": {"alignment": "居中"},
                            "reference_item": {"font": "宋体", "size_cn": "五号"},
                        },
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result = subprocess.run(
                [
                    sys.executable,
                    str(PROJECT_ROOT / "format_docx.py"),
                    "--template",
                    "default",
                    "--override",
                    str(override_path),
                    "--export-debug-rules",
                    str(debug_rules_path),
                    "--merge-rules",
                ],
                cwd=str(PROJECT_ROOT),
                text=True,
                capture_output=True,
                check=False,
            )

            self.assertEqual(result.returncode, 0, msg=result.stderr)
            debug_rules = json.loads(debug_rules_path.read_text(encoding="utf-8"))
            self.assertEqual(debug_rules["normalized_override"]["page"]["left_margin_cm"], 3)
            self.assertEqual(debug_rules["final_rules"]["page"]["left_margin_cm"], 3)
            self.assertEqual(
                debug_rules["final_rules"]["styles"]["reference_title"]["alignment"],
                "center",
            )
            self.assertNotEqual(
                debug_rules["final_rules"]["styles"]["reference_item"]["alignment"],
                "center",
            )

    def test_format_document_applies_left_margin_and_keeps_reference_item_left(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            input_path = temp_path / "input.docx"
            output_path = temp_path / "output.docx"
            report_path = temp_path / "report.json"
            override_path = temp_path / "override.json"

            doc = Document()
            doc.add_paragraph("测试论文标题")
            doc.add_paragraph("参考文献")
            doc.add_paragraph("[1] 张三. 测试文献. 2024.")
            doc.save(str(input_path))

            override_path.write_text(
                json.dumps(
                    {
                        "name": "teacher",
                        "page": {"left_margin_cm": 3},
                        "styles": {
                            "reference_title": {"alignment": "center"},
                            "reference_item": {"font": "宋体", "size_cn": "五号"},
                        },
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result = subprocess.run(
                [
                    sys.executable,
                    str(PROJECT_ROOT / "format_docx.py"),
                    str(input_path),
                    str(output_path),
                    "--template",
                    "default",
                    "--override",
                    str(override_path),
                    "--report",
                    str(report_path),
                    "--overwrite",
                ],
                cwd=str(PROJECT_ROOT),
                text=True,
                capture_output=True,
                check=False,
            )

            self.assertEqual(result.returncode, 0, msg=result.stderr)
            output_doc = Document(str(output_path))
            self.assertAlmostEqual(output_doc.sections[0].left_margin.cm, 3, places=2)
            self.assertEqual(output_doc.paragraphs[1].alignment, WD_ALIGN_PARAGRAPH.CENTER)
            self.assertNotEqual(output_doc.paragraphs[2].alignment, WD_ALIGN_PARAGRAPH.CENTER)

            report = json.loads(report_path.read_text(encoding="utf-8"))
            self.assertEqual(
                report["final_rules_debug"]["final_rules"]["page"]["left_margin_cm"],
                3,
            )
            reference_item = report["final_rules_debug"]["final_rules"]["styles"][
                "reference_item"
            ]
            self.assertNotEqual(reference_item["alignment"], "center")
            self.assertEqual(report["paragraphs"][1]["detected_type"], "reference_title")
            self.assertEqual(report["paragraphs"][2]["detected_type"], "reference_item")
            self.assertEqual(report["paragraphs"][2]["applied_style"], "reference_item")

    def test_format_document_exits_reference_mode_before_appendix(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            input_path = temp_path / "input.docx"
            output_path = temp_path / "output.docx"
            report_path = temp_path / "report.json"
            override_path = temp_path / "override.json"

            doc = Document()
            doc.add_paragraph("测试论文标题")
            doc.add_paragraph("参考文献")
            doc.add_paragraph("[1] 张三. 测试文献. 2024.")
            doc.add_paragraph(
                "1. 注意：这一条在参考文献区域内，虽然以数字编号开头，但不应该被识别成正文标题。"
            )
            doc.add_paragraph("附录")
            doc.add_paragraph("一、访谈提纲")
            doc.add_paragraph("1. 你平时会在哪些学习场景中使用AI工具？")
            doc.add_paragraph("二、问卷题项")
            doc.add_paragraph("本部分包含问卷题项示例，应作为附录正文处理，而不是参考文献条目。")
            doc.save(str(input_path))
            override_path.write_text(
                json.dumps(
                    {
                        "name": "reference_single_spacing",
                        "styles": {
                            "reference_item": {
                                "font": "宋体",
                                "size_cn": "五号",
                                "alignment": "left",
                                "line_spacing": 1,
                            },
                            "body": {"line_spacing": 1.5},
                            "heading_1": {"line_spacing": 1.5},
                        },
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result = subprocess.run(
                [
                    sys.executable,
                    str(PROJECT_ROOT / "format_docx.py"),
                    str(input_path),
                    str(output_path),
                    "--template",
                    "default",
                    "--override",
                    str(override_path),
                    "--report",
                    str(report_path),
                    "--overwrite",
                ],
                cwd=str(PROJECT_ROOT),
                text=True,
                capture_output=True,
                check=False,
            )

            self.assertEqual(result.returncode, 0, msg=result.stderr)
            output_doc = Document(str(output_path))
            report = json.loads(report_path.read_text(encoding="utf-8"))
            detected_types = [
                item["detected_type"]
                for item in report["paragraphs"]
            ]

            self.assertEqual(detected_types[1], "reference_title")
            self.assertEqual(detected_types[2], "reference_item")
            self.assertEqual(detected_types[3], "reference_item")
            self.assertNotEqual(detected_types[4], "reference_item")
            self.assertNotEqual(detected_types[5], "reference_item")
            self.assertNotEqual(detected_types[6], "reference_item")
            self.assertEqual(output_doc.paragraphs[2].paragraph_format.line_spacing, 1.0)
            self.assertEqual(output_doc.paragraphs[3].paragraph_format.line_spacing, 1.0)
            self.assertNotEqual(output_doc.paragraphs[5].paragraph_format.line_spacing, 1.0)
            module_status = report["module_status"]
            self.assertEqual(module_status["reference"]["status"], "detected")
            self.assertEqual(module_status["appendix"]["status"], "detected")
            self.assertEqual(module_status["appendix"]["action"], "boundary_only")


if __name__ == "__main__":
    unittest.main()
