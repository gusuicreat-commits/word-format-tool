import tempfile
import unittest
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement

import format_docx as f
from review_workflow import preview, validate_plan


class StructureReviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name) / "input.docx"
        self.output = Path(self.temp.name) / "output.docx"
        self.doc = Document()
        self.doc.add_paragraph("Research title")
        self.doc.add_paragraph("Experimental setup.")
        self.doc.add_paragraph("Protected figure caption.").runs[0].font.size = Pt(19)
        self.doc.add_table(rows=1, cols=1).cell(0, 0).text = "Protected data"
        self.doc.save(self.source)
        self.rules = f.get_final_format_rules("default", print_warnings=False)

    def test_preview_does_not_modify_source_and_preserves_identity(self):
        original = self.source.read_bytes()
        plan = preview(self.source, self.rules)
        self.assertEqual(self.source.read_bytes(), original)
        self.assertEqual([item["key"] for item in plan["items"]], ["p:0", "p:1", "p:2", "t:0"])
        self.assertFalse(self.output.exists())

    def test_user_roles_and_protection_reach_saved_document(self):
        plan = preview(self.source, self.rules)
        plan["items"][1]["type"] = "heading_3"
        plan["items"][2]["preserve"] = True
        plan["items"][3]["preserve"] = True
        decisions = validate_plan(plan, self.source, self.rules)
        report = f.init_report(self.source, self.output, self.rules)
        f.format_document(self.source, self.output, self.rules, report, review=decisions)
        output = Document(self.output)
        self.assertEqual(f.detect_heading_type_from_style(output.paragraphs[1]), "heading_3")
        self.assertEqual(output.paragraphs[2]._p.xml, self.doc.paragraphs[2]._p.xml)
        self.assertEqual(output.tables[0]._tbl.xml, self.doc.tables[0]._tbl.xml)
        self.assertEqual(output.sections[0]._sectPr.xml, self.doc.sections[0]._sectPr.xml)

    def test_changed_file_rejects_confirmation(self):
        plan = preview(self.source, self.rules)
        self.doc.add_paragraph("Changed content")
        self.doc.save(self.source)
        with self.assertRaises(f.UserFacingError):
            validate_plan(plan, self.source, self.rules)

    def test_automatic_table_protection_preserves_only_affected_section(self):
        self.doc.sections[0].left_margin = Cm(1)
        self.doc.sections[0].right_margin = Cm(1)
        self.doc.tables[0].columns[0].width = Cm(18)
        self.doc.add_section(WD_SECTION_START.NEW_PAGE)
        self.doc.add_paragraph("Body in the next section.")
        self.doc.save(self.source)
        report = f.init_report(self.source, self.output, self.rules)
        f.format_document(self.source, self.output, self.rules, report)
        output = Document(self.output)
        self.assertEqual(output.tables[0]._tbl.xml, self.doc.tables[0]._tbl.xml)
        self.assertEqual(output.sections[0].left_margin, self.doc.sections[0].left_margin)
        self.assertAlmostEqual(output.sections[1].left_margin.cm, 2.5, places=2)
        self.assertEqual(report["automatic_protection"]["tables"], [0])
        self.assertEqual(report["preserved_sections"], [0])
        self.assertEqual(report["module_status"]["page"]["action"], "partial")

    def test_automatic_floating_content_protection(self):
        self.doc.paragraphs[2].runs[0]._r.append(OxmlElement("w:pict"))
        self.doc.save(self.source)
        report = f.init_report(self.source, self.output, self.rules)
        f.format_document(self.source, self.output, self.rules, report)
        output = Document(self.output)
        self.assertEqual(output.paragraphs[2]._p.xml, self.doc.paragraphs[2]._p.xml)
        self.assertEqual(output.sections[0]._sectPr.xml, self.doc.sections[0]._sectPr.xml)
        self.assertEqual(report["automatic_protection"]["paragraphs"], [2])
        self.assertEqual(report["module_status"]["page"]["action"], "skipped")

    def test_plain_paragraph_protection_does_not_freeze_margins(self):
        table = self.doc.tables[0]._tbl
        table.getparent().remove(table)
        self.doc.save(self.source)
        plan = preview(self.source, self.rules)
        plan["items"][2]["preserve"] = True
        report = f.init_report(self.source, self.output, self.rules)
        f.format_document(self.source, self.output, self.rules, report,
                          review=validate_plan(plan, self.source, self.rules))
        output = Document(self.output)
        self.assertEqual(output.paragraphs[2]._p.xml, self.doc.paragraphs[2]._p.xml)
        self.assertEqual(report["preserved_sections"], [])
        self.assertAlmostEqual(output.sections[0].left_margin.cm, 2.5, places=2)

    def test_user_can_demote_heading_to_body(self):
        self.doc.paragraphs[1].style = "Heading 2"
        self.doc.save(self.source)
        plan = preview(self.source, self.rules)
        plan["items"][1]["type"] = "body"
        report = f.init_report(self.source, self.output, self.rules)
        f.format_document(self.source, self.output, self.rules, report,
                          review=validate_plan(plan, self.source, self.rules))
        self.assertIsNone(f.detect_heading_type_from_style(Document(self.output).paragraphs[1]))
        self.assertEqual(report["paragraphs"][1]["detected_type"], "body")

    def test_changed_rules_reject_confirmation(self):
        plan = preview(self.source, self.rules)
        rules = deepcopy(self.rules)
        rules["styles"]["body"]["size_pt"] = 13
        with self.assertRaises(f.UserFacingError):
            validate_plan(plan, self.source, rules)

    def test_duplicate_missing_and_invalid_roles_rejected(self):
        original = preview(self.source, self.rules)
        for change in ("duplicate", "missing", "invalid"):
            plan = deepcopy(original)
            if change == "duplicate": plan["items"][1] = plan["items"][0]
            if change == "missing": plan["items"].pop()
            if change == "invalid": plan["items"][1]["type"] = "execute_code"
            with self.subTest(change=change), self.assertRaises(f.UserFacingError):
                validate_plan(plan, self.source, self.rules)


if __name__ == "__main__":
    unittest.main()
