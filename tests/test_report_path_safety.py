"""CLI report collisions must never damage a real DOCX."""
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document

import format_docx as f
import review_workflow as review


ROOT = Path(__file__).resolve().parents[1]


class ReportPathSafetyTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.source = self.directory / "source.docx"
        self.output = self.directory / "output.docx"
        for path, text in ((self.source, "Original source"), (self.output, "Existing output")):
            doc = Document()
            doc.add_paragraph(text)
            doc.save(path)
        self.originals = {path: path.read_bytes() for path in (self.source, self.output)}
        rules = f.get_final_format_rules("default", print_warnings=False)
        self.plan = self.directory / "plan.json"
        self.plan.write_text(json.dumps(review.preview(self.source, rules)), encoding="utf-8")

    def run_cli(self, mode, report, output=None):
        script = "format_docx.py" if mode == "format" else "review_workflow.py"
        args = [sys.executable, "-B", "-X", "utf8", str(ROOT / script), str(self.source),
                str(output or self.output), "--report", str(report), "--overwrite"]
        if mode == "analyze":
            args.append("--analyze")
        elif mode == "review":
            args.extend(["--review-plan", str(self.plan)])
        return subprocess.run(args, cwd=self.directory, capture_output=True)

    def assert_documents_unchanged(self):
        for path, content in self.originals.items():
            self.assertEqual(path.read_bytes(), content)
            self.assertTrue(Document(path).paragraphs[0].text)

    def check_collision(self, alias):
        for mode in ("format", "analyze", "review"):
            for target in (self.source, self.output):
                with self.subTest(mode=mode, target=target.name):
                    report = target
                    if alias == "relative":
                        report = Path("subdir") / ".." / target.name
                        (self.directory / "subdir").mkdir(exist_ok=True)
                    elif alias == "hardlink":
                        report = self.directory / (target.stem + "-alias.json")
                        if not report.exists():
                            os.link(target, report)
                    result = self.run_cli(mode, report)
                    self.assertNotEqual(result.returncode, 0, result.stdout)
                    self.assertIn("报告文件不能".encode("utf-8"), result.stderr)
                    self.assert_documents_unchanged()

    def test_same_path_rejected(self):
        self.check_collision("same")

    def test_relative_alias_rejected(self):
        self.check_collision("relative")

    def test_hardlink_alias_rejected(self):
        self.check_collision("hardlink")

    def test_new_output_collision_rejected_before_directory_creation(self):
        for mode in ("format", "analyze", "review"):
            with self.subTest(mode=mode):
                output = self.directory / mode / "new.docx"
                result = self.run_cli(mode, output, output=output)
                self.assertNotEqual(result.returncode, 0)
                self.assertFalse(output.parent.exists())
                self.assert_documents_unchanged()

    def test_independent_report_succeeds(self):
        for mode in ("format", "analyze", "review"):
            with self.subTest(mode=mode):
                report = self.directory / (mode + ".json")
                output = self.directory / (mode + ".docx")
                result = self.run_cli(mode, report, output=output)
                self.assertEqual(result.returncode, 0, result.stderr)
                payload = json.loads(report.read_text(encoding="utf-8"))
                if mode == "analyze":
                    self.assertIn("items", payload)
                    self.assertFalse(output.exists())
                else:
                    self.assertIn("stats", payload)
                    self.assertEqual(Document(output).paragraphs[0].text, "Original source")
                self.assert_documents_unchanged()


if __name__ == "__main__":
    unittest.main()
