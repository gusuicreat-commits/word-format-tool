"""Opt-in local audit of downloaded DOCX papers; no network or private fixtures."""

import argparse
import hashlib
import json
import posixpath
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

import format_docx as formatter
from tests.test_document_corpus import protected_snapshot


def normalize_part(name, data):
    if not name.endswith((".xml", ".rels")):
        return data
    root = etree.fromstring(data)
    if name.endswith(".rels"):
        base = posixpath.dirname(posixpath.dirname(name))
        for relation in root:
            if relation.get("TargetMode") != "External":
                target = relation.get("Target", "")
                relation.set("Target", posixpath.normpath(posixpath.join("/", base, target)))
        root[:] = sorted(root, key=lambda item: item.get("Id", ""))
    return etree.tostring(root, method="c14n")


def audit(source, template):
    output = source.with_name(f"{source.stem}-{template}-output.docx")
    rules = formatter.get_final_format_rules(template, print_warnings=False)
    report = formatter.init_report(source, output, rules)
    formatter.format_document(source, output, rules, report, overwrite=True)
    formatter.save_report(report, output.with_suffix(".json"))
    before = protected_snapshot(source)
    after = protected_snapshot(output)
    # XML declarations/attribute ordering may change during a lossless save.
    def normalize_parts(parts):
        return {name: normalize_part(name, data) for name, data in parts.items()}
    before = (*before[:2], normalize_parts(before[2]))
    after = (*after[:2], normalize_parts(after[2]))
    doc = Document(source)
    formatted = Document(output)
    with ZipFile(source) as archive:
        source_files = {n for n in archive.namelist() if not n.endswith("/")}
        root = etree.fromstring(archive.read("word/document.xml"))
        counts = {tag: sum(1 for _ in root.iter(qn(tag))) for tag in
                  ("w:p", "w:tbl", "w:hyperlink", "w:fldChar", "w:drawing", "m:oMath", "w:numPr", "w:footnoteReference", "w:endnoteReference")}
    checks = dict(zip(("paragraph_text_preserved", "protected_xml_preserved", "protected_parts_preserved"),
                      (a == b for a, b in zip(before, after))))
    body = [p for p in report["paragraphs"] if p["detected_type"] == "body" and p["applied_style"] == "body"]
    checks["body_line_spacing_applied"] = all(
        formatted.paragraphs[p["index"]].paragraph_format.line_spacing == rules["styles"]["body"]["line_spacing"]
        for p in body) if body else None
    body_runs = [run for item in body for run in formatted.paragraphs[item["index"]].runs if run.text.strip()]
    checks["body_size_applied"] = all(run.font.size is not None and run.font.size.pt == rules["styles"]["body"]["size_pt"] for run in body_runs) if body_runs else None
    checks["page_margins_applied"] = all(
        abs(getattr(section, field.replace("_cm", "")).cm - value) < 0.01
        for index, section in enumerate(formatted.sections) if index not in report.get("preserved_sections", [])
        for field, value in rules.get("page", {}).items()
        if field in {"top_margin_cm", "bottom_margin_cm", "left_margin_cm", "right_margin_cm"})
    checks["protected_section_margins_preserved"] = all(
        getattr(formatted.sections[index], field) == getattr(doc.sections[index], field)
        for index in report.get("preserved_sections", [])
        for field in ("top_margin", "bottom_margin", "left_margin", "right_margin"))
    with ZipFile(output) as archive:
        output_files = {n for n in archive.namelist() if not n.endswith("/")}
        first_xml = normalize_part("word/document.xml", archive.read("word/document.xml"))
    checks["all_package_files_preserved"] = source_files == output_files
    checks["table_count_matches_xml"] = report["stats"]["table"] == counts["w:tbl"]
    reference_headings = [i for i, p in enumerate(doc.paragraphs)
                          if p.text.strip().rstrip(":").casefold() in ("references", "bibliography", "参考文献")]
    checks["reference_heading_recognized"] = all(
        any(p["index"] == i and p["detected_type"] == "reference_title" for p in report["paragraphs"])
        for i in reference_headings) if reference_headings else None
    recheck_output = output.with_name(f"{source.stem}-{template}-recheck-output.docx")
    recheck_rules = formatter.get_final_format_rules(template, print_warnings=False)
    recheck_report = formatter.init_report(output, recheck_output, recheck_rules)
    formatter.format_document(output, recheck_output, recheck_rules, recheck_report, overwrite=True)
    with ZipFile(recheck_output) as archive:
        checks["second_pass_document_stable"] = first_xml == normalize_part("word/document.xml", archive.read("word/document.xml"))
    checks["second_pass_classification_stable"] = [p["detected_type"] for p in report["paragraphs"]] == [p["detected_type"] for p in recheck_report["paragraphs"]]
    return {"file": source.name, "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "template": template, "structure_counts": counts, "checks": checks,
            "body_paragraphs_checked": len(body), "reference_heading_indices": reference_headings,
            "stats": report["stats"], "output": str(output), "visual_review": "not_performed",
            "manual_review_flags": [label for key, label in (("paper_title", "No paper title identified"), ("heading_1", "No level-one heading identified")) if not report["stats"].get(key)]}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("directory", type=Path)
    args = parser.parse_args()
    sources = sorted(p for p in args.directory.glob("*.docx") if not p.stem.endswith("-output"))
    if not sources:
        parser.error("No input DOCX files found")
    results = [audit(source, template) for source in sources for template in ("default", "course_paper")]
    (args.directory / "audit.json").write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    for result in results:
        print(json.dumps({k: result[k] for k in ("file", "template", "structure_counts", "checks", "stats")}, ensure_ascii=True))
    if any(value is False for result in results for value in result["checks"].values()):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
