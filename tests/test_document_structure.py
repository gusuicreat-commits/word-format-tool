"""Structural analysis is read-only; formatting preserves shared dependencies."""
import json
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Cm, Pt
from docx.opc.constants import RELATIONSHIP_TYPE as RT

import document_structure as structure
import format_docx as formatter
from tests.test_document_corpus import element
from tests.test_recognition_review import numbered


class DocumentStructureTests(unittest.TestCase):
    def process(self, doc, template="default"):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "source.docx", Path(directory) / "output.docx"
            doc.save(source)
            rules = formatter.get_final_format_rules(template, print_warnings=False)
            report = formatter.init_report(source, output, rules)
            formatter.format_document(source, output, rules, report)
            with ZipFile(source) as before, ZipFile(output) as after:
                self.assertEqual(before.read("word/styles.xml"), after.read("word/styles.xml"))
                self.assertEqual(before.read("word/numbering.xml"), after.read("word/numbering.xml"))
            json.dumps(report, ensure_ascii=False)
            return Document(output), report

    def custom_style(self, doc, base="Heading 1"):
        style = doc.styles.add_style("SchoolSection", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles[base]
        return style

    def test_analysis_does_not_mutate_any_package_part(self):
        doc = Document()
        doc.add_paragraph("Study", style=self.custom_style(doc))
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Data"
        before = {str(p.partname): p.blob for p in doc.part.package.parts}
        analysis = structure.analyze_document(doc)
        self.assertEqual(before, {str(p.partname): p.blob for p in doc.part.package.parts})
        self.assertEqual(analysis["paragraphs"][0]["heading"]["source"], "style:Heading1")
        self.assertEqual(analysis["paragraphs"][1]["container"], "table_cell")
        self.assertNotIn("/word/header1.xml", analysis["parts"])
        json.dumps(analysis)

    def test_custom_heading_keeps_original_style(self):
        doc = Document()
        doc.add_paragraph("Study title")
        doc.add_paragraph("Experimental design", style=self.custom_style(doc))
        output, report = self.process(doc)
        self.assertEqual(report["paragraphs"][1]["detected_type"], "heading_1")
        self.assertEqual(output.paragraphs[1].style.style_id, "SchoolSection")
        self.assertEqual(report["paragraphs"][1]["structural_heading_evidence"]["source"], "style:Heading1")

    def test_analysis_does_not_create_missing_styles_part(self):
        doc = Document()
        doc.add_paragraph("Text")
        rel_id = next(key for key, rel in doc.part.rels.items() if rel.reltype == RT.STYLES)
        doc.part.drop_rel(rel_id)
        before = {str(p.partname): p.blob for p in doc.part.package.parts}
        analysis = structure.analyze_document(doc)
        self.assertEqual(analysis["paragraphs"][0]["style_chain"], [])
        self.assertEqual(before, {str(p.partname): p.blob for p in doc.part.package.parts})

    def test_direct_outline_overrides_inherited_heading(self):
        doc = Document()
        p = doc.add_paragraph("Study", style=self.custom_style(doc))
        outline = element("w:outlineLvl", val="2")
        p._p.get_or_add_pPr().append(outline)
        self.assertEqual(formatter.detect_heading_type_from_style(p), "heading_3")
        outline.set(formatter.qn("w:val"), "9")
        self.assertIsNone(formatter.detect_heading_type_from_style(p))

    def test_style_name_substring_is_not_a_heading(self):
        doc = Document()
        p = doc.add_paragraph("Text", style=doc.styles.add_style("NotHeading1", WD_STYLE_TYPE.PARAGRAPH))
        self.assertIsNone(formatter.detect_heading_type_from_style(p))

    def test_style_cycle_is_bounded(self):
        doc = Document()
        first = self.custom_style(doc, "Normal")
        second = doc.styles.add_style("SecondStyle", WD_STYLE_TYPE.PARAGRAPH)
        first.base_style, second.base_style = second, first
        p = doc.add_paragraph("Text", style=first)
        self.assertEqual(len(list(structure.style_chain(p))), 2)
        self.assertIsNone(formatter.detect_heading_type_from_style(p))

    def test_inherited_numbering_is_recognized_without_direct_numpr(self):
        doc = Document()
        doc.add_paragraph("Study title")
        p = numbered(doc, "System architecture")
        style = self.custom_style(doc, "Normal")
        numpr = p._p.pPr.find(formatter.qn("w:numPr"))
        style.element.get_or_add_pPr().append(numpr)
        p.style = style
        definition = next(n for n in doc.part.numbering_part.element.findall(formatter.qn("w:abstractNum"))
                          if n.get(formatter.qn("w:abstractNumId")) == "80")
        definition.find(formatter.qn("w:lvl")).append(element("w:pStyle", val=style.style_id))
        doc.add_paragraph("\uff08\u4e00\uff09Design")
        self.assertEqual(structure.effective_numbering(p)["sources"]["num_id"], "style:SchoolSection")
        output, report = self.process(doc)
        self.assertEqual(report["paragraphs"][1]["detected_type"], "heading_1")
        self.assertEqual(structure.effective_numbering(output.paragraphs[1])["num_id"], "80")
        self.assertEqual(report["structure_analysis"]["paragraphs"][1]["style_chain"][0], "SchoolSection")

    def test_direct_numbering_disable_wins_over_inherited_numbering(self):
        doc = Document()
        p = doc.add_paragraph("Text", style="List Number")
        numpr = element("w:numPr")
        numpr.append(element("w:numId", val="0"))
        p._p.get_or_add_pPr().append(numpr)
        numbering = structure.effective_numbering(p)
        self.assertFalse(numbering["enabled"])
        self.assertIsNone(formatter.automatic_heading_level(p))

    def test_unmapped_style_numbering_is_not_heading_evidence(self):
        doc = Document()
        p = numbered(doc, "Ambiguous list item")
        style = self.custom_style(doc, "Normal")
        style.element.get_or_add_pPr().append(p._p.pPr.find(formatter.qn("w:numPr")))
        p.style = style
        self.assertIsNone(structure.effective_numbering(p)["level"])
        self.assertIsNone(formatter.automatic_heading_level(p))

    def test_heading_style_with_disabled_numbering_adds_no_numpr(self):
        doc = Document()
        doc.add_paragraph("Study title")
        numpr = element("w:numPr")
        numpr.append(element("w:numId", val="0"))
        doc.styles["Heading 1"].element.get_or_add_pPr().append(numpr)
        doc.add_paragraph("Methods")
        output, _ = self.process(doc)
        self.assertIsNone(output.paragraphs[1]._p.pPr.find(formatter.qn("w:numPr")))

    def test_heading_style_cannot_introduce_new_numbering(self):
        doc = Document()
        doc.add_paragraph("Study title")
        p = numbered(doc, "Numbering definition")
        numpr = p._p.pPr.find(formatter.qn("w:numPr"))
        doc.styles["Heading 1"].element.get_or_add_pPr().append(numpr)
        p.text = "Methods"
        output, _ = self.process(doc)
        self.assertIsNone(structure.effective_numbering(output.paragraphs[1]))
        self.assertEqual(formatter.detect_heading_type_from_style(output.paragraphs[1]), "heading_1")

    def test_structural_heading_ends_chinese_abstract(self):
        doc = Document()
        doc.add_paragraph("Study title")
        doc.add_paragraph("\u6458\u8981")
        doc.add_paragraph("Abstract content.")
        doc.add_paragraph("Experimental design", style=self.custom_style(doc))
        _, report = self.process(doc)
        self.assertEqual(report["paragraphs"][3]["detected_type"], "heading_1")

    def test_section_break_belongs_to_previous_section(self):
        doc = Document()
        doc.add_paragraph("First")
        table1 = doc.add_table(rows=1, cols=1)
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        after = doc.add_paragraph("Second")
        table2 = doc.add_table(rows=1, cols=1)
        mapping = structure.section_membership(doc)
        self.assertEqual(mapping[table1._tbl], 0)
        self.assertEqual(mapping[doc.paragraphs[1]._p], 0)
        self.assertEqual(mapping[after._p], 1)
        self.assertEqual(mapping[table2._tbl], 1)

    def test_style_linked_numbering_level_is_not_guessed_as_zero(self):
        doc = Document()
        p = numbered(doc, "1.1.1 Experimental design")
        style = self.custom_style(doc, "Normal")
        numpr = p._p.pPr.find(formatter.qn("w:numPr"))
        numpr.remove(numpr.find(formatter.qn("w:ilvl")))
        style.element.get_or_add_pPr().append(numpr)
        p.style = style
        root = doc.part.numbering_part.element
        definition = next(n for n in root.findall(formatter.qn("w:abstractNum"))
                          if n.get(formatter.qn("w:abstractNumId")) == "80")
        level = definition.find(formatter.qn("w:lvl"))
        level.set(formatter.qn("w:ilvl"), "2")
        level.append(element("w:pStyle", val=style.style_id))
        self.assertEqual(structure.effective_numbering(p)["level"], "2")
        output, _ = self.process(doc)
        self.assertEqual(output.paragraphs[0].style.style_id, style.style_id)
        self.assertEqual(structure.effective_numbering(output.paragraphs[0])["level"], "2")
        self.assertIsNone(output.paragraphs[0]._p.pPr.find(formatter.qn("w:numPr")))

    def test_wrapped_section_is_unknown_instead_of_using_wrong_page(self):
        doc = Document()
        doc.add_paragraph("Study title")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Data"
        doc.add_section(WD_SECTION_START.NEW_PAGE).page_width = Cm(40)
        paragraph = doc.paragraphs[-1]._p
        wrapper, content = element("w:sdt"), element("w:sdtContent")
        paragraph.addprevious(wrapper)
        wrapper.append(content)
        content.append(paragraph)
        before = table._tbl.xml
        margins = doc.sections[0]._sectPr.xml
        self.assertTrue(structure.has_unresolved_sections(doc))
        self.assertIsNone(structure.table_width_budget(doc, table, structure.section_membership(doc)))
        output, report = self.process(doc)
        self.assertEqual(output.tables[0]._tbl.xml, before)
        self.assertEqual(output.sections[0]._sectPr.xml, margins)
        self.assertTrue(report["structure_analysis"]["unresolved_sections"])
        self.assertIsNone(report["tables"][0]["applied_style"])

    def test_preferred_or_cell_width_cannot_hide_behind_narrow_grid(self):
        for source in ("table", "cell"):
            with self.subTest(source=source):
                doc = Document()
                table = doc.add_table(rows=1, cols=1)
                table.cell(0, 0).text = "Data"
                table._tbl.tblGrid.gridCol_lst[0].w = Cm(5)
                width = (table._tbl.tblPr.find(formatter.qn("w:tblW")) if source == "table"
                         else table.cell(0, 0)._tc.tcPr.find(formatter.qn("w:tcW")))
                width.set(formatter.qn("w:type"), "dxa")
                width.set(formatter.qn("w:w"), "15000")
                before = table._tbl.xml
                output, report = self.process(doc)
                self.assertEqual(output.tables[0]._tbl.xml, before)
                self.assertIsNone(report["tables"][0]["applied_style"])

    def test_cross_row_column_width_conflicts_are_protected(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        for column in table._tbl.tblGrid.gridCol_lst:
            column.w = Cm(7)
        for row, widths in zip(table.rows, ((12, 2), (2, 12))):
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
        bound = structure.declared_table_width(table, Cm(16))
        self.assertGreater(bound, Cm(16))
        before = table._tbl.xml
        output, report = self.process(doc)
        self.assertEqual(output.tables[0]._tbl.xml, before)
        self.assertIsNone(report["tables"][0]["applied_style"])

    def test_table_width_uses_its_section_not_largest_paper(self):
        doc = Document()
        doc.add_paragraph("Study title")
        table = doc.add_table(rows=1, cols=1)
        table._tbl.tblGrid.gridCol_lst[0].w = Cm(19)
        table.cell(0, 0).text = "Wide content"
        before = table._tbl.xml
        doc.add_section(WD_SECTION_START.NEW_PAGE).page_width = Cm(40)
        output, report = self.process(doc)
        self.assertEqual(output.tables[0]._tbl.xml, before)
        self.assertIsNone(report["tables"][0]["applied_style"])
        self.assertEqual(report["structure_analysis"]["tables"][0]["section_index"], 0)

    def test_two_columns_reduce_table_width_budget(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        cols = doc.sections[0]._sectPr.find(formatter.qn("w:cols"))
        cols.set(formatter.qn("w:num"), "2")
        section = doc.sections[0]
        budget = structure.table_width_budget(doc, table, structure.section_membership(doc))
        self.assertLess(budget, (section.page_width - section.left_margin - section.right_margin) / 2)

    def test_nested_table_has_no_guessed_page_width(self):
        doc = Document()
        outer = doc.add_table(rows=1, cols=1)
        nested = outer.cell(0, 0).add_table(rows=1, cols=1)
        self.assertIsNone(structure.table_width_budget(doc, nested, structure.section_membership(doc)))
        before = nested._tbl.xml
        output, report = self.process(doc)
        self.assertEqual(output.tables[0].cell(0, 0).tables[0]._tbl.xml, before)
        self.assertIsNone(report["tables"][1]["applied_style"])

    def test_normal_style_and_header_inheritance_are_not_rewritten(self):
        doc = Document()
        doc.styles["Normal"].font.size = Pt(18)
        doc.sections[0].header.paragraphs[0].text = "Header"
        doc.add_paragraph("Study title")
        doc.add_paragraph("Regular body sentence.")
        output, report = self.process(doc)
        self.assertEqual(output.styles["Normal"].font.size.pt, 18)
        self.assertEqual(output.paragraphs[1].runs[0].font.size.pt, 12)
        self.assertEqual(output.sections[0].header.paragraphs[0].text, "Header")
        self.assertEqual(report["verification"]["layout"], "not_performed")

    def test_numbered_table_retains_numbering_through_repeated_processing(self):
        for template in ("default", "course_paper"):
            with self.subTest(template=template):
                doc = Document()
                doc.add_paragraph("Study title")
                p = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
                p.text, p.style = "First experiment", "List Number"
                output, _ = self.process(doc, template)
                second, _ = self.process(output, template)
                self.assertEqual(output.element.xml, second.element.xml)
                self.assertEqual(second.tables[0].cell(0, 0).paragraphs[0].style.name, "List Number")


if __name__ == "__main__":
    unittest.main()
