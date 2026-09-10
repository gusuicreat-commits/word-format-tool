"""Regression for title selection, implicit numbering, and review locations."""
import tempfile
import unittest
from pathlib import Path
from docx import Document
from tests.test_document_corpus import element
import format_docx as f


def numbered(doc, text, fmt="chineseCounting", label="%1、"):
    root = doc.part.numbering_part.element
    abstract = element("w:abstractNum", abstractNumId="80")
    level = element("w:lvl", ilvl="0")
    level.append(element("w:numFmt", val=fmt))
    level.append(element("w:lvlText", val=label))
    abstract.append(level)
    root.append(abstract)
    num = element("w:num", numId="80")
    num.append(element("w:abstractNumId", val="80"))
    root.append(num)
    paragraph = doc.add_paragraph(text)
    numpr = element("w:numPr")
    numpr.append(element("w:ilvl", val="0"))
    numpr.append(element("w:numId", val="80"))
    paragraph._p.get_or_add_pPr().append(numpr)
    return paragraph


class RecognitionReviewTests(unittest.TestCase):
    def test_oversize_table_keeps_original_format_and_reports_location(self):
        doc = Document()
        doc.add_paragraph("研究标题")
        table = doc.add_table(rows=1, cols=2)
        for column in table._tbl.tblGrid.gridCol_lst:
            column.w = f.Cm(20)
        table.cell(0, 0).text = "Wide table content"
        table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Nested content"
        before = table._tbl.xml
        output, report = self.process(doc)
        self.assertEqual(output.tables[0]._tbl.xml, before)
        self.assertIsNone(report["tables"][0]["applied_style"])
        self.assertIsNone(report["tables"][1]["applied_style"])
        self.assertEqual(report["module_status"]["table"]["action"], "skipped")
        warning = next(w for w in report["warnings"] if "table_index" in w)
        self.assertEqual(warning["table_index"], 0)
        self.assertEqual(warning["detected_type"], "table")

    def test_disabled_numbering_survives_caption_cleanup(self):
        doc = Document()
        p = numbered(doc, "图1 系统结构")
        p._p.pPr.find(f.qn("w:numPr")).find(f.qn("w:numId")).set(f.qn("w:val"), "0")
        before = p._p.pPr.find(f.qn("w:numPr")).xml
        output, _ = self.process(doc)
        self.assertEqual(output.paragraphs[0]._p.pPr.find(f.qn("w:numPr")).xml, before)

    def test_split_title_after_repository_cover(self):
        doc = Document()
        doc.add_paragraph("University Open Access Repository")
        doc.add_paragraph("Chapter XI")
        for text in ("THE UK INNOCENCE MOVEMENT:", "PAST, PRESENT, AND FUTURE?"):
            doc.add_paragraph(text).alignment = f.WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("An Author")
        output, report = self.process(doc)
        self.assertEqual([p["index"] for p in report["paragraphs"] if p["detected_type"] == "paper_title"], [2, 3])
        _, second = self.process(output)
        self.assertEqual([p["detected_type"] for p in report["paragraphs"]], [p["detected_type"] for p in second["paragraphs"]])

    def process(self, doc):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory)/"in.docx", Path(directory)/"out.docx"
            doc.save(source)
            rules = f.get_final_format_rules("default", print_warnings=False)
            report = f.init_report(source, output, rules)
            f.format_document(source, output, rules, report)
            return Document(output), report

    def test_title_after_repository_cover(self):
        doc = Document()
        doc.add_paragraph("University repository")
        doc.add_paragraph("Accepted manuscript - copyright retained")
        doc.add_paragraph("Research on document structure", style="Title")
        doc.add_paragraph("Abstract")
        doc.add_paragraph("The study examines document structure.")
        _, report = self.process(doc)
        self.assertEqual([p["index"] for p in report["paragraphs"] if p["detected_type"] == "paper_title"], [2])

    def test_chinese_heading_numbering_is_preserved(self):
        doc = Document()
        doc.add_paragraph("研究标题")
        p = numbered(doc, "系统结构及工作原理")
        doc.add_paragraph("（一）系统总体工作原理")
        before = p._p.pPr.find(f.qn("w:numPr")).xml
        output, report = self.process(doc)
        self.assertEqual(report["paragraphs"][1]["detected_type"], "heading_1")
        self.assertEqual(output.paragraphs[1]._p.pPr.find(f.qn("w:numPr")).xml, before)

    def test_plain_numbered_steps_remain_body(self):
        for fmt, label in (("decimal", "%1."), ("chineseCounting", "%1、")):
            doc = Document()
            doc.add_paragraph("研究标题")
            numbered(doc, "选择输入文件", fmt, label)
            doc.add_paragraph("然后确认输入内容。")
            _, report = self.process(doc)
            self.assertEqual(report["paragraphs"][1]["detected_type"], "body")

    def test_warning_matches_saved_classification(self):
        doc = Document()
        doc.add_paragraph("研究标题")
        doc.add_paragraph("图1为系统的组成，包含三个单元。")
        _, report = self.process(doc)
        warning = next(w for w in report["warnings"] if w["paragraph_index"] == 1)
        self.assertEqual(warning["detected_type"], "body")
        self.assertEqual(warning["applied_style"], "body")
        self.assertTrue(any("一级章节标题" in w["message"] for w in report["warnings"]))


if __name__ == "__main__":
    unittest.main()
