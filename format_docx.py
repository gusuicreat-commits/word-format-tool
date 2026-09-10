"""文格 Lite V2.6.0：本地版 Word 论文格式修改器。

用法：
    python format_docx.py input.docx output.docx --overwrite
    python format_docx.py input.docx output.docx --template default --report reports/report.json --overwrite
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import tempfile
import traceback
from copy import deepcopy
from json import JSONDecodeError
from pathlib import Path
from zipfile import ZipFile

try:
    from document_structure import (
        analyze_document, effective_numbering, heading_evidence,
        section_membership, table_width_budget, declared_table_width,
    )
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from docx.text.paragraph import Paragraph
    from docx.text.run import Run
except ImportError:
    Document = None
    WD_ALIGN_PARAGRAPH = None
    OxmlElement = None
    qn = None
    Cm = None
    Pt = None
    RGBColor = None
    Paragraph = None
    Run = None


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_TEMPLATE_PATH = BASE_DIR / "templates" / "default.json"
VERSION = "V2.6.0"

CHINESE_NUMBER = "一二三四五六七八九十"
PROTECTED_FIRST_PARAGRAPHS = {
    "摘要",
    "摘要：",
    "摘要:",
    "摘 要",
    "关键词",
    "关键词：",
    "关键词:",
    "目录",
    "Abstract",
    "ABSTRACT",
    "Keywords",
    "KEYWORDS",
    "Key words",
    "参考文献",
    "参考文献：",
    "参考文献:",
}
SENTENCE_ENDING_PUNCTUATION = "。？！?!.！"
ALLOWED_ALIGNMENTS = {"left", "center", "right", "justify"}
LATIN_DIGIT_SCOPES = {"global", "body", "abstract", "heading"}
REFERENCE_LATIN_DIGIT_SCOPES = {"reference"}
PAGE_MARGIN_FIELDS = [
    "top_margin_cm",
    "bottom_margin_cm",
    "left_margin_cm",
    "right_margin_cm",
]
TABLE_CAPTION_PATTERN = re.compile(r"^表\s*\d+([-.－—]\d+)?[\.．、\s]*\S+")
FIGURE_CAPTION_PATTERN = re.compile(r"^图\s*\d+([-.－—]\d+)?[\.．、\s]*\S+")

SUPPORTED_STYLE_TYPES = [
    "paper_title",
    "abstract_cn_title",
    "abstract_cn_content",
    "keywords_cn_label",
    "keywords_cn_content",
    "abstract_title",
    "abstract_content",
    "keywords",
    "abstract_en_title",
    "abstract_en_content",
    "keywords_en",
    "keywords_en_label",
    "keywords_en_content",
    "heading_1",
    "heading_2",
    "heading_3",
    "body",
    "table_caption",
    "figure_caption",
    "reference_title",
    "reference_item",
    "table_text",
]

STYLE_FALLBACKS = {
    "abstract_en_title": "abstract_title",
    "abstract_en_content": "abstract_content",
    "keywords_en": "keywords",
}

NUMBERING_CLEAR_PARAGRAPH_TYPES = set(SUPPORTED_STYLE_TYPES) - {
    "body",
    "reference_item",
    "table_text",
    "heading_1",
    "heading_2",
    "heading_3",
}

ALLOWED_STYLE_FIELDS = {
    "font",
    "size_pt",
    "size_cn",
    "color",
    "bold",
    "italic",
    "underline",
    "alignment",
    "line_spacing",
    "first_line_indent_pt",
    "first_line_indent_chars",
    "hanging_indent_chars",
    "space_before_pt",
    "space_after_pt",
    "space_before_lines",
    "space_after_lines",
    "keep_with_next",
    "keep_together",
}

COLOR_MAP = {
    "黑色": "000000",
    "黑": "000000",
    "红色": "FF0000",
    "红": "FF0000",
    "蓝色": "0000FF",
    "蓝": "0000FF",
}
COLOR_PATTERN = re.compile(r"^[0-9A-Fa-f]{6}$")

CHINESE_FONT_SIZE_MAP = {
    "初号": 42,
    "小初": 36,
    "一号": 26,
    "小一": 24,
    "二号": 22,
    "小二": 18,
    "三号": 16,
    "小三": 15,
    "四号": 14,
    "小四": 12,
    "五号": 10.5,
    "小五": 9,
    "六号": 7.5,
    "小六": 6.5,
    "七号": 5.5,
    "八号": 5,
}

ALIGNMENT_MAP = {
    "左对齐": "left",
    "居中": "center",
    "居中对齐": "center",
    "右对齐": "right",
    "两端对齐": "justify",
    "分散对齐": "justify",
}

REPORT_STAT_KEYS = [
    "paper_title",
    "abstract_cn_title",
    "abstract_cn_content",
    "keywords_cn_label",
    "keywords_cn_content",
    "abstract_title",
    "abstract_content",
    "keywords",
    "abstract_en_title",
    "abstract_en_content",
    "keywords_en",
    "keywords_en_label",
    "keywords_en_content",
    "heading_1",
    "heading_2",
    "heading_3",
    "table_caption",
    "figure_caption",
    "body",
    "reference_title",
    "reference_item",
    "table",
    "keywords_cn_label_count",
    "keywords_cn_content_count",
    "keywords_en_label_count",
    "keywords_en_content_count",
    "reference_latin_digit_format_count",
    "reference_hanging_indent_count",
    "skipped_complex_paragraph_count",
    "style_fallback_count",
    "warning_count",
]

MODULE_STATUS_KEYS = [
    "paper_title",
    "abstract_cn",
    "keywords_cn",
    "abstract_en",
    "keywords_en",
    "toc",
    "heading",
    "body",
    "table",
    "figure_caption",
    "table_caption",
    "reference",
    "appendix",
    "page",
    "latin_digit_format",
    "reference_latin_digit_format",
    "header_footer",
    "page_number",
]

MODULE_DEFAULT_NOTES = {
    "paper_title": "未检测到论文标题，未处理",
    "abstract_cn": "未检测到中文摘要，未处理",
    "keywords_cn": "未检测到中文关键词，未处理",
    "abstract_en": "未检测到英文摘要，未处理",
    "keywords_en": "未检测到英文关键词，未处理",
    "toc": "未检测到目录区域，未处理",
    "heading": "未检测到标题层级，未处理",
    "body": "未检测到正文段落，未处理",
    "table": "未检测到表格，未处理",
    "figure_caption": "未检测到图题，未处理",
    "table_caption": "未检测到表题，未处理",
    "reference": "未检测到参考文献，未处理",
    "appendix": "未检测到附录，未处理",
    "latin_digit_format": "老师要求中未出现英文/数字格式要求，未处理",
    "reference_latin_digit_format": "老师要求中未出现参考文献英文/数字格式要求，未处理",
    "header_footer": "老师要求中未出现页眉页脚要求，未处理",
    "page_number": "老师要求中未出现页码要求，未处理",
}

DEFAULT_STYLE = {
    "font": "宋体",
    "size_pt": 12,
    "color": "000000",
    "bold": False,
    "italic": False,
    "underline": False,
    "alignment": "left",
    "line_spacing": 1.5,
    "first_line_indent_pt": 0,
    "space_before_pt": 0,
    "space_after_pt": 0,
}


class UserFacingError(Exception):
    """可以直接展示给普通用户的错误。"""


class TemplateError(UserFacingError):
    """模板读取或校验失败。"""


class ReportError(UserFacingError):
    """报告保存失败。"""


def parse_args(argv=None) -> argparse.Namespace:
    """解析命令行参数。"""
    parser = argparse.ArgumentParser(
        description="文格 Lite V2.6.0：基于规则识别段落，按 JSON 模板修改 .docx，并可生成处理报告。"
    )
    parser.add_argument("input_docx", nargs="?", help="输入 Word 文件路径，例如 samples/input.docx")
    parser.add_argument("output_docx", nargs="?", help="输出 Word 文件路径，例如 samples/output.docx")
    parser.add_argument(
        "--template",
        default=None,
        help="模板名称或路径，例如 default、course_paper 或 templates/default.json",
    )
    parser.add_argument(
        "--report",
        default=None,
        help="可选的 JSON 报告路径，例如 reports/report.json",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="允许覆盖已经存在的输出 .docx 文件",
    )
    parser.add_argument(
        "--debug",
        action="store_true",
        help="出错时显示 Python traceback，方便开发者调试",
    )
    parser.add_argument(
        "--validate-template",
        default=None,
        help="只校验模板 JSON，不处理 Word，例如 templates/default.json",
    )
    parser.add_argument(
        "--override",
        default=None,
        help="可选的局部覆盖规则 JSON，例如 overrides/teacher_override.json",
    )
    parser.add_argument(
        "--export-final-rules",
        default=None,
        help="可选：导出基础模板和覆盖规则合并后的最终完整规则 JSON",
    )
    parser.add_argument(
        "--export-debug-rules",
        default=None,
        help="可选：导出规则调试 JSON，包含原始覆盖、规范化覆盖、最终规则和调试摘要",
    )
    parser.add_argument(
        "--merge-rules",
        action="store_true",
        help="只合并模板和覆盖规则，不处理 Word，通常配合 --export-final-rules 使用",
    )
    return parser.parse_args(argv)


def print_error(message: str) -> None:
    """统一输出中文错误提示。"""
    print(f"错误：{message}", file=sys.stderr)


def ensure_dependency_installed() -> None:
    """检查 python-docx 是否已经安装。"""
    if Document is None:
        raise UserFacingError("缺少依赖 python-docx，请先运行：pip install -r requirements.txt")


def validate_input_file(input_path: Path) -> None:
    """校验输入文件。"""
    if not input_path.exists():
        raise UserFacingError(f"输入文件不存在：{input_path}")

    if not input_path.is_file():
        raise UserFacingError(f"输入路径不是文件：{input_path}")

    suffix = input_path.suffix.lower()
    if suffix == ".doc":
        raise UserFacingError(
            "当前仅支持 .docx 文件，不支持 .doc。请先用 Word/WPS 另存为 .docx 后再试。"
        )

    if suffix != ".docx":
        raise UserFacingError("输入文件必须是 .docx 格式。")

    try:
        if input_path.stat().st_size == 0:
            raise UserFacingError("输入文件为空，无法处理。")
    except OSError as exc:
        raise UserFacingError("无法读取输入文件，请检查文件权限。") from exc

    try:
        with input_path.open("rb") as file:
            file.read(1)
    except OSError as exc:
        raise UserFacingError("无法读取输入文件，请检查文件权限。") from exc


def validate_output_path(
    input_path: Path, output_path: Path, overwrite: bool = False
) -> None:
    """校验输出路径，并在需要时创建输出目录。"""
    if output_path.exists() and output_path.is_dir():
        raise UserFacingError(f"输出路径不能是文件夹：{output_path}")

    if output_path.suffix.lower() != ".docx":
        raise UserFacingError("输出文件名必须以 .docx 结尾。")

    if input_path.resolve() == output_path.resolve():
        raise UserFacingError("为了避免覆盖原文件，输出文件不能和输入文件相同。")

    output_dir = output_path.parent
    if output_dir and not output_dir.exists():
        try:
            output_dir.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            raise UserFacingError(f"无法创建输出目录：{output_dir}。请检查路径和权限。") from exc

    if output_dir and not output_dir.is_dir():
        raise UserFacingError(f"输出目录不是文件夹：{output_dir}")

    if output_path.exists() and not overwrite:
        raise UserFacingError(f"输出文件已存在：{output_path}\n如需覆盖，请添加 --overwrite 参数。")


def validate_report_path(report_path, input_path: Path, output_path: Path) -> None:
    """Reject report aliases before any CLI output is written."""
    if report_path is None:
        return
    path = Path(report_path)
    try:
        for document_path in (input_path, output_path):
            if path.resolve() == document_path.resolve() or (
                path.exists() and document_path.exists() and path.samefile(document_path)
            ):
                raise UserFacingError("报告文件不能与输入或输出 Word 文件相同。")
    except (OSError, RuntimeError) as exc:
        raise UserFacingError("无法确认报告路径安全，请检查路径和文件权限。") from exc


def validate_paths(input_path: Path, output_path: Path, overwrite: bool = False) -> None:
    """兼容旧调用：同时校验输入文件和输出路径。"""
    validate_input_file(input_path)
    validate_output_path(input_path, output_path, overwrite=overwrite)


def resolve_template_path(template_arg: str | None) -> Path:
    """把模板参数解析成具体 JSON 文件路径。"""
    if not template_arg:
        return DEFAULT_TEMPLATE_PATH

    arg_path = Path(template_arg)
    if arg_path.suffix.lower() == ".json":
        return arg_path if arg_path.is_absolute() else BASE_DIR / arg_path

    return BASE_DIR / "templates" / f"{template_arg}.json"


def read_template_json(template_path: Path):
    """读取模板 JSON 文件，并把常见读取错误转换成中文提示。"""
    if not template_path.exists():
        raise TemplateError(f"模板文件不存在：{template_path}")

    if not template_path.is_file():
        raise TemplateError(f"模板路径不是文件：{template_path}")

    try:
        with template_path.open("r", encoding="utf-8") as file:
            template = json.load(file)
    except JSONDecodeError as exc:
        raise TemplateError(
            "模板 JSON 格式不正确，请检查是否缺少逗号、引号或括号。"
            f"文件：{template_path}，第 {exc.lineno} 行第 {exc.colno} 列。"
        ) from exc
    except OSError as exc:
        raise TemplateError(f"读取模板失败：{template_path}，请检查文件权限。") from exc

    return template


def is_number(value) -> bool:
    """bool 在 Python 中是 int 的子类，这里显式排除。"""
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def normalize_font_size(size_cn: str) -> int | float:
    """把中文字号转换成磅值，例如“小四”转换为 12。"""
    if not isinstance(size_cn, str):
        raise TemplateError("中文字号 size_cn 应为字符串，例如“小四”。")

    normalized = size_cn.strip()
    if normalized not in CHINESE_FONT_SIZE_MAP:
        raise TemplateError(f"未知中文字号：{size_cn}。")
    return CHINESE_FONT_SIZE_MAP[normalized]


def normalize_alignment(alignment) -> str:
    """把中文对齐方式转换成内部英文枚举。"""
    if not isinstance(alignment, str):
        raise TemplateError("alignment 应为字符串，例如 center 或 居中。")

    normalized = alignment.strip()
    if normalized in ALIGNMENT_MAP:
        return ALIGNMENT_MAP[normalized]
    if normalized in ALLOWED_ALIGNMENTS:
        return normalized
    raise TemplateError("alignment 只能是 left、center、right、justify，或支持的中文写法。")


def normalize_color(color):
    """把常见中文颜色转换成 6 位十六进制颜色值。"""
    if not isinstance(color, str):
        return color

    normalized = color.strip()
    if normalized in COLOR_MAP:
        return COLOR_MAP[normalized]

    if normalized.startswith("#"):
        normalized = normalized[1:]

    if COLOR_PATTERN.match(normalized):
        return normalized.upper()

    return normalized


def normalize_line_spacing(value) -> int | float:
    """把常见行距写法转换成数字。"""
    if is_number(value):
        return value
    if not isinstance(value, str):
        raise TemplateError("line_spacing 应为数字，或类似“1.5倍”“单倍行距”的字符串。")

    normalized = value.strip()
    spacing_map = {
        "单倍": 1.0,
        "单倍行距": 1.0,
        "双倍": 2.0,
        "双倍行距": 2.0,
    }
    if normalized in spacing_map:
        return spacing_map[normalized]

    number_match = re.match(r"^(\d+(?:\.\d+)?)\s*(?:倍|倍行距)?$", normalized)
    if number_match:
        return float(number_match.group(1))

    raise TemplateError(f"无法识别行距写法：{value}。")


def normalize_indent(value) -> int | float:
    """把首行缩进常见写法转换成磅值，约定 1 个中文字符为 12 磅。"""
    if is_number(value):
        return value
    if not isinstance(value, str):
        raise TemplateError("first_line_indent_pt 应为数字，或类似“2字符”“不缩进”的字符串。")

    normalized = value.strip()
    if normalized in {"无", "不缩进"}:
        return 0
    if normalized == "两个字符":
        return 24

    character_match = re.match(r"^(\d+(?:\.\d+)?)\s*字符$", normalized)
    if character_match:
        return float(character_match.group(1)) * 12

    number_match = re.match(r"^\d+(?:\.\d+)?$", normalized)
    if number_match:
        return float(normalized)

    raise TemplateError(f"无法识别首行缩进写法：{value}。")


def normalize_page_margin(value) -> int | float:
    """把页边距写法转换成厘米数字，例如 3cm 转换为 3。"""
    if is_number(value):
        return value
    if not isinstance(value, str):
        raise TemplateError("页边距应为数字，或类似“3cm”“3厘米”的字符串。")

    normalized = value.strip().lower()
    match = re.match(r"^(\d+(?:\.\d+)?)\s*(?:cm|厘米|公分)?$", normalized)
    if not match:
        raise TemplateError(f"无法识别页边距写法：{value}。")

    margin = float(match.group(1))
    return int(margin) if margin.is_integer() else margin


def normalize_boolean(value, field_name: str) -> bool:
    """把可选中文布尔写法转换成 bool。"""
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        normalized = value.strip()
        if normalized in {"是", "true", "True", "TRUE"}:
            return True
        if normalized in {"否", "false", "False", "FALSE"}:
            return False
    raise TemplateError(f"{field_name} 应为布尔值 true 或 false。")


def normalize_style_indent_fields(style_name: str, normalized_style: dict, warnings: list[str]) -> dict | None:
    """Normalize first-line indent chars and points for one style."""
    if "first_line_indent_chars" not in normalized_style:
        if "first_line_indent_pt" in normalized_style:
            normalized_style["first_line_indent_pt"] = normalize_indent(
                normalized_style["first_line_indent_pt"]
            )
        return None

    indent_chars = normalized_style["first_line_indent_chars"]
    if not is_number(indent_chars):
        raise TemplateError(
            f"styles.{style_name}.first_line_indent_chars 应为数字，例如 2。"
        )

    size_pt = normalized_style.get("size_pt", DEFAULT_STYLE["size_pt"])
    if not is_number(size_pt):
        size_pt = DEFAULT_STYLE["size_pt"]

    converted_pt = float(indent_chars) * float(size_pt)
    existing_pt = None
    status = "generated"
    if "first_line_indent_pt" in normalized_style:
        existing_pt = normalize_indent(normalized_style["first_line_indent_pt"])
        if abs(float(existing_pt) - converted_pt) <= 0.5:
            normalized_style["first_line_indent_pt"] = existing_pt
            status = "equivalent"
        else:
            normalized_style["first_line_indent_pt"] = existing_pt
            status = "mismatch"
            warnings.append(
                f"styles.{style_name}.first_line_indent_chars={indent_chars} "
                f"按 {size_pt} pt 字号约等于 {converted_pt:g} pt，"
                f"与 first_line_indent_pt={existing_pt:g} 不一致，已保留 pt 值。"
            )
    else:
        normalized_style["first_line_indent_pt"] = converted_pt

    return {
        "style": style_name,
        "first_line_indent_chars": float(indent_chars),
        "size_pt": float(size_pt),
        "converted_pt": converted_pt,
        "first_line_indent_pt": normalized_style.get("first_line_indent_pt"),
        "status": status,
        **({"existing_pt": existing_pt} if existing_pt is not None else {}),
    }


def normalize_format_rules(template, fill_defaults=True):
    """把模板中的中文表达和非标准表达转换成内部标准格式。"""
    if not isinstance(template, dict):
        return template

    normalized_template = deepcopy(template)
    warnings = []
    indent_normalization = []

    page = normalized_template.get("page")
    if isinstance(page, dict):
        normalized_page = dict(page)
        for field in PAGE_MARGIN_FIELDS:
            if field in normalized_page:
                normalized_page[field] = normalize_page_margin(normalized_page[field])
        normalized_template["page"] = normalized_page

    latin_digit_format = normalized_template.get("latin_digit_format")
    if isinstance(latin_digit_format, dict):
        normalized_latin_digit_format = dict(latin_digit_format)
        if "size_cn" in normalized_latin_digit_format:
            size_cn_value = normalized_latin_digit_format["size_cn"]
            size_cn_pt = normalize_font_size(size_cn_value)
            if (
                "size_pt" in normalized_latin_digit_format
                and is_number(normalized_latin_digit_format["size_pt"])
            ):
                if (
                    abs(float(normalized_latin_digit_format["size_pt"]) - float(size_cn_pt))
                    > 0.01
                ):
                    warnings.append(
                        f"英文/数字格式 size_cn={size_cn_value} 与 "
                        f"size_pt={normalized_latin_digit_format['size_pt']} 不一致，"
                        "已优先使用 size_pt。"
                    )
            else:
                normalized_latin_digit_format["size_pt"] = size_cn_pt
        normalized_template["latin_digit_format"] = normalized_latin_digit_format

    reference_latin_digit_format = normalized_template.get("reference_latin_digit_format")
    if isinstance(reference_latin_digit_format, dict):
        normalized_reference_latin_digit_format = dict(reference_latin_digit_format)
        if "size_cn" in normalized_reference_latin_digit_format:
            size_cn_value = normalized_reference_latin_digit_format["size_cn"]
            size_cn_pt = normalize_font_size(size_cn_value)
            if (
                "size_pt" in normalized_reference_latin_digit_format
                and is_number(normalized_reference_latin_digit_format["size_pt"])
            ):
                if (
                    abs(
                        float(normalized_reference_latin_digit_format["size_pt"])
                        - float(size_cn_pt)
                    )
                    > 0.01
                ):
                    warnings.append(
                        f"参考文献英文/数字格式 size_cn={size_cn_value} 与 "
                        f"size_pt={normalized_reference_latin_digit_format['size_pt']} 不一致，"
                        "已优先使用 size_pt。"
                    )
            else:
                normalized_reference_latin_digit_format["size_pt"] = size_cn_pt
        normalized_template["reference_latin_digit_format"] = (
            normalized_reference_latin_digit_format
        )

    styles = normalized_template.get("styles")
    if not isinstance(styles, dict):
        normalized_template["_template_warnings"] = warnings
        normalized_template["_indent_normalization"] = indent_normalization
        return normalized_template

    for legacy, specific in (("abstract_title", "abstract_cn_title"), ("abstract_content", "abstract_cn_content")):
        if isinstance(styles.get(specific), dict):
            inherited = styles.get(legacy, {})
            styles[specific] = {**(inherited if isinstance(inherited, dict) else {}), **styles[specific]}

    for style_name, style_config in list(styles.items()):
        if not isinstance(style_config, dict):
            continue

        normalized_style = dict(style_config)

        if "size_cn" in normalized_style:
            size_cn_value = normalized_style["size_cn"]
            size_cn_pt = normalize_font_size(size_cn_value)
            if "size_pt" in normalized_style and is_number(normalized_style["size_pt"]):
                if abs(float(normalized_style["size_pt"]) - float(size_cn_pt)) > 0.01:
                    warnings.append(
                        f"模板样式 {style_name} 中 size_cn={size_cn_value} 与 "
                        f"size_pt={normalized_style['size_pt']} 不一致，已优先使用 size_pt。"
                    )
            elif "size_pt" not in normalized_style:
                normalized_style["size_pt"] = size_cn_pt

        if "alignment" in normalized_style:
            normalized_style["alignment"] = normalize_alignment(
                normalized_style["alignment"]
            )

        if "color" in normalized_style:
            normalized_style["color"] = normalize_color(normalized_style["color"])

        if "line_spacing" in normalized_style:
            normalized_style["line_spacing"] = normalize_line_spacing(
                normalized_style["line_spacing"]
            )

        indent_summary = normalize_style_indent_fields(
            style_name,
            normalized_style,
            warnings,
        )
        if indent_summary is not None:
            indent_normalization.append(indent_summary)

        for field in ["underline", "keep_with_next", "keep_together"]:
            if field in normalized_style:
                normalized_style[field] = normalize_boolean(
                    normalized_style[field],
                    f"styles.{style_name}.{field}",
                )

        if fill_defaults:
            for field, default_value in DEFAULT_STYLE.items():
                if field == "underline":
                    continue
                normalized_style.setdefault(field, default_value)

            normalized_style.setdefault("space_before_pt", DEFAULT_STYLE["space_before_pt"])
            normalized_style.setdefault("space_after_pt", DEFAULT_STYLE["space_after_pt"])
        styles[style_name] = normalized_style

    normalized_template["_template_warnings"] = warnings
    normalized_template["_indent_normalization"] = indent_normalization
    return normalized_template


def validate_format_rules(template):
    """检查规范化后的模板是否合法，返回 errors 和 warnings。"""
    errors = []
    warnings = []

    if not isinstance(template, dict):
        return ["模板内容必须是一个 JSON 对象。"], warnings

    if "version" in template and not isinstance(template["version"], str):
        errors.append("模板字段 version 应为字符串，例如 1.0。")

    if "name" not in template or not isinstance(template.get("name"), str):
        errors.append("模板缺少必要字段 name，且 name 应为字符串。")

    if "description" in template and not isinstance(template["description"], str):
        errors.append("模板字段 description 应为字符串。")

    has_structural_error = False

    if "page" not in template:
        errors.append("模板缺少必要字段 page。")
        has_structural_error = True

    if "styles" not in template:
        errors.append("模板缺少必要字段 styles。")
        has_structural_error = True

    if "page" in template and not isinstance(template["page"], dict):
        errors.append("模板字段 page 应为对象。")
        has_structural_error = True

    if "styles" in template and not isinstance(template["styles"], dict):
        errors.append("模板字段 styles 应为对象。")
        has_structural_error = True

    if "latin_digit_format" in template:
        latin_errors, latin_warnings = validate_latin_digit_format(
            template["latin_digit_format"],
            "模板字段 latin_digit_format",
            LATIN_DIGIT_SCOPES,
        )
        errors.extend(latin_errors)
        warnings.extend(latin_warnings)

    if "reference_latin_digit_format" in template:
        reference_latin_errors, reference_latin_warnings = validate_latin_digit_format(
            template["reference_latin_digit_format"],
            "模板字段 reference_latin_digit_format",
            REFERENCE_LATIN_DIGIT_SCOPES,
        )
        errors.extend(reference_latin_errors)
        warnings.extend(reference_latin_warnings)

    toc_errors = validate_toc_config(template.get("toc"), "模板字段 toc")
    errors.extend(toc_errors)

    unsupported_errors = validate_unsupported_modules(
        template.get("unsupported_modules"),
        "模板字段 unsupported_modules",
    )
    errors.extend(unsupported_errors)

    if has_structural_error:
        return errors, warnings

    if "body" not in template["styles"]:
        errors.append("模板缺少必要字段 styles.body。")

    for field in PAGE_MARGIN_FIELDS:
        if field in template["page"] and not is_number(template["page"][field]):
            errors.append(f"模板字段 page.{field} 应为数字，例如 2.5。")

    for style_name, style_config in template["styles"].items():
        if style_name not in SUPPORTED_STYLE_TYPES:
            warnings.append(
                f"模板包含未知样式类型：{style_name}，将被忽略或保留但不会被规则识别使用。"
            )

        if not isinstance(style_config, dict):
            errors.append(f"模板字段 styles.{style_name} 应为对象。")
            continue

        style_errors, style_warnings = validate_style_config(style_name, style_config)
        errors.extend(style_errors)
        warnings.extend(style_warnings)

    return errors, warnings


def validate_template(template) -> None:
    """兼容旧调用：发现模板错误时直接抛出异常。"""
    errors, _ = validate_format_rules(template)
    if errors:
        raise TemplateError("模板校验失败：\n- " + "\n- ".join(errors))


def validate_override_rules(override):
    """校验局部覆盖规则。override 可以只包含 page 或 styles 的部分字段。"""
    errors = []
    warnings = []

    if not isinstance(override, dict):
        return ["自定义覆盖规则必须是一个 JSON 对象。"], warnings

    if "version" in override and not isinstance(override["version"], str):
        errors.append("自定义覆盖规则字段 version 应为字符串。")

    if "name" in override and not isinstance(override["name"], str):
        errors.append("自定义覆盖规则字段 name 应为字符串。")

    if "description" in override and not isinstance(override["description"], str):
        errors.append("自定义覆盖规则字段 description 应为字符串。")

    has_page = "page" in override
    has_styles = "styles" in override
    has_latin_digit_format = "latin_digit_format" in override
    has_reference_latin_digit_format = "reference_latin_digit_format" in override
    has_toc = "toc" in override
    has_unsupported_modules = "unsupported_modules" in override
    if (
        not has_page
        and not has_styles
        and not has_latin_digit_format
        and not has_reference_latin_digit_format
        and not has_toc
        and not has_unsupported_modules
    ):
        warnings.append(
            "自定义覆盖规则中没有 page、styles、latin_digit_format 或 V0.4 解析字段，不会产生实际覆盖效果。"
        )

    if has_page:
        if not isinstance(override["page"], dict):
            errors.append("自定义覆盖规则字段 page 应为对象。")
        else:
            for field in PAGE_MARGIN_FIELDS:
                if field in override["page"] and not is_number(override["page"][field]):
                    errors.append(f"自定义覆盖规则字段 page.{field} 应为数字，例如 2.5。")

    if has_styles:
        if not isinstance(override["styles"], dict):
            errors.append("自定义覆盖规则字段 styles 应为对象。")
        else:
            for style_name, style_config in override["styles"].items():
                if style_name not in SUPPORTED_STYLE_TYPES:
                    warnings.append(
                        f"自定义覆盖规则包含未知样式类型：{style_name}，将被忽略。"
                    )

                if not isinstance(style_config, dict):
                    errors.append(f"自定义覆盖规则字段 styles.{style_name} 应为对象。")
                    continue

                style_errors, style_warnings = validate_style_config(style_name, style_config)
                errors.extend(
                    error.replace("模板字段", "自定义覆盖规则字段")
                    for error in style_errors
                )
                warnings.extend(
                    warning.replace("模板字段", "自定义覆盖规则字段")
                    for warning in style_warnings
                )

    if has_latin_digit_format:
        latin_errors, latin_warnings = validate_latin_digit_format(
            override["latin_digit_format"],
            "自定义覆盖规则字段 latin_digit_format",
            LATIN_DIGIT_SCOPES,
        )
        errors.extend(latin_errors)
        warnings.extend(latin_warnings)

    if has_reference_latin_digit_format:
        reference_latin_errors, reference_latin_warnings = validate_latin_digit_format(
            override["reference_latin_digit_format"],
            "自定义覆盖规则字段 reference_latin_digit_format",
            REFERENCE_LATIN_DIGIT_SCOPES,
        )
        errors.extend(reference_latin_errors)
        warnings.extend(reference_latin_warnings)

    if has_toc:
        errors.extend(validate_toc_config(override["toc"], "自定义覆盖规则字段 toc"))

    if has_unsupported_modules:
        errors.extend(
            validate_unsupported_modules(
                override["unsupported_modules"],
                "自定义覆盖规则字段 unsupported_modules",
            )
        )

    return errors, warnings


def validate_latin_digit_format(config, field_path: str, allowed_scopes=None):
    """检查英文/数字字符格式配置。"""
    errors = []
    warnings = []
    allowed_fields = {"font", "size_pt", "size_cn", "scope"}
    allowed_scopes = allowed_scopes or LATIN_DIGIT_SCOPES

    if not isinstance(config, dict):
        return [f"{field_path} 应为对象。"], warnings

    for field in config:
        if field not in allowed_fields:
            warnings.append(f"{field_path}.{field} 不是标准字段，将被保留但不会被使用。")

    if "font" in config and not isinstance(config["font"], str):
        errors.append(f"{field_path}.font 应为字符串，例如 \"Times New Roman\"。")
    if "size_pt" in config and not is_number(config["size_pt"]):
        errors.append(f"{field_path}.size_pt 应为数字，例如 12。")
    if "size_cn" in config:
        size_cn = config["size_cn"]
        if not isinstance(size_cn, str):
            errors.append(f"{field_path}.size_cn 应为字符串，例如 \"小四\"。")
        elif size_cn.strip() not in CHINESE_FONT_SIZE_MAP:
            errors.append(f"{field_path}.size_cn 使用了未知中文字号：{size_cn}。")
    if "scope" in config:
        scope = config["scope"]
        if not isinstance(scope, str) or scope not in allowed_scopes:
            allowed_scope_text = "、".join(sorted(allowed_scopes))
            errors.append(f"{field_path}.scope 只能是 {allowed_scope_text}。")

    return errors, warnings


def validate_toc_config(config, field_path: str):
    """检查目录处理策略。"""
    if config is None:
        return []
    errors = []
    if not isinstance(config, dict):
        return [f"{field_path} 应为对象。"]
    for field in config:
        if field != "action":
            errors.append(f"{field_path}.{field} 不是支持的目录字段。")
    action = config.get("action")
    if action is not None and action != "protect":
        errors.append(f"{field_path}.action 只能是 protect。")
    return errors


UNSUPPORTED_MODULE_KEYS = {
    "header_footer",
    "page_number",
    "footnote",
    "endnote",
    "table_three_line",
    "formula",
    "figure_caption",
    "table_caption",
}


def validate_unsupported_modules(config, field_path: str):
    """检查检测到但暂不支持的模块描述。"""
    if config is None:
        return []
    errors = []
    if not isinstance(config, dict):
        return [f"{field_path} 应为对象。"]
    for module_key, module_config in config.items():
        module_path = f"{field_path}.{module_key}"
        if module_key not in UNSUPPORTED_MODULE_KEYS:
            errors.append(f"{module_path} 不是支持的 unsupported module。")
            continue
        if not isinstance(module_config, dict):
            errors.append(f"{module_path} 应为对象。")
            continue
        if module_config.get("status") != "detected_but_not_supported":
            errors.append(f"{module_path}.status 只能是 detected_but_not_supported。")
        if module_config.get("action") != "warning_only":
            errors.append(f"{module_path}.action 只能是 warning_only。")
        note = module_config.get("note")
        if not isinstance(note, str) or not note.strip():
            errors.append(f"{module_path}.note 应为非空字符串。")
        for field in module_config:
            if field not in {"status", "action", "note"}:
                errors.append(f"{module_path}.{field} 不是支持字段。")
    return errors


def validate_style_config(style_name: str, style_config: dict):
    """检查单个样式配置的字段类型。"""
    errors = []
    warnings = []
    field_path = f"styles.{style_name}"

    for field in style_config:
        if field not in ALLOWED_STYLE_FIELDS:
            warnings.append(f"模板字段 {field_path}.{field} 不是标准样式字段，将被保留但不会被使用。")

    if "font" in style_config and not isinstance(style_config["font"], str):
        errors.append(f"模板字段 {field_path}.font 应为字符串，例如 \"宋体\"。")

    if "color" in style_config:
        color = style_config["color"]
        if not isinstance(color, str) or not COLOR_PATTERN.match(color):
            errors.append(
                f"模板字段 {field_path}.color 应为 6 位十六进制颜色值，例如 000000。"
            )

    if "size_pt" in style_config and not is_number(style_config["size_pt"]):
        errors.append(f"模板字段 {field_path}.size_pt 应为数字，例如 12。")

    if "size_cn" in style_config:
        size_cn = style_config["size_cn"]
        if not isinstance(size_cn, str):
            errors.append(f"模板字段 {field_path}.size_cn 应为字符串，例如 \"小四\"。")
        elif size_cn.strip() not in CHINESE_FONT_SIZE_MAP:
            errors.append(f"模板字段 {field_path}.size_cn 使用了未知中文字号：{size_cn}。")

    if "first_line_indent_chars" in style_config and not is_number(
        style_config["first_line_indent_chars"]
    ):
        errors.append(f"模板字段 {field_path}.first_line_indent_chars 应为数字，例如 2。")

    if "bold" in style_config and not isinstance(style_config["bold"], bool):
        errors.append(f"模板字段 {field_path}.bold 应为布尔值 true 或 false。")

    if "italic" in style_config and not isinstance(style_config["italic"], bool):
        errors.append(f"模板字段 {field_path}.italic 应为布尔值 true 或 false。")

    if "underline" in style_config and not isinstance(style_config["underline"], bool):
        errors.append(f"模板字段 {field_path}.underline 应为布尔值 true 或 false。")

    for field in ["keep_with_next", "keep_together"]:
        if field in style_config and not isinstance(style_config[field], bool):
            errors.append(f"模板字段 {field_path}.{field} 应为布尔值 true 或 false。")

    if "alignment" in style_config:
        alignment = style_config["alignment"]
        if not isinstance(alignment, str) or alignment not in ALLOWED_ALIGNMENTS:
            errors.append(
                f"模板字段 {field_path}.alignment 只能是 left、center、right、justify。"
            )

    numeric_fields = [
        "line_spacing",
        "first_line_indent_pt",
        "hanging_indent_chars",
        "space_before_pt",
        "space_after_pt",
        "space_before_lines",
        "space_after_lines",
    ]
    for field in numeric_fields:
        if field in style_config and not is_number(style_config[field]):
            errors.append(f"模板字段 {field_path}.{field} 应为数字。")

    return errors, warnings


def prepare_template(template_arg, print_warnings=True):
    """读取、规范化并校验模板，返回可用于排版的标准模板。"""
    template_path = resolve_template_path(template_arg)
    raw_template = read_template_json(template_path)
    normalized_template = normalize_format_rules(raw_template)
    normalization_warnings = normalized_template.get("_template_warnings", [])
    errors, validation_warnings = validate_format_rules(normalized_template)
    template_warnings = normalization_warnings + validation_warnings

    if errors:
        raise TemplateError("模板校验失败：\n- " + "\n- ".join(errors))

    normalized_template["_template_path"] = str(template_path)
    normalized_template["_template_warnings"] = template_warnings
    normalized_template["_runtime"] = {
        "fallback_count": 0,
        "warned_missing_styles": set(),
    }

    if print_warnings:
        for warning in template_warnings:
            print(f"警告：{warning}")

    return normalized_template


def load_template(template_arg):
    """读取模板配置，并执行规范化与校验。"""
    return prepare_template(template_arg, print_warnings=True)


def resolve_override_path(override_path: str | None) -> Path | None:
    """把 override 参数解析成具体 JSON 路径。"""
    if not override_path:
        return None

    path = Path(override_path)
    return path if path.is_absolute() else BASE_DIR / path


def read_override_json(override_path):
    """读取覆盖规则 JSON，返回原始对象和实际路径。"""
    path = resolve_override_path(override_path)
    if path is None:
        return None, None

    if not path.exists():
        raise UserFacingError(f"自定义覆盖规则文件不存在：{path}")

    if not path.is_file():
        raise UserFacingError(f"自定义覆盖规则路径不是文件：{path}")

    try:
        with path.open("r", encoding="utf-8") as file:
            raw_override = json.load(file)
    except JSONDecodeError as exc:
        raise UserFacingError(
            f"自定义覆盖规则 JSON 格式不正确：{path}，第 {exc.lineno} 行第 {exc.colno} 列。"
        ) from exc
    except OSError as exc:
        raise UserFacingError(f"读取自定义覆盖规则失败：{path}，请检查文件权限。") from exc

    return raw_override, path


def normalize_override_rules(raw_override, path: Path | None = None):
    """规范化并校验局部覆盖规则。"""
    try:
        normalized_override = normalize_format_rules(raw_override, fill_defaults=False)
    except TemplateError as exc:
        raise UserFacingError(f"自定义覆盖规则规范化失败：{exc}") from exc

    normalization_warnings = normalized_override.get("_template_warnings", [])
    errors, validation_warnings = validate_override_rules(normalized_override)
    warnings = normalization_warnings + validation_warnings

    if errors:
        raise UserFacingError("自定义覆盖规则校验失败：\n- " + "\n- ".join(errors))

    normalized_override["_override_path"] = str(path or "")
    normalized_override["_override_warnings"] = warnings
    return normalized_override


def load_override_rules(override_path):
    """读取、规范化并校验局部覆盖规则。"""
    raw_override, path = read_override_json(override_path)
    if raw_override is None:
        return None
    return normalize_override_rules(raw_override, path)


def merge_format_rules(base_rules, override_rules):
    """深度合并基础模板和局部覆盖规则，并记录覆盖字段。"""
    if override_rules is None:
        return deepcopy(base_rules), {
            "enabled": False,
            "overridden_fields": [],
            "warnings": [],
        }

    merged = deepcopy(base_rules)
    overridden_fields = []
    warnings = list(override_rules.get("_override_warnings", []))
    if isinstance(override_rules.get("warnings"), list):
        warnings.extend(
            warning
            for warning in override_rules["warnings"]
            if isinstance(warning, str)
        )

    override_page = override_rules.get("page")
    if isinstance(override_page, dict):
        merged.setdefault("page", {})
        for field, value in override_page.items():
            merged["page"][field] = value
            overridden_fields.append(f"page.{field}")

    override_latin_digit_format = override_rules.get("latin_digit_format")
    if isinstance(override_latin_digit_format, dict):
        merged.setdefault("latin_digit_format", {})
        for field, value in override_latin_digit_format.items():
            if field == "size_cn" or field.startswith("_"):
                continue
            merged["latin_digit_format"][field] = value
            overridden_fields.append(f"latin_digit_format.{field}")

    override_reference_latin_digit_format = override_rules.get(
        "reference_latin_digit_format"
    )
    if isinstance(override_reference_latin_digit_format, dict):
        merged.setdefault("reference_latin_digit_format", {})
        for field, value in override_reference_latin_digit_format.items():
            if field == "size_cn" or field.startswith("_"):
                continue
            merged["reference_latin_digit_format"][field] = value
            overridden_fields.append(f"reference_latin_digit_format.{field}")

    override_toc = override_rules.get("toc")
    if isinstance(override_toc, dict):
        merged["toc"] = {
            key: value
            for key, value in override_toc.items()
            if not str(key).startswith("_")
        }
        overridden_fields.extend(
            f"toc.{field}"
            for field in merged["toc"]
            if not str(field).startswith("_")
        )

    override_unsupported_modules = override_rules.get("unsupported_modules")
    if isinstance(override_unsupported_modules, dict):
        merged["unsupported_modules"] = deepcopy(override_unsupported_modules)
        for module_key, module_config in override_unsupported_modules.items():
            if isinstance(module_config, dict):
                for field in module_config:
                    if not str(field).startswith("_"):
                        overridden_fields.append(
                            f"unsupported_modules.{module_key}.{field}"
                        )

    merged["_indent_normalization"] = [
        *list(base_rules.get("_indent_normalization", [])),
        *list(override_rules.get("_indent_normalization", [])),
    ]

    for metadata_field in ("parser_metadata", "module_requirements"):
        metadata_value = override_rules.get(metadata_field)
        if isinstance(metadata_value, dict):
            merged[metadata_field] = deepcopy(metadata_value)

    override_styles = override_rules.get("styles")
    if isinstance(override_styles, dict):
        merged.setdefault("styles", {})
        for style_name, style_config in override_styles.items():
            if style_name not in SUPPORTED_STYLE_TYPES:
                warnings.append(f"自定义覆盖规则包含未知样式类型：{style_name}，已忽略。")
                continue
            if not isinstance(style_config, dict):
                continue

            merged["styles"].setdefault(style_name, {})
            for field, value in style_config.items():
                if field == "size_cn":
                    continue
                if field.startswith("_"):
                    continue
                merged["styles"][style_name][field] = value
                overridden_fields.append(f"styles.{style_name}.{field}")

    base_name = base_rules.get("name", "base")
    override_name = override_rules.get("name", "override")
    base_description = base_rules.get("description", base_name)
    override_description = override_rules.get("description", override_name)

    merged["name"] = f"{base_name}_with_{override_name}"
    merged["description"] = f"{base_description} + {override_description}"
    merged["version"] = base_rules.get("version", "1.0")

    normalized_merged = normalize_format_rules(merged, fill_defaults=True)
    errors, validation_warnings = validate_format_rules(normalized_merged)
    warnings.extend(normalized_merged.get("_template_warnings", []))
    warnings.extend(validation_warnings)
    warnings = list(dict.fromkeys(warnings))

    if errors:
        raise UserFacingError("最终格式规则校验失败：\n- " + "\n- ".join(errors))

    normalized_merged["_template_path"] = base_rules.get("_template_path")
    normalized_merged["_template_warnings"] = list(
        dict.fromkeys(list(base_rules.get("_template_warnings", [])) + warnings)
    )
    normalized_merged["_runtime"] = {
        "fallback_count": 0,
        "warned_missing_styles": set(),
    }
    normalized_merged["_base_template"] = {
        "name": base_rules.get("name", ""),
        "description": base_rules.get("description", ""),
        "path": base_rules.get("_template_path", ""),
    }
    normalized_merged["_override"] = {
        "enabled": True,
        "path": override_rules.get("_override_path", ""),
        "name": override_rules.get("name", ""),
        "description": override_rules.get("description", ""),
        "overridden_fields": overridden_fields,
        "warnings": warnings,
        "warning_count": len(warnings),
    }
    normalized_merged["_final_rules"] = {
        "name": normalized_merged.get("name", ""),
        "description": normalized_merged.get("description", ""),
    }

    return normalized_merged, normalized_merged["_override"]


def get_final_format_rules(template_arg=None, override_path=None, print_warnings=True):
    """加载基础模板和可选 override，返回最终可执行格式规则。"""
    base_rules = prepare_template(template_arg, print_warnings=print_warnings)
    if not override_path:
        base_rules["_base_template"] = {
            "name": base_rules.get("name", ""),
            "description": base_rules.get("description", ""),
            "path": base_rules.get("_template_path", ""),
        }
        base_rules["_override"] = {"enabled": False}
        base_rules["_final_rules"] = {
            "name": base_rules.get("name", ""),
            "description": base_rules.get("description", ""),
        }
        base_rules["_raw_override"] = None
        base_rules["_normalized_override"] = None
        return base_rules

    raw_override, resolved_override_path = read_override_json(override_path)
    override_rules = normalize_override_rules(raw_override, resolved_override_path)
    final_rules, override_metadata = merge_format_rules(base_rules, override_rules)
    final_rules["_raw_override"] = raw_override
    final_rules["_normalized_override"] = override_rules

    if print_warnings:
        for warning in override_metadata.get("warnings", []):
            print(f"警告：{warning}")

    return final_rules


def clean_format_rules_for_export(rules):
    """移除运行期私有字段，生成可导出的完整规则 JSON。"""
    def clean_value(value):
        if isinstance(value, dict):
            return {
                key: clean_value(item)
                for key, item in value.items()
                if not str(key).startswith("_") and key != "size_cn"
            }
        if isinstance(value, list):
            return [clean_value(item) for item in value]
        if isinstance(value, set):
            return sorted(value)
        return value

    return clean_value(rules)


def export_final_rules(rules, export_path):
    """导出最终完整格式规则 JSON。"""
    path = Path(export_path)
    if path.exists() and path.is_dir():
        raise UserFacingError(f"最终规则导出路径不能是文件夹：{path}")

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("w", encoding="utf-8") as file:
            json.dump(clean_format_rules_for_export(rules), file, ensure_ascii=False, indent=2)
            file.write("\n")
    except OSError as exc:
        raise UserFacingError(f"保存最终规则失败：{path}，请检查目录写入权限。") from exc


def get_nested_value(data, dotted_path: str):
    """按 a.b.c 路径从字典里取值，用于调试摘要。"""
    current = data
    for part in dotted_path.split("."):
        if not isinstance(current, dict) or part not in current:
            return None
        current = current[part]
    return current


def build_debug_summary(rules) -> dict:
    """生成关键规则链路摘要，便于定位覆盖是否进入最终规则。"""
    final_rules = clean_format_rules_for_export(rules)
    override_metadata = rules.get("_override", {"enabled": False})
    overridden_fields = override_metadata.get("overridden_fields", [])
    override_warnings = list(override_metadata.get("warnings", []))
    parser_metadata = rules.get("parser_metadata", {})
    if not isinstance(parser_metadata, dict):
        parser_metadata = {}
    return {
        "override_enabled": bool(override_metadata.get("enabled")),
        "overridden_field_count": len(overridden_fields),
        "overridden_fields": [
            {
                "field": field,
                "final_value": get_nested_value(final_rules, field),
            }
            for field in overridden_fields
        ],
        "page_left_margin_cm": get_nested_value(final_rules, "page.left_margin_cm"),
        "reference_title_alignment": get_nested_value(
            final_rules,
            "styles.reference_title.alignment",
        ),
        "reference_item_alignment": get_nested_value(
            final_rules,
            "styles.reference_item.alignment",
        ),
        "override_warning_count": len(override_warnings),
        "override_warnings": override_warnings,
        "parser_mode": parser_metadata.get("parser_mode"),
        "raw_ai_rules": parser_metadata.get("raw_ai_rules"),
        "validated_rules": parser_metadata.get("validated_rules"),
        "unsupported_modules": final_rules.get("unsupported_modules", {}),
        "conflict_result": parser_metadata.get("conflict_result"),
        "indent_normalization": rules.get("_indent_normalization", []),
    }


def build_rules_debug_payload(rules) -> dict:
    """生成规则调试 JSON，覆盖 raw -> normalized -> final 的主链路。"""
    normalized_override = rules.get("_normalized_override")
    return {
        "base_template": rules.get("_base_template", {}),
        "override": rules.get("_override", {"enabled": False}),
        "raw_override": rules.get("_raw_override"),
        "normalized_override": (
            clean_format_rules_for_export(normalized_override)
            if isinstance(normalized_override, dict)
            else None
        ),
        "final_rules": clean_format_rules_for_export(rules),
        "parser_metadata": rules.get("parser_metadata", {}),
        "debug_summary": build_debug_summary(rules),
    }


def export_debug_rules(rules, export_path):
    """导出规则调试 JSON。"""
    path = Path(export_path)
    if path.exists() and path.is_dir():
        raise UserFacingError(f"规则调试导出路径不能是文件夹：{path}")

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("w", encoding="utf-8") as file:
            json.dump(build_rules_debug_payload(rules), file, ensure_ascii=False, indent=2)
            file.write("\n")
    except OSError as exc:
        raise UserFacingError(f"保存规则调试信息失败：{path}，请检查目录写入权限。") from exc


def count_override_fields(rules) -> int:
    """返回 override 实际覆盖字段数量。"""
    return len(rules.get("_override", {}).get("overridden_fields", []))


def init_module_status() -> dict:
    """初始化所有模块状态，确保未检测到的模块也出现在报告中。"""
    status = {}
    for module_key in MODULE_STATUS_KEYS:
        if module_key in {
            "header_footer",
            "page_number",
            "latin_digit_format",
            "reference_latin_digit_format",
        }:
            module_status = "not_requested"
        else:
            module_status = "not_detected"
        status[module_key] = {
            "status": module_status,
            "count": 0,
            "action": "skipped",
            "note": MODULE_DEFAULT_NOTES.get(module_key, "未检测到，未处理"),
        }
    return status


def init_report(input_path, output_path, template):
    """初始化格式处理报告对象。"""
    template_warnings = list(template.get("_template_warnings", []))
    override_metadata = template.get("_override", {"enabled": False})
    return {
        "version": VERSION,
        "report_schema_version": 1,
        "input_file": str(input_path),
        "output_file": str(output_path),
        "template": {
            "name": template.get("name", "未命名模板"),
            "description": template.get("description", "无"),
        },
        "base_template": template.get(
            "_base_template",
            {
                "name": template.get("name", "未命名模板"),
                "description": template.get("description", "无"),
                "path": template.get("_template_path", ""),
            },
        ),
        "override": override_metadata,
        "final_rules": template.get(
            "_final_rules",
            {
                "name": template.get("name", "未命名模板"),
                "description": template.get("description", "无"),
            },
        ),
        "final_rules_debug": build_rules_debug_payload(template),
        "stats": {key: 0 for key in REPORT_STAT_KEYS},
        "module_status": init_module_status(),
        "paragraphs": [],
        "tables": [],
        "template_warnings": template_warnings,
        "warnings": [
            {
                "paragraph_index": None,
                "text_preview": "",
                "message": warning,
            }
            for warning in template_warnings
        ],
    }


def make_text_preview(text, max_len=50):
    """生成报告中的文本预览，最多保留 max_len 个字符。"""
    clean_text = (text or "").strip()
    if len(clean_text) <= max_len:
        return clean_text
    return clean_text[:max_len] + "..."


def add_paragraph_report(
    report, index, text, detected_type, applied_style, warning=None, pagination=None
):
    """添加段落识别和样式应用记录。"""
    item = {
        "index": index,
        "text_preview": make_text_preview(text),
        "detected_type": detected_type,
        "applied_style": applied_style,
        "warning": warning,
    }
    if pagination is not None:
        item["pagination"] = pagination
    report["paragraphs"].append(item)


def add_table_report(
    report, table_index, row_count, cell_count, applied_style, warning=None
):
    """添加表格处理记录。"""
    item = {
        "table_index": table_index,
        "row_count": row_count,
        "cell_count": cell_count,
        "applied_style": applied_style,
    }
    if warning:
        item["warning"] = warning
    report["tables"].append(item)


def add_warning(report, paragraph_index, text, message):
    """添加警告信息。"""
    report["warnings"].append(
        {
            "paragraph_index": paragraph_index,
            "text_preview": make_text_preview(text),
            "message": message,
        }
    )


def set_module_status(
    module_status: dict,
    module_key: str,
    status: str,
    count: int,
    action: str,
    note: str,
) -> None:
    """写入单个模块状态。"""
    if module_key not in module_status:
        module_status[module_key] = {}
    module_status[module_key].update(
        {
            "status": status,
            "count": count,
            "action": action,
            "note": note,
        }
    )


def increment_module_count(context: dict, module_key: str, amount: int = 1) -> None:
    """累计模块命中次数。"""
    counts = context.setdefault("module_counts", {})
    counts[module_key] = counts.get(module_key, 0) + amount


def mark_module_flag(context: dict, flag_name: str) -> None:
    """记录模块级布尔状态，例如目录保护或附录边界判断。"""
    context.setdefault("module_flags", {})[flag_name] = True


def get_module_requirement_flags(template) -> dict:
    """从解析出的 override 元数据中读取高级模块要求标记。"""
    flags = {}
    candidates = [
        template.get("_normalized_override"),
        template.get("_raw_override"),
        template,
    ]
    for candidate in candidates:
        if not isinstance(candidate, dict):
            continue
        module_requirements = candidate.get("module_requirements") or candidate.get(
            "_module_requirements"
        )
        if isinstance(module_requirements, dict):
            for key in ("header_footer", "page_number"):
                if isinstance(module_requirements.get(key), bool):
                    flags[key] = module_requirements[key]
        unsupported_modules = candidate.get("unsupported_modules")
        if isinstance(unsupported_modules, dict):
            for key in ("header_footer", "page_number"):
                module_config = unsupported_modules.get(key)
                if (
                    isinstance(module_config, dict)
                    and module_config.get("status") == "detected_but_not_supported"
                ):
                    flags[key] = True
    return flags


def finalize_module_status(report, context, template) -> dict:
    """根据本次检测结果生成保守式模块触发状态报告。"""
    module_status = init_module_status()
    stats = context.get("stats", {})
    counts = context.get("module_counts", {})
    flags = context.get("module_flags", {})

    paper_title_count = stats.get("paper_title", 0)
    if paper_title_count:
        set_module_status(
            module_status,
            "paper_title",
            "detected",
            paper_title_count,
            "formatted",
            "检测到论文标题并应用标题样式",
        )

    abstract_cn_count = stats.get("abstract_title", 0) + stats.get("abstract_content", 0)
    if abstract_cn_count:
        set_module_status(
            module_status,
            "abstract_cn",
            "detected",
            abstract_cn_count,
            "formatted",
            "检测到中文摘要并应用摘要样式",
        )

    keywords_cn_count = stats.get("keywords", 0)
    if keywords_cn_count:
        keywords_cn_label_count = counts.get("keywords_cn_label_count", 0)
        keywords_cn_content_count = counts.get("keywords_cn_content_count", 0)
        keywords_cn_note = "检测到中文关键词并应用关键词样式"
        if keywords_cn_label_count or keywords_cn_content_count:
            keywords_cn_note = (
                f"检测到中文关键词，并已分别格式化 label {keywords_cn_label_count} 处、"
                f"content {keywords_cn_content_count} 处"
            )
        set_module_status(
            module_status,
            "keywords_cn",
            "detected",
            keywords_cn_count,
            "formatted",
            keywords_cn_note,
        )

    abstract_en_formatted_count = (
        stats.get("abstract_en_title", 0) + stats.get("abstract_en_content", 0)
    )
    abstract_en_count = max(counts.get("abstract_en", 0), abstract_en_formatted_count)
    if abstract_en_count:
        abstract_en_action = "formatted" if abstract_en_formatted_count else "warning_only"
        abstract_en_note = "检测到英文摘要并应用英文摘要样式"
        if not abstract_en_formatted_count:
            abstract_en_note = "检测到英文摘要，但未执行单独格式化"
        set_module_status(
            module_status,
            "abstract_en",
            "detected",
            abstract_en_count,
            abstract_en_action,
            abstract_en_note,
        )

    keywords_en_count = max(counts.get("keywords_en", 0), stats.get("keywords_en", 0))
    if keywords_en_count:
        keywords_en_label_count = counts.get("keywords_en_label_count", 0)
        keywords_en_content_count = counts.get("keywords_en_content_count", 0)
        keywords_en_action = "formatted" if stats.get("keywords_en", 0) else "warning_only"
        keywords_en_note = "检测到英文关键词并应用英文关键词样式"
        if keywords_en_label_count or keywords_en_content_count:
            keywords_en_action = "formatted"
            keywords_en_note = (
                f"检测到英文关键词，并已分别格式化 label {keywords_en_label_count} 处、"
                f"content {keywords_en_content_count} 处"
            )
        elif not stats.get("keywords_en", 0):
            keywords_en_note = "检测到英文关键词，但未执行单独格式化"
        set_module_status(
            module_status,
            "keywords_en",
            "detected",
            keywords_en_count,
            keywords_en_action,
            keywords_en_note,
        )

    toc_count = counts.get("toc", 0)
    if toc_count:
        set_module_status(
            module_status,
            "toc",
            "detected",
            toc_count,
            "protected",
            "检测到目录区域，当前版本不生成或更新目录，仅保护不强行改写",
        )

    heading_count = (
        stats.get("heading_1", 0)
        + stats.get("heading_2", 0)
        + stats.get("heading_3", 0)
    )
    heading_count = max(heading_count, counts.get("heading", 0))
    if heading_count:
        set_module_status(
            module_status,
            "heading",
            "detected",
            heading_count,
            "formatted",
            "检测到标题层级并应用标题样式",
        )

    body_count = stats.get("body", 0)
    if body_count:
        set_module_status(
            module_status,
            "body",
            "detected",
            body_count,
            "formatted",
            "检测到正文段落并应用正文样式",
        )

    table_count = stats.get("table", 0)
    if table_count:
        skipped_tables = sum(item.get("applied_style") is None for item in report["tables"])
        set_module_status(
            module_status,
            "table",
            "detected",
            table_count,
            "skipped" if skipped_tables == table_count else "formatted",
            f"检测到 {table_count} 个表格，应用格式 {table_count - skipped_tables} 个，保留原格式 {skipped_tables} 个。",
        )

    figure_caption_count = stats.get("figure_caption", 0)
    if figure_caption_count:
        set_module_status(
            module_status,
            "figure_caption",
            "detected",
            figure_caption_count,
            "formatted",
            "检测到图题并应用图题样式",
        )

    table_caption_count = stats.get("table_caption", 0)
    if table_caption_count:
        set_module_status(
            module_status,
            "table_caption",
            "detected",
            table_caption_count,
            "formatted",
            "检测到表题并应用表题样式",
        )

    reference_count = stats.get("reference_title", 0) + stats.get("reference_item", 0)
    if reference_count:
        reference_latin_count = counts.get("reference_latin_digit_format", 0)
        reference_hanging_count = counts.get("reference_hanging_indent", 0)
        reference_note = "检测到参考文献标题或条目并应用参考文献样式"
        reference_details = []
        if reference_latin_count:
            reference_details.append(f"参考文献英文/数字已处理 {reference_latin_count} 处")
        if reference_hanging_count:
            reference_details.append(f"悬挂缩进已应用 {reference_hanging_count} 段")
        if reference_details:
            reference_note = "检测到参考文献并应用参考文献样式；" + "；".join(reference_details)
        set_module_status(
            module_status,
            "reference",
            "detected",
            reference_count,
            "formatted",
            reference_note,
        )

    appendix_count = counts.get("appendix", 0)
    if appendix_count:
        appendix_note = "检测到附录，并用于退出参考文献模式"
        if not flags.get("appendix_boundary"):
            appendix_note = "检测到附录，当前版本仅用于边界判断，不单独格式化"
        set_module_status(
            module_status,
            "appendix",
            "detected",
            appendix_count,
            "boundary_only",
            appendix_note,
        )

    page_field_count = len(template.get("page", {})) if isinstance(template.get("page"), dict) else 0
    if page_field_count:
        set_module_status(
            module_status,
            "page",
            "detected",
            page_field_count,
            "formatted",
            "已应用页边距等页面基础设置",
        )

    latin_digit_config = get_latin_digit_format(template)
    if latin_digit_config is not None:
        latin_digit_count = counts.get("latin_digit_format", 0)
        latin_digit_skipped_count = counts.get("latin_digit_format_skipped", 0)
        latin_digit_font = latin_digit_config.get("font", "当前字符字体")
        if latin_digit_count:
            latin_note = (
                f"检测到英文/数字格式要求，并已应用 {latin_digit_font}。"
            )
            if latin_digit_skipped_count:
                latin_note += f" 另有 {latin_digit_skipped_count} 个复杂片段保守跳过。"
            set_module_status(
                module_status,
                "latin_digit_format",
                "detected",
                latin_digit_count,
                "formatted",
                latin_note,
            )
        else:
            latin_note = "检测到英文/数字格式要求，但未找到可安全处理的英文或数字文本。"
            if latin_digit_skipped_count:
                latin_note = (
                    "检测到英文/数字格式要求，但目标位于复杂结构中，已保守跳过。"
                )
            set_module_status(
                module_status,
                "latin_digit_format",
                "detected",
                0,
                "skipped",
                latin_note,
            )

    reference_latin_digit_config = get_reference_latin_digit_format(template)
    if reference_latin_digit_config is not None:
        reference_latin_count = counts.get("reference_latin_digit_format", 0)
        reference_latin_skipped_count = counts.get(
            "reference_latin_digit_format_skipped",
            0,
        )
        reference_latin_font = reference_latin_digit_config.get("font", "当前字符字体")
        if reference_latin_count:
            reference_latin_note = (
                f"检测到参考文献英文/数字格式要求，并已应用 {reference_latin_font}。"
            )
            if reference_latin_skipped_count:
                reference_latin_note += (
                    f" 另有 {reference_latin_skipped_count} 个复杂片段保守跳过。"
                )
            set_module_status(
                module_status,
                "reference_latin_digit_format",
                "detected",
                reference_latin_count,
                "formatted",
                reference_latin_note,
            )
        else:
            reference_latin_note = (
                "检测到参考文献英文/数字格式要求，但未找到可安全处理的参考文献英文或数字文本。"
            )
            if reference_latin_skipped_count:
                reference_latin_note = (
                    "检测到参考文献英文/数字格式要求，但目标位于复杂结构中，已保守跳过。"
                )
            set_module_status(
                module_status,
                "reference_latin_digit_format",
                "detected",
                0,
                "skipped",
                reference_latin_note,
            )

    requirement_flags = get_module_requirement_flags(template)
    if requirement_flags.get("header_footer"):
        set_module_status(
            module_status,
            "header_footer",
            "detected_but_not_supported",
            1,
            "warning_only",
            "检测到页眉页脚相关要求，但当前版本暂不处理，请人工确认",
        )
        add_warning(
            report,
            None,
            "",
            "检测到页眉页脚相关要求，但当前版本暂不处理，请人工确认",
        )

    if requirement_flags.get("page_number"):
        set_module_status(
            module_status,
            "page_number",
            "detected_but_not_supported",
            1,
            "warning_only",
            "检测到页码相关要求，但当前版本暂不处理，请人工确认",
        )
        add_warning(
            report,
            None,
            "",
            "检测到页码相关要求，但当前版本暂不处理，请人工确认",
        )

    report["module_status"] = module_status
    debug_summary = report.setdefault("final_rules_debug", {}).setdefault(
        "debug_summary", {}
    )
    execution_counts = dict(counts)
    debug_summary["execution_counts"] = execution_counts
    debug_summary["keyword_label_content"] = {
        "keywords_cn_label_count": counts.get("keywords_cn_label_count", 0),
        "keywords_cn_content_count": counts.get("keywords_cn_content_count", 0),
        "keywords_en_label_count": counts.get("keywords_en_label_count", 0),
        "keywords_en_content_count": counts.get("keywords_en_content_count", 0),
        "skipped_complex_paragraph_count": counts.get("skipped_complex_paragraph_count", 0),
    }
    debug_summary["reference_execution"] = {
        "reference_latin_digit_format_count": counts.get(
            "reference_latin_digit_format",
            0,
        ),
        "reference_hanging_indent_count": counts.get("reference_hanging_indent", 0),
        "reference_latin_digit_format_skipped": counts.get(
            "reference_latin_digit_format_skipped",
            0,
        ),
    }
    debug_summary["indent_normalization"] = template.get("_indent_normalization", [])
    debug_summary["unsupported_modules"] = clean_format_rules_for_export(template).get(
        "unsupported_modules",
        {},
    )
    debug_summary["heading_style_conflicts"] = context.get(
        "heading_style_conflicts",
        [],
    )
    debug_summary["heading_numbering_system"] = context.get(
        "heading_numbering_system",
        {},
    )
    debug_summary["arabic_heading_level1_skipped"] = context.get(
        "arabic_heading_level1_skipped",
        [],
    )
    debug_summary["heading_style_sync"] = context.get("heading_style_sync", [])
    debug_summary["heading_style_sync_failures"] = context.get(
        "heading_style_sync_failures",
        [],
    )
    debug_summary["module_status"] = module_status
    debug_summary["module_status_summary"] = {
        "detected_count": sum(
            1
            for item in module_status.values()
            if item.get("status") in {"detected", "detected_but_not_supported"}
        ),
        "warning_only_count": sum(
            1 for item in module_status.values() if item.get("action") == "warning_only"
        ),
    }
    return module_status


def save_report(report, report_path):
    """保存 JSON 报告文件。"""
    path = Path(report_path)
    if path.exists() and path.is_dir():
        raise ReportError(f"报告路径不能是文件夹：{path}")

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("w", encoding="utf-8") as file:
            json.dump(report, file, ensure_ascii=False, indent=2)
            file.write("\n")
    except OSError as exc:
        raise ReportError(f"保存报告失败：{path}，请检查目录写入权限。") from exc


def is_explicit_keywords(text: str) -> bool:
    """判断是否是关键词段落。"""
    normalized = text.strip()
    return (
        normalized == "关键词"
        or CHINESE_KEYWORDS_LABEL_PATTERN.match(normalized) is not None
    )


def is_chinese_abstract_title(text: str) -> bool:
    """判断是否是中文摘要标题。"""
    return re.fullmatch(r"(?:【\s*摘\s*要\s*】|摘\s*要)\s*[:：]?", text.strip()) is not None


def is_english_abstract(text: str) -> bool:
    """保守识别英文摘要标题或摘要起始段。"""
    normalized = re.sub(r"\s+", " ", text.strip())
    return re.match(r"^abstract\s*[:：]?(?:\s|$)", normalized, re.IGNORECASE) is not None


def is_english_keywords(text: str) -> bool:
    """保守识别英文关键词标题或关键词起始段。"""
    normalized = re.sub(r"\s+", " ", text.strip())
    return re.match(
        r"^(keywords|key\s+words)\s*[:：]?(?:\s|$)",
        normalized,
        re.IGNORECASE,
    ) is not None


def is_toc_title(text: str) -> bool:
    """保守识别目录标题。"""
    compact = re.sub(r"\s+", "", text.strip())
    return compact == "目录"


def is_toc_entry(text: str) -> bool:
    """保守识别目录条目，不生成或更新目录。"""
    normalized = re.sub(r"\s+", " ", text.strip())
    if not normalized:
        return False
    if re.search(r"\.{2,}\s*\d+$", normalized):
        return True
    if re.search(r"…+\s*\d+$", normalized):
        return True
    if re.search(r"\s{2,}\d+$", normalized):
        return True
    if re.match(rf"^[{CHINESE_NUMBER}]+[\u3001.\uFF0E]\s*\S.+\s+\d{{1,4}}$", normalized):
        return True
    if re.match(r"^\d+(?:\.\d+)+\s+\S.+\s+\d{1,4}$", normalized):
        return True
    return False


def is_reference_title(text: str) -> bool:
    """判断是否是参考文献标题。"""
    return text.strip().rstrip(":：").strip().casefold() in {"参考文献", "references", "bibliography"}


def is_reference_exit_title(text: str) -> bool:
    """判断参考文献区域后续章节标题，遇到后退出参考文献模式。"""
    normalized = text.strip()
    if re.match(r"^(附录(?:[一二三四五六七八九十\dA-Za-z])?|附件|致谢)(?:[：:\s].*)?$", normalized):
        return True
    return re.match(
        r"^(appendix|acknowledgements?|acknowledgments?)(?:[：:\s].*)?$",
        normalized,
        re.IGNORECASE,
    ) is not None


def is_title_candidate(text: str) -> bool:
    """标题保护规则：过长或明显句子结尾的段落，不当作标题。"""
    if len(text) > 40:
        return False
    if text.endswith(tuple(SENTENCE_ENDING_PUNCTUATION)):
        return False
    return True


def is_paper_title_candidate(text, paragraph) -> bool:
    if is_title_candidate(text):
        return True
    if paragraph is None or not 40 < len(text) <= 240:
        return False
    if re.search(r"https?://|@|_{3,}|\b(?:repository|correspondence|author affiliations)\b", text, re.I):
        return False
    runs = [run for run in paragraph.runs if run.text.strip()]
    # Long first-paragraph titles need explicit typographic evidence.
    return bool(runs) and all(run.bold is True for run in runs) and not text.endswith(("。", ";", "；"))


def detect_named_heading_type(text):
    normalized = re.sub(r"\s+", " ", text.strip()).rstrip(":：").casefold()
    level_one = {"introduction", "background", "method", "methods", "materials and methods",
                 "results", "discussion", "results and discussion", "conclusion", "conclusions",
                 "recommendations", "declarations", "full methods", "引言", "绪论", "结论", "总结"}
    level_two = {"participants", "materials", "procedure", "procedures", "statistical analysis"}
    if normalized in level_one:
        return "heading_1"
    if normalized in level_two:
        return "heading_2"
    return None


def automatic_heading_level(paragraph):
    """Resolve Chinese chapter numbering without modifying numbering definitions."""
    if paragraph is None:
        return None
    numbering = effective_numbering(paragraph)
    if not numbering or not numbering["enabled"]:
        return None
    level = numbering["level"]
    if level is None:
        return None
    try:
        root = paragraph.part.numbering_part.element
    except (KeyError, NotImplementedError):
        return None
    nums = root.xpath('./w:num')
    num = next((n for n in nums if n.get(qn("w:numId")) == numbering["num_id"]), None)
    if num is None:
        return None
    abstract = num.find(qn("w:abstractNumId"))
    if abstract is None:
        return None
    definitions = root.xpath('./w:abstractNum')
    definition = next((n for n in definitions if n.get(qn("w:abstractNumId")) == abstract.get(qn("w:val"))), None)
    if definition is None:
        return None
    levels = [n for n in definition.findall(qn("w:lvl")) if n.get(qn("w:ilvl")) == level]
    for override in num.findall(qn("w:lvlOverride")):
        if override.get(qn("w:ilvl")) == level and override.find(qn("w:lvl")) is not None:
            levels = [override.find(qn("w:lvl"))]
    if not levels:
        return None
    fmt = levels[0].find(qn("w:numFmt"))
    label = levels[0].find(qn("w:lvlText"))
    if fmt is None or label is None or fmt.get(qn("w:val")) not in {"chineseCounting", "chineseCountingThousand", "ideographTraditional"}:
        return None
    pattern = label.get(qn("w:val"), "")
    if re.fullmatch(r"%\d[、．.]", pattern) or re.fullmatch(r"第%\d章", pattern):
        return "heading_1"
    if re.fullmatch(r"[（(]%\d[）)]", pattern):
        return "heading_2"
    return None


def find_document_title(paragraphs, protected):
    """Select an explicit title before the abstract, including after a cover sheet."""
    candidates = []
    first_content = next((i for i, p in enumerate(paragraphs) if p.text.strip() and i not in protected), None)
    repository_cover = False
    chapter_marker = False
    for index, paragraph in enumerate(paragraphs[:80]):
        text = paragraph.text.strip()
        if not text or index in protected:
            continue
        repository_cover = repository_cover or bool(re.search(r"open access repository", text, re.I))
        if repository_cover and re.fullmatch(r"Chapter\s+[IVXLCDM]+", text, re.I):
            chapter_marker = True
            continue
        if chapter_marker and paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER and text.isupper() and 12 <= len(text) <= 160:
            return index
        if is_chinese_abstract_title(text) or is_english_abstract(text) or re.match(r"^(?:【摘要】|摘要[:：])", text):
            break
        if re.search(r"repository|accepted manuscript|copyright|https?://|学号|指导教师|毕业论文|学位论文|独创性声明", text, re.I):
            continue
        style = get_paragraph_style_name(paragraph).casefold()
        if style in {"title", "标题"} and is_paper_title_candidate(text, paragraph):
            return index
        if index == first_content and detect_heading_type_from_style(paragraph) == "heading_1" and is_paper_title_candidate(text, paragraph) and not detect_text_heading_type(text, paragraph=paragraph):
            candidates.append(index)
    return candidates[0] if len(candidates) == 1 else None


def is_bold_numbered_heading(text, paragraph):
    if paragraph is None or len(text) > 160 or text.endswith((".", "。", ";", "；")):
        return False
    if not re.match(r"^\d+[.．]\s+[A-Za-z]", text):
        return False
    runs = [run for run in paragraph.runs if run.text.strip()]
    return bool(runs) and all(run.bold is True for run in runs)


def is_chapter_heading(text: str) -> bool:
    """Return True for explicit chapter headings like 第1章 or 第一章."""
    normalized = text.strip()
    if not is_title_candidate(normalized):
        return False
    return (
        re.match(rf"^第(?:\d+|[{CHINESE_NUMBER}]+)章(?!为)\s*\S+", normalized)
        is not None
    )


def is_arabic_level1_heading_candidate(text: str) -> bool:
    """Return True for single-level Arabic numbered heading candidates."""
    normalized = text.strip()
    return re.match(r"^\d+(?:[、.．]|\s+)\s*\S+", normalized) is not None


def get_arabic_level1_heading_number(text: str) -> str | None:
    """Return the leading Arabic level-1 number, if present."""
    match = re.match(r"^\s*(\d+)(?:[、.．]|\s+)\s*\S+", text.strip())
    return match.group(1) if match else None


def get_numeric_subheading_root_number(text: str) -> str | None:
    """Return the root number for 1.1 / 1.1.1 style headings."""
    match = re.match(r"^\s*(\d+)\.\d+(?:\.\d+)*\s*\S+", text.strip())
    return match.group(1) if match else None


def is_strong_arabic_level1_heading_text(text: str) -> bool:
    """Conservatively allow common short chapter titles without a full hierarchy."""
    normalized = text.strip()
    match = re.match(r"^\d+(?:[、.．]|\s+)\s*(\S.*)$", normalized)
    if not match:
        return False
    title_text = match.group(1).strip()
    if len(title_text) > 16:
        return False
    return bool(
        re.search(
            r"(绪论|引言|概述|理论|现状|方法|内容|结论|总结|展望|参考文献)",
            title_text,
        )
    )


def detect_number_heading_type(text: str) -> str | None:
    """识别 1、1.1、1.1.1 这类数字编号标题。"""
    normalized = text.strip()
    if is_chapter_heading(normalized):
        return "heading_1"
    if re.match(r"^\d+\.\d+\.\d+(?:\.\d+)*\s*\S+", normalized):
        return "heading_3"
    if re.match(r"^\d+\.\d+(?!\.)\s*\S+", normalized):
        return "heading_2"

    return None


def record_arabic_heading_level1_skipped(
    context: dict | None,
    text: str,
    reason: str,
) -> None:
    """Record a cautious non-classification for Arabic single-level numbering."""
    if context is None:
        return
    skipped = context.setdefault("arabic_heading_level1_skipped", [])
    skipped.append(
        {
            "text_preview": make_text_preview(text),
            "reason": reason,
        }
    )


def detect_text_heading_type(text: str, context: dict | None = None, paragraph=None) -> str | None:
    """Use explicit paragraph numbering as the highest-priority heading signal."""
    normalized = text.strip()
    if re.match(rf"^[{CHINESE_NUMBER}]+[、.．]\s*\S+", normalized):
        return "heading_1"
    if re.match(rf"^[（(][{CHINESE_NUMBER}]+[）)]\s*\S*", normalized):
        return "heading_2"
    number_heading_type = detect_number_heading_type(normalized)
    if number_heading_type:
        return number_heading_type
    if is_arabic_level1_heading_candidate(normalized):
        number = get_arabic_level1_heading_number(normalized)
        allowed_numbers = set((context or {}).get("arabic_heading_level1_numbers", []))
        seen_numbers = (context or {}).setdefault(
            "seen_arabic_heading_level1_numbers",
            set(),
        ) if context is not None else set()
        if (
            detect_heading_type_from_style(paragraph)
            or (number in allowed_numbers and number not in seen_numbers)
            or is_strong_arabic_level1_heading_text(normalized)
            or is_bold_numbered_heading(normalized, paragraph)
        ):
            if context is not None and number is not None:
                seen_numbers.add(number)
            return "heading_1"
        record_arabic_heading_level1_skipped(
            context,
            normalized,
            "single_level_arabic_number_without_heading_context",
        )
    return detect_named_heading_type(normalized)


def record_heading_style_conflict(
    context: dict,
    paragraph,
    index: int,
    text: str,
    text_pattern_level: str | None,
) -> None:
    """Record a Word-style/text-numbering heading conflict for report/debug."""
    if text_pattern_level is None:
        return
    original_style_level = detect_heading_type_from_style(paragraph)
    if original_style_level is None or original_style_level == text_pattern_level:
        return

    conflict = {
        "paragraph_index": index,
        "text_preview": make_text_preview(text),
        "original_style": get_paragraph_style_name(paragraph),
        "original_style_id": get_paragraph_style_id(paragraph),
        "original_style_level": original_style_level,
        "text_pattern_level": text_pattern_level,
        "final_level": text_pattern_level,
        "reason": "text_pattern_priority",
    }
    context.setdefault("heading_style_conflicts", []).append(conflict)
    increment_module_count(context, "heading_style_conflict")


def detect_caption_type(text: str) -> str | None:
    """识别表题和图题，例如“表1 xxx”“图 1-1 xxx”。"""
    # References to figures inside prose are not captions and must retain numbering.
    if re.match(r"^[图表]\s*\d+(?:[-.－—]\d+)?\s*(?:为|是|中|所示|显示|表明|说明|给出)", text) or len(text) > 120:
        return None
    if TABLE_CAPTION_PATTERN.match(text):
        return "table_caption"
    if FIGURE_CAPTION_PATTERN.match(text):
        return "figure_caption"
    return None


def detect_paragraph_type(text, index, context, paragraph=None):
    """用正则和上下文识别段落类型，不使用 AI 判断。"""
    if not text:
        return "empty"

    if index in context.get("paper_title_indices", set()):
        return "paper_title"

    if is_toc_title(text):
        context["in_toc_section"] = True
        context["in_abstract_section"] = False
        context["in_english_abstract_section"] = False
        return "body"

    if context.get("in_toc_section"):
        if is_toc_entry(text):
            return "body"
        context["in_toc_section"] = False

    if is_reference_title(text):
        context["in_reference_section"] = True
        context["in_abstract_section"] = False
        context["in_english_abstract_section"] = False
        return "reference_title"

    if context.get("in_reference_section"):
        if is_reference_exit_title(text):
            context["in_reference_section"] = False
        else:
            return "reference_item"

    if is_english_abstract(text):
        context["in_english_abstract_section"] = True
        context["english_abstract_content_seen"] = False
        context["structured_english_abstract"] = False
        context["in_abstract_section"] = False
        return "abstract_en_title"

    if is_english_keywords(text):
        context["in_english_abstract_section"] = False
        context["in_abstract_section"] = False
        return "keywords_en"

    if is_chinese_abstract_title(text):
        context["in_abstract_section"] = True
        context["in_english_abstract_section"] = False
        return "abstract_title"

    if re.match(r"^(?:【\s*摘\s*要\s*】\s*[:：]?|摘\s*要\s*[:：])", text):
        context["in_abstract_section"] = False
        return "abstract_content"

    if is_explicit_keywords(text):
        context["in_abstract_section"] = False
        context["in_english_abstract_section"] = False
        return "keywords"

    if context.get("in_abstract_section"):
        if detect_heading_type_from_style(paragraph) or detect_text_heading_type(text, context, paragraph):
            context["in_abstract_section"] = False
        else:
            return "abstract_content"

    if context.get("in_english_abstract_section"):
        named_heading = detect_named_heading_type(text)
        if named_heading and not context.get("english_abstract_content_seen"):
            context["structured_english_abstract"] = True
        leave_abstract = text.strip().casefold() == "introduction" or (
            named_heading and context.get("english_abstract_content_seen") and not context.get("structured_english_abstract")
        )
        if leave_abstract or detect_heading_type_from_style(paragraph) or (not named_heading and detect_text_heading_type(text, context, paragraph)):
            context["in_english_abstract_section"] = False
        else:
            context["english_abstract_content_seen"] = True
            return "abstract_en_content"

    caption_type = detect_caption_type(text)
    if caption_type:
        return caption_type

    text_heading_type = context.get("automatic_headings", {}).get(index) or detect_text_heading_type(text, context, paragraph)
    if text_heading_type:
        record_heading_style_conflict(
            context,
            paragraph,
            index,
            text,
            text_heading_type,
        )
        return text_heading_type

    if not context.get("first_non_empty_seen") and context.get("paper_title_index") is None:
        if text not in PROTECTED_FIRST_PARAGRAPHS and is_paper_title_candidate(text, paragraph):
            return "paper_title"

    heading_style_type = detect_heading_type_from_style(paragraph)
    if heading_style_type:
        return heading_style_type

    if not is_title_candidate(text):
        return "body"

    return "body"


def detect_paragraph_warning(text: str, paragraph_type: str) -> str | None:
    """识别可能存在误判风险的段落。"""
    if paragraph_type == "body" and (TABLE_CAPTION_PATTERN.match(text) or FIGURE_CAPTION_PATTERN.match(text)) and detect_caption_type(text) is None:
        return "该段以图表编号开头，但更像说明正文或长题注，已保留编号并按正文处理，请人工确认。"
    if paragraph_type == "table_caption" and len(text) > 50:
        return "该段类似表题，但长度较长，请人工确认"

    if paragraph_type == "figure_caption" and len(text) > 50:
        return "该段类似图题，但长度较长，请人工确认"

    if paragraph_type != "body" or len(text) <= 40:
        return None

    if re.match(r"^\d+(?:\.\d+)*[\.、\s]", text):
        return "该段以数字编号开头但长度较长，已按正文处理"

    if re.match(rf"^[{CHINESE_NUMBER}]+[、.．]", text):
        return "该段类似中文一级标题，但长度较长，请人工确认"

    return None


def get_style_config(
    template, paragraph_type, report=None, paragraph_index=None, text=""
):
    """根据段落类型获取模板样式；缺失时回退到 body。"""
    styles = template["styles"]
    specific = {"abstract_title": "abstract_cn_title", "abstract_content": "abstract_cn_content"}.get(paragraph_type)
    if specific in styles:
        return styles[specific], specific, None
    if paragraph_type in styles:
        return styles[paragraph_type], paragraph_type, None
    fallback_type = STYLE_FALLBACKS.get(paragraph_type)
    if fallback_type in styles:
        return styles[fallback_type], fallback_type, None

    runtime = template.setdefault("_runtime", {})
    runtime["fallback_count"] = runtime.get("fallback_count", 0) + 1
    warned = runtime.setdefault("warned_missing_styles", set())
    message = f"未找到 {paragraph_type} 样式，已使用 body 样式代替"

    if paragraph_type not in warned:
        print(f"警告：{message}")
        warned.add(paragraph_type)

    if report is not None:
        add_warning(report, paragraph_index, text, message)

    return styles["body"], "body", message


def get_alignment(alignment_name: str | None):
    """把模板里的对齐字符串映射为 python-docx 枚举。"""
    alignments = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    return alignments.get((alignment_name or "left").lower(), WD_ALIGN_PARAGRAPH.LEFT)


def get_paragraph_style_name(paragraph) -> str:
    """Return a paragraph style name without raising on unusual documents."""
    try:
        return getattr(paragraph.style, "name", "") or ""
    except (KeyError, ValueError, AttributeError):
        return ""


def get_paragraph_style_id(paragraph) -> str:
    """Return a paragraph style id from python-docx or direct XML."""
    try:
        style_id = getattr(paragraph.style, "style_id", "") or ""
        if style_id:
            return style_id
    except (KeyError, ValueError, AttributeError):
        pass

    p_pr = paragraph._p.pPr
    if p_pr is None:
        return ""
    p_style = p_pr.find(qn("w:pStyle"))
    if p_style is None:
        return ""
    return p_style.get(qn("w:val"), "") or ""


def detect_heading_type_from_style(paragraph) -> str | None:
    """Resolve outline levels and inherited styles before text heuristics."""
    if paragraph is None:
        return None
    return heading_evidence(paragraph)["type"]


def is_word_heading_paragraph(paragraph) -> bool:
    """Return True when a paragraph already carries a Word Heading style."""
    return detect_heading_type_from_style(paragraph) is not None


HEADING_PARAGRAPH_STYLES = {
    "heading_1": "Heading 1",
    "heading_2": "Heading 2",
    "heading_3": "Heading 3",
}


def get_heading_style_conflict(context: dict, paragraph_index: int) -> dict | None:
    """Return recorded heading style conflict for one paragraph, if any."""
    for conflict in context.get("heading_style_conflicts", []):
        if conflict.get("paragraph_index") == paragraph_index:
            return conflict
    return None


def get_heading_style_conflict_message(conflict: dict) -> str:
    """Build a report warning for a heading style/text-numbering conflict."""
    original_style = conflict.get("original_style") or conflict.get("original_style_id") or "未知样式"
    return (
        "标题层级冲突：原 Word 样式 "
        f"{original_style} 对应 {conflict.get('original_style_level')}，"
        f"但文本编号对应 {conflict.get('text_pattern_level')}，"
        "已按文本编号层级处理。"
    )


def set_local_heading_level(paragraph, paragraph_type):
    properties = paragraph._p.get_or_add_pPr()
    outline = properties.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        properties.insert_element_before(outline, "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
    level = int(paragraph_type[-1]) - 1 if paragraph_type in HEADING_PARAGRAPH_STYLES else 9
    outline.set(qn("w:val"), str(level))


def sync_heading_paragraph_style(
    paragraph,
    paragraph_type: str,
    context: dict,
) -> None:
    """Set Word paragraph style to the detected Heading level when possible."""
    target_style = HEADING_PARAGRAPH_STYLES.get(paragraph_type)
    if not target_style:
        return
    if detect_heading_type_from_style(paragraph) == paragraph_type:
        return
    before_style = get_paragraph_style_name(paragraph)
    before_style_id = get_paragraph_style_id(paragraph)
    previous_style = paragraph.style
    numbering = effective_numbering(paragraph)
    if numbering and (numbering["sources"].get("num_id", "").startswith("style:") or numbering["level"] is None):
        # Keep the style-to-numbering link intact; set the outline locally instead.
        set_local_heading_level(paragraph, paragraph_type)
        return
    try:
        paragraph.style = target_style
    except (KeyError, ValueError, AttributeError) as exc:
        context.setdefault("heading_style_sync_failures", []).append(
            {
                "target_style": target_style,
                "original_style": before_style,
                "original_style_id": before_style_id,
                "error": str(exc),
            }
        )
        increment_module_count(context, "heading_style_sync_failed")
        return

    # Changing pStyle must not disconnect an inherited list from its definition.
    if numbering is not None:
        numpr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        numpr.get_or_add_numId().val = int(numbering["num_id"])
        numpr.get_or_add_ilvl().val = int(numbering["level"])
    elif (target_numbering := effective_numbering(paragraph)) and target_numbering["enabled"]:
        paragraph.style = previous_style
        set_local_heading_level(paragraph, paragraph_type)
    outline = paragraph._p.pPr.find(qn("w:outlineLvl"))
    if outline is not None:
        outline.set(qn("w:val"), str(int(paragraph_type[-1]) - 1))

    after_style = get_paragraph_style_name(paragraph)
    if after_style != before_style:
        context.setdefault("heading_style_sync", []).append(
            {
                "target_style": target_style,
                "original_style": before_style,
                "original_style_id": before_style_id,
                "final_style": after_style,
                "final_style_id": get_paragraph_style_id(paragraph),
            }
        )
        increment_module_count(context, "heading_style_synced")


def apply_run_font(run, style_config) -> None:
    """设置 run 的字体、字号、加粗，并确保中文 eastAsia 字体生效。"""
    font_name = style_config.get("font", DEFAULT_STYLE["font"])
    font_size = style_config.get("size_pt", DEFAULT_STYLE["size_pt"])
    color = style_config.get("color", DEFAULT_STYLE["color"])
    bold = style_config.get("bold", DEFAULT_STYLE["bold"])
    italic = style_config.get("italic", DEFAULT_STYLE["italic"])
    underline = style_config.get("underline", DEFAULT_STYLE["underline"])

    run.font.name = font_name
    run.font.size = Pt(font_size)
    if color:
        run.font.color.rgb = RGBColor.from_string(str(color))
    run.bold = bold
    run.italic = italic
    if "underline" in style_config:
        run.font.underline = bool(underline)

    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)

    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:cs"), font_name)


def resolve_paragraph_space_pt(style_config: dict, pt_field: str, lines_field: str, default_value: int | float) -> int | float:
    """Resolve paragraph spacing in points, accepting explicit line units."""
    if lines_field in style_config and is_number(style_config[lines_field]):
        font_size = style_config.get("size_pt", DEFAULT_STYLE["size_pt"])
        if not is_number(font_size):
            font_size = DEFAULT_STYLE["size_pt"]
        return float(style_config[lines_field]) * float(font_size)
    if pt_field in style_config:
        return style_config.get(pt_field, default_value)
    return default_value


def apply_paragraph_style(paragraph, style_config) -> None:
    """根据模板 style_config 设置段落格式。"""
    paragraph_format = paragraph.paragraph_format

    paragraph.alignment = get_alignment(style_config.get("alignment"))
    paragraph_format.line_spacing = style_config.get(
        "line_spacing", DEFAULT_STYLE["line_spacing"]
    )
    # A source document grid can override the template's explicit line spacing.
    properties = paragraph._p.get_or_add_pPr()
    snap_to_grid = properties.find(qn("w:snapToGrid"))
    if snap_to_grid is None:
        snap_to_grid = OxmlElement("w:snapToGrid")
        properties.insert_element_before(
            snap_to_grid, "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
            "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
            "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
            "w:sectPr", "w:pPrChange",
        )
    snap_to_grid.set(qn("w:val"), "0")
    paragraph_format.first_line_indent = Pt(
        style_config.get("first_line_indent_pt", DEFAULT_STYLE["first_line_indent_pt"])
    )
    paragraph_format.space_before = Pt(
        resolve_paragraph_space_pt(
            style_config,
            "space_before_pt",
            "space_before_lines",
            DEFAULT_STYLE["space_before_pt"],
        )
    )
    paragraph_format.space_after = Pt(
        resolve_paragraph_space_pt(
            style_config,
            "space_after_pt",
            "space_after_lines",
            DEFAULT_STYLE["space_after_pt"],
        )
    )
    if "keep_with_next" in style_config:
        paragraph_format.keep_with_next = bool(style_config["keep_with_next"])
    if "keep_together" in style_config:
        paragraph_format.keep_together = bool(style_config["keep_together"])

    for run in paragraph.runs:
        apply_run_font(run, style_config)


LATIN_DIGIT_TEXT_PATTERN = re.compile(r"[A-Za-z0-9]+(?:[._/-][A-Za-z0-9]+)*")
REFERENCE_LATIN_DIGIT_TEXT_PATTERN = re.compile(r"[A-Za-z0-9,.:;()\[\]\-/]+")
CHINESE_KEYWORDS_LABEL_PATTERN = re.compile(r"^\s*(?:【\s*关键词\s*】\s*[:：]?|关键词\s*[:：])")
ENGLISH_KEYWORDS_LABEL_PATTERN = re.compile(
    r"^\s*(?:keywords|key\s+words)[:：]",
    re.IGNORECASE,
)
LATIN_DIGIT_TARGET_PARAGRAPH_TYPES = {
    "paper_title",
    "abstract_cn_title",
    "abstract_cn_content",
    "keywords_cn_label",
    "keywords_cn_content",
    "abstract_title",
    "abstract_content",
    "keywords",
    "abstract_en_title",
    "abstract_en_content",
    "keywords_en",
    "keywords_en_label",
    "keywords_en_content",
    "heading_1",
    "heading_2",
    "heading_3",
    "body",
    "reference_item",
}


def get_latin_digit_format(template) -> dict | None:
    """Return a usable top-level English/digit character format rule."""
    config = template.get("latin_digit_format")
    if not isinstance(config, dict):
        return None
    if not isinstance(config.get("font"), str) and not is_number(config.get("size_pt")):
        return None
    return config


def get_reference_latin_digit_format(template) -> dict | None:
    """Return a usable reference English/digit character format rule."""
    config = template.get("reference_latin_digit_format")
    if not isinstance(config, dict):
        return None
    if not isinstance(config.get("font"), str) and not is_number(config.get("size_pt")):
        return None
    if config.get("scope", "reference") != "reference":
        return None
    return config


def get_xml_local_name(element) -> str:
    """Return the local name for an OOXML element tag."""
    tag = getattr(element, "tag", "") or ""
    return tag.rsplit("}", 1)[-1]


def paragraph_has_complex_latin_digit_structure(paragraph) -> bool:
    """Protect fields, hyperlinks, equations, and other complex text containers."""
    complex_local_names = {
        "hyperlink",
        "fldSimple",
        "fldChar",
        "instrText",
        "oMath",
        "oMathPara",
    }
    return any(
        get_xml_local_name(element) in complex_local_names
        for element in paragraph._p.iter()
    )


def run_is_plain_text_for_latin_digit_split(run) -> bool:
    """Only plain text runs can be split without changing Word structures."""
    return all(
        get_xml_local_name(child) in {"rPr", "t"}
        for child in run._r
    )


def paragraph_has_only_plain_text_runs(paragraph) -> bool:
    """Return True when every run can be safely cloned and split."""
    return all(run_is_plain_text_for_latin_digit_split(run) for run in paragraph.runs)


def split_latin_digit_text(
    text: str,
    pattern=LATIN_DIGIT_TEXT_PATTERN,
) -> list[tuple[str, bool]]:
    """Split text into unchanged non-Latin pieces and Latin/digit pieces."""
    pieces = []
    start = 0
    for match in pattern.finditer(text):
        if match.start() > start:
            pieces.append((text[start:match.start()], False))
        pieces.append((match.group(0), True))
        start = match.end()
    if start < len(text):
        pieces.append((text[start:], False))
    return [piece for piece in pieces if piece[0]]


def set_latin_digit_run_font(run, config: dict) -> None:
    """Set ascii/hAnsi font and optional size without changing eastAsia."""
    font_name = config.get("font")
    if isinstance(font_name, str) and font_name.strip():
        r_pr = run._element.get_or_add_rPr()
        r_fonts = r_pr.rFonts
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.append(r_fonts)
        r_fonts.set(qn("w:ascii"), font_name)
        r_fonts.set(qn("w:hAnsi"), font_name)

    if is_number(config.get("size_pt")):
        run.font.size = Pt(config["size_pt"])


def format_plain_text_run_latin_digits(
    run,
    paragraph,
    config: dict,
    pattern=LATIN_DIGIT_TEXT_PATTERN,
) -> int:
    """Format Latin/digit fragments in one safe run, splitting if mixed."""
    pieces = split_latin_digit_text(run.text, pattern)
    latin_piece_count = sum(1 for _, is_latin_piece in pieces if is_latin_piece)
    if latin_piece_count == 0:
        return 0
    if len(pieces) == 1 and pieces[0][1]:
        set_latin_digit_run_font(run, config)
        return 1

    parent = run._r.getparent()
    if parent is None or Run is None:
        return 0

    insert_index = parent.index(run._r)
    for piece_text, is_latin_piece in pieces:
        cloned_run_element = deepcopy(run._r)
        cloned_run = Run(cloned_run_element, paragraph)
        cloned_run.text = piece_text
        if is_latin_piece:
            set_latin_digit_run_font(cloned_run, config)
        parent.insert(insert_index, cloned_run_element)
        insert_index += 1
    parent.remove(run._r)
    return latin_piece_count


def apply_latin_digit_config_to_paragraph(
    paragraph,
    config: dict,
    report,
    context: dict,
    paragraph_index: int,
    text: str,
    count_key: str,
    skipped_key: str,
    warning_prefix: str,
    pattern=LATIN_DIGIT_TEXT_PATTERN,
) -> None:
    """Apply one English/digit config to a plain paragraph without changing text."""
    if not pattern.search(text):
        return

    if paragraph_has_complex_latin_digit_structure(paragraph):
        increment_module_count(context, skipped_key)
        increment_module_count(context, "skipped_complex_paragraph_count")
        add_warning(
            report,
            paragraph_index,
            text,
            f"{warning_prefix}处理已跳过字段、超链接或公式等复杂结构，请人工确认。",
        )
        return

    formatted_count = 0
    skipped_run_count = 0
    for run in list(paragraph.runs):
        if not pattern.search(run.text):
            continue
        if not run_is_plain_text_for_latin_digit_split(run):
            skipped_run_count += 1
            continue
        formatted_count += format_plain_text_run_latin_digits(
            run,
            paragraph,
            config,
            pattern,
        )

    if formatted_count:
        increment_module_count(context, count_key, formatted_count)
    if skipped_run_count:
        increment_module_count(context, skipped_key, skipped_run_count)
        increment_module_count(context, "skipped_complex_paragraph_count")
        add_warning(
            report,
            paragraph_index,
            text,
            f"{warning_prefix}处理跳过了包含复杂 run 内容的段落片段，请人工确认。",
        )


def apply_latin_digit_format_to_paragraph(
    paragraph,
    paragraph_type: str,
    template,
    report,
    context: dict,
    paragraph_index: int,
    text: str,
    is_toc_context: bool,
) -> None:
    """Apply a conservative English/digit rule to eligible plain text runs."""
    config = get_latin_digit_format(template)
    if config is None or paragraph_type not in LATIN_DIGIT_TARGET_PARAGRAPH_TYPES:
        return
    scope = config.get("scope", "global")
    if scope == "body" and paragraph_type != "body":
        return
    if scope == "abstract" and paragraph_type not in {
        "abstract_title",
        "abstract_content",
        "abstract_cn_title",
        "abstract_cn_content",
        "abstract_en_title",
        "abstract_en_content",
    }:
        return
    if scope == "heading" and paragraph_type not in {"heading_1", "heading_2", "heading_3"}:
        return
    if is_toc_context or is_toc_title(text) or is_toc_entry(text):
        return

    apply_latin_digit_config_to_paragraph(
        paragraph,
        config,
        report,
        context,
        paragraph_index,
        text,
        "latin_digit_format",
        "latin_digit_format_skipped",
        "英文/数字格式",
        LATIN_DIGIT_TEXT_PATTERN,
    )


def apply_reference_latin_digit_format_to_paragraph(
    paragraph,
    paragraph_type: str,
    template,
    report,
    context: dict,
    paragraph_index: int,
    text: str,
) -> None:
    """Apply reference-only English/digit formatting to reference items."""
    if paragraph_type != "reference_item":
        return
    config = get_reference_latin_digit_format(template)
    if config is None:
        return
    apply_latin_digit_config_to_paragraph(
        paragraph,
        config,
        report,
        context,
        paragraph_index,
        text,
        "reference_latin_digit_format",
        "reference_latin_digit_format_skipped",
        "参考文献英文/数字格式",
        REFERENCE_LATIN_DIGIT_TEXT_PATTERN,
    )


def get_keyword_label_end(paragraph_type: str, text: str) -> int | None:
    """Return the label/content split point for keyword paragraphs."""
    if paragraph_type == "keywords":
        match = CHINESE_KEYWORDS_LABEL_PATTERN.match(text)
    elif paragraph_type == "keywords_en":
        match = ENGLISH_KEYWORDS_LABEL_PATTERN.match(text)
    else:
        match = None
    return match.end() if match is not None else None


def clone_run_with_text(run, paragraph, text: str):
    """Clone one run and replace its plain text."""
    cloned_run_element = deepcopy(run._r)
    cloned_run = Run(cloned_run_element, paragraph)
    cloned_run.text = text
    return cloned_run_element, cloned_run


def split_run_by_keyword_boundary(run, paragraph, start_index: int, label_end: int, label_style: dict, content_style: dict) -> tuple[bool, bool]:
    """Split one plain run around the keyword label boundary."""
    run_text = run.text
    run_end = start_index + len(run_text)
    if not run_text:
        return False, False

    if run_end <= label_end:
        apply_run_font(run, label_style)
        return True, False
    if start_index >= label_end:
        apply_run_font(run, content_style)
        return False, True

    label_text = run_text[: label_end - start_index]
    content_text = run_text[label_end - start_index :]
    parent = run._r.getparent()
    if parent is None or Run is None:
        return False, False

    insert_index = parent.index(run._r)
    if label_text:
        label_run_element, label_run = clone_run_with_text(run, paragraph, label_text)
        apply_run_font(label_run, label_style)
        parent.insert(insert_index, label_run_element)
        insert_index += 1
    if content_text:
        content_run_element, content_run = clone_run_with_text(run, paragraph, content_text)
        apply_run_font(content_run, content_style)
        parent.insert(insert_index, content_run_element)
    parent.remove(run._r)
    return bool(label_text), bool(content_text)


def apply_keyword_label_content_format(
    paragraph,
    paragraph_type: str,
    template,
    report,
    context: dict,
    paragraph_index: int,
) -> None:
    """Apply separate run styles to keyword label and content when safe."""
    full_text = paragraph.text
    label_end = get_keyword_label_end(paragraph_type, full_text)
    if label_end is None:
        return

    styles = template.get("styles", {})
    if paragraph_type == "keywords":
        label_key = "keywords_cn_label"
        content_key = "keywords_cn_content"
        skipped_key = "keywords_cn_label_content_skipped"
        warning_prefix = "中文关键词"
    else:
        label_key = "keywords_en_label"
        content_key = "keywords_en_content"
        skipped_key = "keywords_en_label_content_skipped"
        warning_prefix = "英文关键词"

    label_style = styles.get(label_key)
    content_style = styles.get(content_key)
    if not isinstance(label_style, dict) or not isinstance(content_style, dict):
        return

    if (
        paragraph_has_complex_latin_digit_structure(paragraph)
        or not paragraph_has_only_plain_text_runs(paragraph)
    ):
        increment_module_count(context, skipped_key)
        increment_module_count(context, "skipped_complex_paragraph_count")
        add_warning(
            report,
            paragraph_index,
            full_text,
            f"{warning_prefix}段落包含超链接、域、公式或复杂 run，已保守跳过 label/content 拆分。",
        )
        return

    original_text = paragraph.text
    label_formatted = False
    content_formatted = False
    cursor = 0
    for run in list(paragraph.runs):
        run_length = len(run.text)
        run_label, run_content = split_run_by_keyword_boundary(
            run,
            paragraph,
            cursor,
            label_end,
            label_style,
            content_style,
        )
        label_formatted = label_formatted or run_label
        content_formatted = content_formatted or run_content
        cursor += run_length

    if paragraph.text != original_text:
        increment_module_count(context, skipped_key)
        increment_module_count(context, "skipped_complex_paragraph_count")
        add_warning(
            report,
            paragraph_index,
            original_text,
            f"{warning_prefix}段落拆分后文本校验失败，请人工确认。",
        )
        return

    if label_formatted:
        increment_module_count(context, f"{label_key}_count")
    if content_formatted and label_end < len(original_text):
        increment_module_count(context, f"{content_key}_count")


def apply_reference_hanging_indent(
    paragraph,
    paragraph_type: str,
    style_config: dict,
    context: dict,
) -> None:
    """Apply reference-item hanging indent in character units."""
    if paragraph_type != "reference_item":
        return
    hanging_chars = style_config.get("hanging_indent_chars")
    if not is_number(hanging_chars):
        return
    size_pt = style_config.get("size_pt", DEFAULT_STYLE["size_pt"])
    if not is_number(size_pt):
        size_pt = DEFAULT_STYLE["size_pt"]
    indent_pt = float(hanging_chars) * float(size_pt)
    paragraph_format = paragraph.paragraph_format
    paragraph_format.left_indent = Pt(indent_pt)
    paragraph_format.first_line_indent = Pt(-indent_pt)
    increment_module_count(context, "reference_hanging_indent")


def remove_child_by_tag(parent, tag_name: str) -> bool:
    """Remove a direct XML child by tag name."""
    child = parent.find(qn(tag_name))
    if child is None:
        return False
    parent.remove(child)
    return True


def paragraph_style_is_list_style(style) -> bool:
    """Return True when a paragraph style can carry Word list numbering."""
    seen = set()
    current = style
    while current is not None and id(current) not in seen:
        seen.add(id(current))
        style_id = (getattr(current, "style_id", "") or "").lower()
        style_name = (getattr(current, "name", "") or "").lower()
        if any(
            marker in style_id or marker in style_name
            for marker in ("list", "bullet", "number", "项目符号", "编号", "列表")
        ):
            return True

        p_pr = current._element.find(qn("w:pPr"))
        if p_pr is not None and p_pr.find(qn("w:numPr")) is not None:
            return True

        current = getattr(current, "base_style", None)

    return False


def clear_paragraph_numbering(paragraph) -> None:
    """Clear Word automatic numbering/list formatting without changing text."""
    if is_word_heading_paragraph(paragraph):
        return

    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    num_id = num_pr.find(qn("w:numId")) if num_pr is not None else None
    if num_id is not None and num_id.get(qn("w:val")) == "0":
        return
    remove_child_by_tag(p_pr, "w:numPr")

    try:
        if paragraph_style_is_list_style(paragraph.style):
            paragraph.style = "Normal"
    except (KeyError, ValueError, AttributeError):
        remove_child_by_tag(p_pr, "w:pStyle")


def clear_paragraph_list_formatting(paragraph) -> None:
    """Backward-compatible wrapper for numbering cleanup."""
    clear_paragraph_numbering(paragraph)


def apply_caption_pagination_defaults(paragraph, paragraph_type: str) -> None:
    """图题和表题默认与后续内容保持在同一页。"""
    if paragraph_type not in {"table_caption", "figure_caption"}:
        return

    clear_paragraph_numbering(paragraph)
    paragraph_format = paragraph.paragraph_format
    paragraph_format.keep_with_next = True
    paragraph_format.keep_together = True


def get_caption_pagination_report(paragraph, paragraph_type: str) -> dict | None:
    """生成图题/表题分页控制报告。"""
    if paragraph_type not in {"table_caption", "figure_caption"}:
        return None

    paragraph_format = paragraph.paragraph_format
    return {
        "keep_with_next": paragraph_format.keep_with_next is True,
        "keep_together": paragraph_format.keep_together is True,
    }


def apply_normal_style(doc, template) -> None:
    """用 body 样式设置 Word Normal 样式。"""
    body_style, _, _ = get_style_config(template, "body")
    style = doc.styles["Normal"]
    style.font.name = body_style.get("font", DEFAULT_STYLE["font"])
    style.font.size = Pt(body_style.get("size_pt", DEFAULT_STYLE["size_pt"]))
    if body_style.get("color"):
        style.font.color.rgb = RGBColor.from_string(str(body_style["color"]))

    r_pr = style._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)

    font_name = body_style.get("font", DEFAULT_STYLE["font"])
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:cs"), font_name)


def apply_page_settings(doc, template, preserved_sections=None) -> None:
    """根据模板 page 配置设置页面边距。"""
    page = template["page"]
    top_margin = Cm(page.get("top_margin_cm", 2.5))
    bottom_margin = Cm(page.get("bottom_margin_cm", 2.5))
    left_margin = Cm(page.get("left_margin_cm", 2.5))
    right_margin = Cm(page.get("right_margin_cm", 2.5))

    for index, section in enumerate(doc.sections):
        if index in (preserved_sections or set()):
            continue
        section.top_margin = top_margin
        section.bottom_margin = bottom_margin
        section.left_margin = left_margin
        section.right_margin = right_margin


def increment_stat(context: dict, paragraph_type: str) -> None:
    """记录每类段落处理数量。"""
    stats = context["stats"]
    stats[paragraph_type] = stats.get(paragraph_type, 0) + 1


def analyze_heading_numbering_system(paragraphs) -> dict:
    """Pre-scan document heading numbering signals for cautious Arabic level-1 detection."""
    arabic_level1_count = 0
    numeric_subheading_count = 0
    chapter_heading_count = 0
    arabic_level1_numbers = set()
    numeric_subheading_root_numbers = set()
    for paragraph in paragraphs:
        text = paragraph.text.strip()
        if not text or is_toc_title(text) or is_toc_entry(text):
            continue
        if is_chapter_heading(text):
            chapter_heading_count += 1
            continue
        root_number = get_numeric_subheading_root_number(text)
        if root_number is not None:
            numeric_subheading_root_numbers.add(root_number)
        number_heading_type = detect_number_heading_type(text)
        if number_heading_type in {"heading_2", "heading_3"}:
            numeric_subheading_count += 1
        elif is_arabic_level1_heading_candidate(text):
            arabic_level1_count += 1
            level1_number = get_arabic_level1_heading_number(text)
            if level1_number is not None:
                arabic_level1_numbers.add(level1_number)

    allowed_level1_numbers = sorted(
        arabic_level1_numbers.intersection(numeric_subheading_root_numbers),
        key=lambda value: int(value),
    )

    return {
        "arabic_level1_count": arabic_level1_count,
        "numeric_subheading_count": numeric_subheading_count,
        "chapter_heading_count": chapter_heading_count,
        "arabic_heading_level1_numbers": allowed_level1_numbers,
        "numeric_subheading_root_numbers": sorted(
            numeric_subheading_root_numbers,
            key=lambda value: int(value),
        ),
        "arabic_heading_system": bool(allowed_level1_numbers),
    }


def clear_empty_paragraph_numbering_near_captions(paragraphs, index: int) -> None:
    """Clear list markers from an empty paragraph adjacent to a caption."""
    paragraph = paragraphs[index]
    if paragraph.text.strip():
        return

    adjacent_indexes = [index - 1, index + 1]
    for adjacent_index in adjacent_indexes:
        if adjacent_index < 0 or adjacent_index >= len(paragraphs):
            continue
        if detect_caption_type(paragraphs[adjacent_index].text.strip()):
            clear_paragraph_numbering(paragraph)
            return


def record_paragraph_module_detection(
    context: dict,
    text: str,
    paragraph_type: str,
    was_reference_section: bool,
    was_toc_section: bool,
) -> None:
    """记录段落级模块命中，不影响格式处理决策。"""
    if not text:
        return

    if is_english_abstract(text):
        increment_module_count(context, "abstract_en")
    if is_english_keywords(text):
        increment_module_count(context, "keywords_en")

    if is_toc_title(text):
        increment_module_count(context, "toc")
        mark_module_flag(context, "toc_protected")
    elif (was_toc_section or context.get("in_toc_section")) and is_toc_entry(text):
        increment_module_count(context, "toc")
        mark_module_flag(context, "toc_protected")

    if paragraph_type in {"heading_1", "heading_2", "heading_3"}:
        increment_module_count(context, "heading")
    elif (
        not was_reference_section
        and not was_toc_section
        and not context.get("in_toc_section")
        and (
            detect_number_heading_type(text)
            or re.match(rf"^[{CHINESE_NUMBER}]+[\u3001.\uFF0E]\s*\S+", text)
        )
    ):
        increment_module_count(context, "heading")

    if is_reference_exit_title(text):
        increment_module_count(context, "appendix")
        if was_reference_section:
            mark_module_flag(context, "appendix_boundary")


def record_toc_module_detection_from_xml(doc, context: dict) -> None:
    """Detect Word TOC blocks that python-docx does not expose in doc.paragraphs."""
    if context.get("module_counts", {}).get("toc"):
        return

    toc_count = 0
    for paragraph_element in doc._element.body.iter(qn("w:p")):
        p_pr = paragraph_element.find(qn("w:pPr"))
        style_element = p_pr.find(qn("w:pStyle")) if p_pr is not None else None
        style_id = style_element.get(qn("w:val")) if style_element is not None else ""
        text = "".join(
            text_element.text or ""
            for text_element in paragraph_element.iter(qn("w:t"))
        ).strip()

        if style_id.upper().startswith("TOC") or is_toc_title(text) or is_toc_entry(text):
            toc_count += 1

    if toc_count:
        for _ in range(toc_count):
            increment_module_count(context, "toc")
        mark_module_flag(context, "toc_protected")


def get_protected_toc_indices(paragraphs) -> set[int]:
    """Locate TOC styles and fields, including fields spanning paragraphs."""
    protected = set()
    fields = []
    in_text_toc = False
    for index, paragraph in enumerate(paragraphs):
        if get_paragraph_style_id(paragraph).upper().startswith("TOC"):
            protected.add(index)
        text = paragraph.text.strip()
        if is_toc_title(text):
            in_text_toc = True
            protected.add(index)
        elif in_text_toc and is_toc_entry(text):
            protected.add(index)
        elif text:
            in_text_toc = False
        for node in paragraph._p.iter():
            if node.tag == qn("w:fldSimple"):
                if re.match(r"^\s*TOC\b", node.get(qn("w:instr"), ""), re.I):
                    protected.add(index)
            elif node.tag == qn("w:fldChar"):
                kind = node.get(qn("w:fldCharType"))
                if kind == "begin":
                    fields.append([index, ""])
                elif kind == "end" and fields:
                    start, instruction = fields.pop()
                    if re.match(r"^\s*TOC\b", instruction, re.I):
                        protected.update(range(start, index + 1))
            elif node.tag == qn("w:instrText") and fields:
                fields[-1][1] += node.text or ""
    for start, instruction in fields:
        if re.match(r"^\s*TOC\b", instruction, re.I):
            protected.update(range(start, len(paragraphs)))
    return protected


def format_normal_paragraphs(doc, template, report, review=None) -> dict:
    """遍历普通段落，识别论文结构并按模板套用格式。"""
    paragraphs = doc.paragraphs
    heading_numbering_system = analyze_heading_numbering_system(paragraphs)
    context = {
        "first_non_empty_seen": False,
        "in_reference_section": False,
        "in_abstract_section": False,
        "in_english_abstract_section": False,
        "in_toc_section": False,
        "stats": {},
        "module_counts": {},
        "module_flags": {},
        "non_empty_count": 0,
        "heading_numbering_system": heading_numbering_system,
        "arabic_heading_system": heading_numbering_system["arabic_heading_system"],
        "arabic_heading_level1_numbers": set(
            heading_numbering_system.get("arabic_heading_level1_numbers", [])
        ),
    }

    protected_toc_indices = get_protected_toc_indices(paragraphs)
    context["paper_title_index"] = find_document_title(paragraphs, protected_toc_indices)
    title_index = context["paper_title_index"]
    context["paper_title_indices"] = {title_index} if title_index is not None else set()
    if title_index is not None and paragraphs[title_index].text.strip().isupper():
        for following_index in range(title_index + 1, min(title_index + 3, len(paragraphs))):
            p = paragraphs[following_index]
            if p.alignment == WD_ALIGN_PARAGRAPH.CENTER and p.text.strip().isupper() and 12 <= len(p.text.strip()) <= 160:
                context["paper_title_indices"].add(following_index)
            else:
                break
    automatic = {}
    confirmed_numbering = set()
    for index, paragraph in enumerate(paragraphs):
        if index in protected_toc_indices or not is_title_candidate(paragraph.text.strip()):
            continue
        level = automatic_heading_level(paragraph)
        following = next((p.text.strip() for p in paragraphs[index + 1:index + 5] if p.text.strip()), "")
        if level and (detect_number_heading_type(following) or re.match(rf"^[（(][{CHINESE_NUMBER}]+[）)]", following)):
            automatic[index] = level
            numbering = effective_numbering(paragraph)
            confirmed_numbering.add((numbering["num_id"], numbering["level"]))
    for index, paragraph in enumerate(paragraphs):
        numbering = effective_numbering(paragraph)
        if index not in protected_toc_indices and numbering is not None and (numbering["num_id"], numbering["level"]) in confirmed_numbering and is_title_candidate(paragraph.text.strip()):
            automatic[index] = automatic_heading_level(paragraph)
    context["automatic_headings"] = automatic
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if paragraph._p.xpath('.//wp:anchor | .//w:pict'):
            add_warning(report, index, text, "该段包含浮动图片或旧式图形；位置随正文重排可能变化，请在 Word/WPS 中检查重叠、越界及图文分页。")
        if index in protected_toc_indices:
            increment_module_count(context, "toc")
            mark_module_flag(context, "toc_protected")
            context["in_toc_section"] = True
            context["in_abstract_section"] = False
            context["in_english_abstract_section"] = False
            if text:
                context["non_empty_count"] += 1
                context["first_non_empty_seen"] = True
            increment_stat(context, "body" if text else "empty")
            add_paragraph_report(report, index, text, "body", None, "目录内容已保护，未修改段落格式。")
            continue
        was_reference_section = bool(context.get("in_reference_section"))
        was_toc_section = bool(context.get("in_toc_section"))
        paragraph_type = detect_paragraph_type(text, index, context, paragraph)
        decision = (review or {}).get("paragraphs", {}).get(index, {})
        if decision.get("type", "auto") != "auto":
            paragraph_type = decision["type"]
            context["heading_style_conflicts"] = [c for c in context.get("heading_style_conflicts", []) if c.get("paragraph_index") != index]
            for flag in ("in_reference_section", "in_abstract_section", "in_english_abstract_section", "in_toc_section"):
                context[flag] = False
        record_paragraph_module_detection(
            context, text, paragraph_type, was_reference_section, was_toc_section
        )
        if decision.get("preserve"):
            add_paragraph_report(report, index, text, paragraph_type, None, "已按保护策略保留原格式。")
            increment_stat(context, paragraph_type)
            if text:
                context["non_empty_count"] += 1
                context["first_non_empty_seen"] = True
            continue
        warning_messages = []
        applied_style = None

        if paragraph_type != "empty":
            context["non_empty_count"] += 1

            paragraph_warning = detect_paragraph_warning(text, paragraph_type)
            if paragraph_warning:
                warning_messages.append(paragraph_warning)
                add_warning(report, index, text, paragraph_warning)

            heading_conflict = get_heading_style_conflict(context, index)
            if heading_conflict:
                heading_conflict_message = get_heading_style_conflict_message(
                    heading_conflict
                )
                heading_conflict["message"] = heading_conflict_message
                warning_messages.append(heading_conflict_message)
                add_warning(report, index, text, heading_conflict_message)

            style_config, applied_style, style_warning = get_style_config(
                template, paragraph_type, report, index, text
            )
            if style_warning:
                warning_messages.append(style_warning)

            if decision.get("type", "auto") != "auto" and paragraph_type not in HEADING_PARAGRAPH_STYLES and is_word_heading_paragraph(paragraph):
                set_local_heading_level(paragraph, paragraph_type)
            if (
                paragraph_type in NUMBERING_CLEAR_PARAGRAPH_TYPES
                and not is_word_heading_paragraph(paragraph)
            ):
                clear_paragraph_numbering(paragraph)
            sync_heading_paragraph_style(paragraph, paragraph_type, context)
            apply_paragraph_style(paragraph, style_config)
            apply_caption_pagination_defaults(paragraph, paragraph_type)
            apply_reference_hanging_indent(
                paragraph,
                paragraph_type,
                style_config,
                context,
            )
            apply_keyword_label_content_format(
                paragraph,
                paragraph_type,
                template,
                report,
                context,
                index,
            )
            apply_latin_digit_format_to_paragraph(
                paragraph,
                paragraph_type,
                template,
                report,
                context,
                index,
                text,
                was_toc_section or bool(context.get("in_toc_section")),
            )
            apply_reference_latin_digit_format_to_paragraph(
                paragraph,
                paragraph_type,
                template,
                report,
                context,
                index,
                text,
            )
            add_paragraph_report(
                report,
                index,
                text,
                paragraph_type,
                applied_style,
                "；".join(warning_messages) if warning_messages else None,
                pagination=get_caption_pagination_report(paragraph, paragraph_type),
            )
        else:
            clear_empty_paragraph_numbering_near_captions(doc.paragraphs, index)

        increment_stat(context, paragraph_type)

        if paragraph_type != "empty":
            context["first_non_empty_seen"] = True

    if context["non_empty_count"] == 0:
        add_warning(report, None, "", "未检测到有效正文段落")

    for kind, label in (("paper_title", "论文标题"), ("heading_1", "一级章节标题")):
        if not context["stats"].get(kind):
            add_warning(report, None, "", f"未识别到{label}，请核查文档结构；附件或纯表格可忽略此项。")
    by_index = {item["index"]: item for item in report["paragraphs"]}
    for warning in report["warnings"]:
        item = by_index.get(warning.get("paragraph_index"))
        warning["detected_type"] = item["detected_type"] if item else None
        warning["applied_style"] = item["applied_style"] if item else None

    record_toc_module_detection_from_xml(doc, context)

    return context


def iter_unique_table_cells(table):
    """Merged cells appear in multiple grid positions but share one XML cell."""
    seen = set()
    for row in table.rows:
        for cell in row.cells:
            if cell._tc not in seen:
                seen.add(cell._tc)
                yield cell


def iter_tables(tables):
    """递归获取普通表格和嵌套表格。"""
    for table in tables:
        yield table
        for cell in iter_unique_table_cells(table):
            yield from iter_tables(cell.tables)


def set_table_rows_cant_split(table) -> None:
    """设置表格行不跨页断行，不强制整个表格同页。"""
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))


def get_previous_paragraph_for_table(table):
    """Return the paragraph immediately before a table in document order."""
    if Paragraph is None:
        return None

    previous = table._element.getprevious()
    if previous is None or previous.tag != qn("w:p"):
        return None

    return Paragraph(previous, table._parent)


def get_next_paragraph_for_table(table):
    """Return the paragraph immediately after a table in document order."""
    if Paragraph is None:
        return None

    next_element = table._element.getnext()
    if next_element is None or next_element.tag != qn("w:p"):
        return None

    return Paragraph(next_element, table._parent)


def get_previous_table_caption(table):
    """Return the table caption paragraph directly before a table, if any."""
    paragraph = get_previous_paragraph_for_table(table)
    if paragraph is None:
        return None

    if detect_caption_type(paragraph.text.strip()) == "table_caption":
        return paragraph

    return None


def clear_empty_paragraph_numbering_near_table(table) -> None:
    """Clear list markers from empty paragraphs directly adjacent to a table."""
    for paragraph in (
        get_previous_paragraph_for_table(table),
        get_next_paragraph_for_table(table),
    ):
        if paragraph is not None and not paragraph.text.strip():
            clear_paragraph_numbering(paragraph)


def apply_table_pagination_basics(table) -> bool:
    """Keep only low-risk pagination controls around tables."""
    caption = get_previous_table_caption(table)
    if caption is not None:
        apply_caption_pagination_defaults(caption, "table_caption")

    clear_empty_paragraph_numbering_near_table(table)
    return caption is not None


def format_tables(doc, template, report, review=None) -> int:
    """统一设置表格内文字格式，并记录表格报告。"""
    table_count = 0
    preserved_tables = set()
    membership = section_membership(doc)
    user_tables = set()
    for index, table in enumerate(iter_tables(doc.tables)):
        if index in (review or {}).get("tables", set()):
            user_tables.add(table._tbl)
            user_tables.update(p for p in table._tbl.iterancestors() if p.tag == qn("w:tbl"))

    for table_index, table in enumerate(iter_tables(doc.tables)):
        table_count += 1
        row_count = len(table.rows)
        cell_count = sum(len(row.cells) for row in table.rows)
        budget = table_width_budget(doc, table, membership)
        width = declared_table_width(table, budget)
        inside_preserved = table._tbl in user_tables or any(parent in preserved_tables for parent in table._tbl.iterancestors())
        if inside_preserved or budget is None or width is None or width > budget:
            preserved_tables.add(table._tbl)
            warning = ("已按保护策略保留该表格及其嵌套内容的原格式。" if table._tbl in user_tables else
                       "所属外层表格已受保护，此嵌套表格保留原格式。" if inside_preserved else
                       "无法确定表格容器或声明宽度，已保留原格式；请人工检查嵌套或分节布局。" if budget is None or width is None else
                       "表格声明宽度超过所在节或栏的可用宽度，已保留原格式；请人工检查列宽、缩进及文字越界。")
            add_table_report(report, table_index, row_count, cell_count, None, warning)
            add_warning(report, None, "", warning)
            report["warnings"][-1].update(table_index=table_index, detected_type="table", applied_style=None)
            continue
        set_table_rows_cant_split(table)
        style_config, applied_style, style_warning = get_style_config(
            template, "table_text", report, None, f"表格 {table_index}"
        )
        cell_style_config = {
            field: value
            for field, value in style_config.items()
            if field not in {"keep_with_next", "keep_together"}
        }

        for cell in iter_unique_table_cells(table):
            for paragraph in cell.paragraphs:
                apply_paragraph_style(paragraph, cell_style_config)

        apply_table_pagination_basics(table)
        add_table_report(
            report,
            table_index,
            row_count,
            cell_count,
            applied_style,
            warning=style_warning,
        )

    return table_count


def build_report_stats(stats, report, template) -> dict:
    """生成报告和命令行共用的统计信息。"""
    report_stats = {key: stats.get(key, 0) for key in REPORT_STAT_KEYS}
    report_stats["table"] = stats.get("table", 0)
    module_counts = report.get("final_rules_debug", {}).get("debug_summary", {}).get(
        "execution_counts",
        {},
    )
    report_stats["keywords_cn_label_count"] = module_counts.get(
        "keywords_cn_label_count",
        0,
    )
    report_stats["keywords_cn_content_count"] = module_counts.get(
        "keywords_cn_content_count",
        0,
    )
    report_stats["keywords_en_label_count"] = module_counts.get(
        "keywords_en_label_count",
        0,
    )
    report_stats["keywords_en_content_count"] = module_counts.get(
        "keywords_en_content_count",
        0,
    )
    report_stats["reference_latin_digit_format_count"] = module_counts.get(
        "reference_latin_digit_format",
        0,
    )
    report_stats["reference_hanging_indent_count"] = module_counts.get(
        "reference_hanging_indent",
        0,
    )
    report_stats["skipped_complex_paragraph_count"] = module_counts.get(
        "skipped_complex_paragraph_count",
        0,
    )
    report_stats["style_fallback_count"] = template.get("_runtime", {}).get(
        "fallback_count", 0
    )
    report_stats["warning_count"] = len(report["warnings"])
    return report_stats


def open_document(input_path: Path):
    """打开 Word 文档，并把常见打开失败转换成中文提示。"""
    try:
        with ZipFile(input_path) as archive:
            entries = archive.infolist()
            if len(entries) > 10000 or sum(item.file_size for item in entries) > 100 * 1024 * 1024:
                raise ValueError("DOCX expanded size exceeds processing limit")
        return Document(str(input_path))
    except Exception as exc:
        raise UserFacingError(
            "无法打开 Word 文件。请确认它是有效的 .docx 文件，并且没有被 Word/WPS 占用。"
        ) from exc


def safe_save_document(doc, output_path: Path, overwrite: bool = False) -> None:
    """先保存到临时文件，再替换为最终输出文件，避免半成品。"""
    temp_path = None
    try:
        if output_path.exists() and not overwrite:
            raise UserFacingError(f"输出文件已存在：{output_path}\n如需覆盖，请添加 --overwrite 参数。")

        with tempfile.NamedTemporaryFile(
            prefix=f".{output_path.stem}.",
            suffix=".tmp.docx",
            dir=str(output_path.parent),
            delete=False,
        ) as temp_file:
            temp_path = Path(temp_file.name)

        doc.save(str(temp_path))

        if output_path.exists() and not overwrite:
            raise UserFacingError(f"输出文件已存在：{output_path}\n如需覆盖，请添加 --overwrite 参数。")

        os.replace(temp_path, output_path)
        temp_path = None
    except UserFacingError:
        raise
    except Exception as exc:
        raise UserFacingError(
            "保存输出文件失败。请确认目标文件没有被 Word/WPS 打开，并且当前目录有写入权限。"
        ) from exc
    finally:
        if temp_path and temp_path.exists():
            try:
                temp_path.unlink()
            except OSError:
                pass


def format_document(input_path: Path, output_path: Path, template, report, overwrite=False, review=None) -> dict:
    """读取 Word 文件，修改格式，安全保存新文件，并返回统计结果。"""
    doc = open_document(input_path)

    report["structure_analysis"] = analyze_document(doc)
    membership = section_membership(doc)
    review = deepcopy(review or {})
    review.setdefault("paragraphs", {})
    review.setdefault("tables", set())
    proposed = deepcopy(doc)
    apply_page_settings(proposed, template)
    proposed_membership = section_membership(proposed)
    automatic_tables = set()
    for index, table in enumerate(iter_tables(proposed.tables)):
        budget = table_width_budget(proposed, table, proposed_membership)
        width = declared_table_width(table, budget)
        if budget is None or width is None or width > budget:
            automatic_tables.add(index)
    review["tables"].update(automatic_tables)
    automatic_paragraphs = []
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph._p.xpath('.//wp:anchor | .//w:pict'):
            review["paragraphs"].setdefault(index, {})["preserve"] = True
            automatic_paragraphs.append(index)
    preserved_sections = set()
    for index, paragraph in enumerate(doc.paragraphs):
        if (review or {}).get("paragraphs", {}).get(index, {}).get("preserve") and paragraph._p.xpath('.//w:drawing | .//w:pict'):
            preserved_sections.add(membership.get(paragraph._p))
    for index, table in enumerate(iter_tables(doc.tables)):
        if index in (review or {}).get("tables", set()):
            preserved_sections.add(membership.get(table._tbl))
    if report["structure_analysis"]["unresolved_sections"]:
        add_warning(report, None, "", "存在未支持的包装分节结构，已跳过页面设置和表格处理；请人工检查文档布局。")
    else:
        apply_page_settings(doc, template, preserved_sections)
    if preserved_sections:
        add_warning(report, None, "", "为保护复杂版面或用户指定内容，未套用所在节的页边距要求；相邻正文重排仍可能改变位置，请检查最终页面。")
    report["preserved_sections"] = sorted(i for i in preserved_sections if i is not None)
    report["automatic_protection"] = {"tables": sorted(automatic_tables), "paragraphs": automatic_paragraphs}

    context = format_normal_paragraphs(doc, template, report, review)
    context["stats"]["table"] = format_tables(doc, template, report, review)

    source_paragraphs = {item["paragraph_index"]: item for item in report["structure_analysis"]["paragraphs"]
                         if item["paragraph_index"] is not None}
    for item in report["paragraphs"]:
        source = source_paragraphs.get(item["index"])
        if source is not None:
            item["source_node_id"] = source["node_id"]
            item["section_index"] = source["section_index"]
            item["structural_heading_evidence"] = source["heading"]
    report["verification"] = {"structure_inventory": "completed", "layout": "not_performed"}
    if any(
        source["left_margin_emu"] != current.left_margin
        or source["right_margin_emu"] != current.right_margin
        or source["top_margin_emu"] != current.top_margin
        or source["bottom_margin_emu"] != current.bottom_margin
        for source, current in zip(report["structure_analysis"]["sections"], doc.sections)
    ):
        add_warning(report, None, "", "已调整节的页边距；保留段落或表格 XML 不代表分页和位置不变，布局尚未通过渲染验收。")

    finalize_module_status(report, context, template)
    if preserved_sections or report["structure_analysis"]["unresolved_sections"]:
        report["module_status"]["page"] = {
            "status": "detected", "count": len(report["preserved_sections"]),
            "action": "skipped" if len(preserved_sections) >= len(doc.sections) or report["structure_analysis"]["unresolved_sections"] else "partial",
            "note": "复杂版面或用户保护内容所在节未套用页边距；其他可处理节按模板执行。",
        }
    report_stats = build_report_stats(context["stats"], report, template)
    report["stats"] = report_stats

    safe_save_document(doc, output_path, overwrite=overwrite)
    return report_stats


def print_report(
    output_path: str,
    template: dict,
    stats: dict,
    report_path: str | None = None,
    overwritten: bool = False,
) -> None:
    """输出处理完成信息、模板信息和统计报告。"""
    abstract_count = (
        stats.get("abstract_title", 0)
        + stats.get("abstract_content", 0)
    )
    keywords_count = stats.get("keywords", 0)
    reference_count = stats.get("reference_title", 0) + stats.get("reference_item", 0)
    override_metadata = template.get("_override", {"enabled": False})

    print(f"格式修改完成：{output_path}")
    if override_metadata.get("enabled"):
        base_template = template.get("_base_template", {})
        print(f"使用基础模板：{base_template.get('name', '未命名模板')}")
        print(f"使用自定义覆盖：{override_metadata.get('path', '')}")
        print(f"覆盖字段数量：{len(override_metadata.get('overridden_fields', []))}")
        print(f"使用最终规则：{template.get('name', '未命名模板')}")
    else:
        print(f"使用模板：{template.get('name', '未命名模板')}")
        print(f"模板说明：{template.get('description', '无')}")
    if report_path:
        print(f"报告文件：{report_path}")
    if overwritten:
        print("提示：已覆盖已有输出文件。")
    print()
    print("处理统计：")
    print(f"- 论文标题：{stats.get('paper_title', 0)}")
    print(f"- 一级标题：{stats.get('heading_1', 0)}")
    print(f"- 二级标题：{stats.get('heading_2', 0)}")
    print(f"- 三级标题：{stats.get('heading_3', 0)}")
    print(f"- 表题：{stats.get('table_caption', 0)}")
    print(f"- 图题：{stats.get('figure_caption', 0)}")
    print(f"- 正文段落：{stats.get('body', 0)}")
    print(f"- 摘要相关段落：{abstract_count}")
    print(f"- 关键词段落：{keywords_count}")
    print(f"- 参考文献：{reference_count}")
    print(f"- 表格：{stats.get('table', 0)}")
    print(f"- 使用默认 body 回退次数：{stats.get('style_fallback_count', 0)}")
    print(f"- 警告数量：{stats.get('warning_count', 0)}")


def validate_template_command(template_arg: str) -> int:
    """只校验模板 JSON，不进入 Word 处理流程。"""
    template_path = resolve_template_path(template_arg)

    try:
        raw_template = read_template_json(template_path)
        normalized_template = normalize_format_rules(raw_template)
        normalization_warnings = normalized_template.get("_template_warnings", [])
        errors, validation_warnings = validate_format_rules(normalized_template)
        warnings = normalization_warnings + validation_warnings
    except TemplateError as exc:
        print(f"模板校验失败：{template_arg}")
        print("错误：")
        print(f"- {exc}")
        return 1

    if errors:
        print(f"模板校验失败：{template_arg}")
        print("错误：")
        for error in errors:
            print(f"- {error}")
        if warnings:
            print()
            print("警告：")
            for warning in warnings:
                print(f"- {warning}")
        return 1

    supported_style_count = sum(
        1
        for style_name in normalized_template.get("styles", {})
        if style_name in SUPPORTED_STYLE_TYPES
    )
    print(f"模板校验通过：{template_arg}")
    print(f"模板名称：{normalized_template.get('description') or normalized_template.get('name')}")
    print(f"模板版本：{normalized_template.get('version', '未声明')}")
    print(f"支持样式数量：{supported_style_count}")
    if warnings:
        print()
        print("警告：")
        for warning in warnings:
            print(f"- {warning}")
    return 0


def merge_rules_command(args) -> int:
    """只合并格式规则，不处理 Word。"""
    rules = get_final_format_rules(
        template_arg=args.template,
        override_path=args.override,
        print_warnings=True,
    )

    if args.export_final_rules:
        export_final_rules(rules, args.export_final_rules)
    if args.export_debug_rules:
        export_debug_rules(rules, args.export_debug_rules)

    override_metadata = rules.get("_override", {"enabled": False})
    print("格式规则合并完成。")
    print(f"基础模板：{rules.get('_base_template', {}).get('name', rules.get('name', 'default'))}")
    print(f"自定义覆盖：{override_metadata.get('path', '未使用') if override_metadata.get('enabled') else '未使用'}")
    print(f"最终规则：{args.export_final_rules or '未导出'}")
    print(f"调试规则：{args.export_debug_rules or '未导出'}")
    print(f"覆盖字段数量：{len(override_metadata.get('overridden_fields', []))}")
    print(f"警告数量：{len(override_metadata.get('warnings', []))}")
    return 0


def main(argv=None) -> int:
    args = parse_args(argv)

    if args.validate_template:
        return validate_template_command(args.validate_template)

    if args.merge_rules:
        return merge_rules_command(args)

    if not args.input_docx or not args.output_docx:
        raise UserFacingError(
            "请提供 input.docx 和 output.docx，或使用 --validate-template / --merge-rules。"
        )

    input_path = Path(args.input_docx)
    output_path = Path(args.output_docx)
    output_existed = output_path.exists()

    ensure_dependency_installed()
    validate_report_path(args.report, input_path, output_path)
    validate_paths(input_path, output_path, overwrite=args.overwrite)

    template = get_final_format_rules(
        template_arg=args.template,
        override_path=args.override,
        print_warnings=True,
    )
    if args.export_final_rules:
        export_final_rules(template, args.export_final_rules)
    if args.export_debug_rules:
        export_debug_rules(template, args.export_debug_rules)

    report = init_report(args.input_docx, args.output_docx, template)
    stats = format_document(
        input_path,
        output_path,
        template,
        report,
        overwrite=args.overwrite,
    )
    if args.report:
        save_report(report, args.report)

    print_report(
        args.output_docx,
        template,
        stats,
        args.report,
        overwritten=args.overwrite and output_existed,
    )
    return 0


if __name__ == "__main__":
    debug_enabled = "--debug" in sys.argv
    try:
        sys.exit(main())
    except UserFacingError as exc:
        print_error(str(exc))
        if debug_enabled:
            traceback.print_exc()
        sys.exit(1)
    except Exception:
        print_error("程序遇到未知异常。可使用 --debug 查看详细信息。")
        if debug_enabled:
            traceback.print_exc()
        sys.exit(1)
